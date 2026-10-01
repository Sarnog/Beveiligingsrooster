"""De urenformules van de Excel-export tegenover bereken_uren/uren_voor (Deel 1c, versie 1.6.0).

Per instelling een proefwerkboek: de echte bladen Lijsten, Feestdagen en Rekenhulp uit de
export, plus een blad 'Proef' met per regel een datum, begin/eind, opmerkingtijden en precies
de formules die ook in de weekbladen staan. 'formulas' rekent het werkboek uit (zie
tests/excel_formules.py); elke uitkomst moet gelijk zijn aan die van de app.
"""

import io
import random
from datetime import date, datetime, time, timedelta
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.extensions import db
from app.models import Dienstcode
from app.services import excel_export as ex
from app.services import instellingen
from app.services.feestdagen import feestdagen_in_periode
from app.services.rooster import UrenContext, uren_voor
from app.services.urenberekening import uren_uit_minuten

from .excel_formules import bereken

JAAR = 2026
ZATERDAG, ZONDAG, MAANDAG = date(2026, 3, 7), date(2026, 3, 8), date(2026, 3, 9)
KONINGSDAG = date(2026, 4, 27)  # maandag
TWEEDE_PAASDAG, HEMELVAART = date(2026, 4, 6), date(2026, 5, 14)


def _t(minuten: int) -> str:
    return f"{minuten // 60 % 24:02d}:{minuten % 60:02d}"


def _gevallen(seed: int) -> list[tuple]:
    """(datum, begin, eind, opm_begin, opm_eind): randgevallen plus een willekeurige steekproef."""
    rng = random.Random(seed)
    dagen = [MAANDAG, ZATERDAG, ZONDAG, KONINGSDAG, TWEEDE_PAASDAG, HEMELVAART, date(2026, 12, 26)]
    vast = [
        (MAANDAG, "07:15", "15:45", None, None),  # weekdag
        (ZATERDAG, "07:15", "15:45", None, None),  # zaterdag
        (ZONDAG, "07:15", "15:45", None, None),  # zondag
        (KONINGSDAG, "07:15", "15:45", None, None),  # feestdag (maandag)
        (date(2026, 12, 26), "07:15", "15:45", None, None),  # 2e kerstdag op zaterdag: hoogste factor
        (MAANDAG, "22:00", "06:30", None, None),  # nachtdienst
        (ZATERDAG, "22:00", "06:30", None, None),
        (MAANDAG, "08:00", "13:30", None, None),  # precies 5,5 uur (kommagetal exact)
        (MAANDAG, "13:00", "18:30", None, None),  # precies 5,5 uur, kommagetal net erboven (VBA)
        (MAANDAG, "06:00", "15:00", None, None),  # precies 9 uur
        (MAANDAG, "08:00", "17:00", None, None),  # 9 uur, kommagetal net erboven
        (ZATERDAG, "00:00", "01:15", None, None),  # half kwartier exact: 7,5 -> 8
        (ZATERDAG, "10:00", "11:15", None, None),  # 7,4999...: VBA rondt af naar 7
        (ZATERDAG, "12:00", "13:15", None, None),  # 7,5000...1: naar 8
        (ZATERDAG, "07:00", "15:15", None, None),  # 8:15 op zaterdag: half kwartier
        (ZATERDAG, "07:15", "15:30", None, None),
        (MAANDAG, "07:15", "13:00", "13:00", "17:00"),  # opmerkingtijden
        (MAANDAG, None, None, "13:00", "17:00"),  # alleen opmerkingtijden
        (MAANDAG, "07:15", None, None, None),  # alleen een begintijd: geen uren
        (MAANDAG, "07:15", "07:15", None, None),  # begin = eind: 0 uur
    ]
    for _ in range(140):
        begin = rng.randrange(1440)
        eind = (begin + rng.choice([rng.randrange(1, 1440), rng.randrange(60, 720, 5), 330, 540, 75, 495]))
        opm = rng.random() < 0.2
        opm_begin = rng.randrange(1440)
        vast.append((rng.choice(dagen), _t(begin), _t(eind),
                     _t(opm_begin) if opm else None, _t(opm_begin + rng.randrange(15, 300)) if opm else None))
    return vast


def _proefwerkboek(gevallen) -> bytes:
    """Lijsten/Feestdagen/Rekenhulp zoals de export ze maakt, plus het blad Proef."""
    eerste, laatste = date(JAAR - 1, 12, 29), date(JAAR + 1, 1, 3)
    feestdagen = feestdagen_in_periode(eerste, laatste)
    codes = Dienstcode.query.order_by(Dienstcode.nummer).all()
    toeslagen = instellingen.toeslagen()
    pauze = instellingen.pauze_instelling()
    staffel = tuple(pauze["regels"]) if pauze["aan"] else ()
    factoren = {1.0, toeslagen["factor_zaterdag"], toeslagen["factor_zondag"]} | (
        {toeslagen["factor_feestdag"]} if toeslagen["factor_feestdag"] else set())
    tabel = sorted({c for f in factoren for c in ex.correcties(f, staffel)})
    formules = ex._Formules(len(pauze["regels"]), len(tabel))

    boek = Workbook()
    proef = boek.active
    proef.title = "Proef"
    ex._lijsten(boek.create_sheet("Lijsten"), JAAR, [], codes, toeslagen, pauze, ex._Stijlen())
    ex._feestdagen(boek.create_sheet("Feestdagen"), feestdagen)
    ex._rekenhulp(boek.create_sheet("Rekenhulp"), tabel, pauze)
    for rij, (dag, begin, eind, opm_begin, opm_eind) in enumerate(gevallen, start=1):
        proef.cell(rij, 1).value = datetime.combine(dag, time())
        for kolom, tekst in ((2, begin), (3, eind), (7, opm_begin), (8, opm_eind)):
            proef.cell(rij, kolom).value = time.fromisoformat(tekst) if tekst else None
        factor = f"$D${rij}"
        ex._formule(proef, rij, 4, formules.factor(f"$A${rij}", len(feestdagen)))
        ex._formule(proef, rij, 5, formules.kwartieren_ruw(f"B{rij}", f"C{rij}", factor))
        ex._formule(proef, rij, 6, formules.uren(f"E{rij}", f"B{rij}", f"C{rij}", factor))
        ex._formule(proef, rij, 9, formules.kwartieren_ruw(f"G{rij}", f"H{rij}", factor))
        ex._formule(proef, rij, 10, formules.uren(f"I{rij}", f"G{rij}", f"H{rij}", factor))
        ex._formule(proef, rij, 11, formules.dagtotaal(f"F{rij}", f"J{rij}"))  # = de urencel van dienst 1
    uitvoer = io.BytesIO()
    boek.save(uitvoer)
    return uitvoer.getvalue()


def _vergelijk(gevallen) -> None:
    uitkomst = bereken(_proefwerkboek(gevallen))
    context = UrenContext(date(JAAR - 1, 12, 29), date(JAAR + 1, 1, 3))
    fout = []
    for rij, (dag, begin, eind, opm_begin, opm_eind) in enumerate(gevallen, start=1):
        dienst = SimpleNamespace(uren_handmatig=None, datum=dag, begin=begin, eind=eind,
                                 opmerking_begin=opm_begin, opmerking_eind=opm_eind)
        app = uren_voor(dienst, context)
        excel = uitkomst[f"Proef!K{rij}"]
        if (app is None and excel != "") or (app is not None and excel != app):
            fout.append(f"{dag} {begin}-{eind} opm {opm_begin}-{opm_eind}: app {app}, Excel {excel!r}")
    assert fout == [], f"{len(fout)} verschillen, bijv.: " + "; ".join(fout[:5])


@pytest.fixture
def codes(app):
    from app.services.voorbeeldpakket import laad_voorbeeldpakket

    laad_voorbeeldpakket()
    db.session.commit()


def test_standaardinstellingen(app, codes):
    _vergelijk(_gevallen(1))


def test_feestdagtoeslag_en_opmerkingtijden(app, codes):
    instellingen.schrijf("toeslag_feestdag", "2.5")
    instellingen.schrijf("toeslag_zaterdag", "1.25")
    instellingen.schrijf("opmerkingtijden_meetellen", "1")
    db.session.commit()
    _vergelijk(_gevallen(2))


def test_pauze_uit(app, codes):
    instellingen.schrijf("pauze", instellingen.pauze_json(False, [(5.5, 0.5)]))
    db.session.commit()
    _vergelijk(_gevallen(3))


def test_meerdere_pauzestaffels(app, codes):
    instellingen.schrijf("pauze", instellingen.pauze_json(True, [(4, 0.25), (5.5, 0.5), (9, 0.75)]))
    instellingen.schrijf("toeslag_feestdag", "2")
    db.session.commit()
    _vergelijk(_gevallen(4))


def test_correcties_alle_begintijden_bij_een_half_kwartier(app, codes):
    """Zaterdag 8:15 uur (7,75 × 1,5 × 4 = 46,5): voor élke begintijd gelijk aan de app."""
    gevallen = [(ZATERDAG, _t(b), _t(b + 495), None, None) for b in range(0, 1440, 11)]
    gevallen += [(MAANDAG, _t(b), _t(b + 330), None, None) for b in range(3, 1440, 17)]  # precies 5,5 uur
    _vergelijk(gevallen)
    # Zonder correctie zou een deel afwijken (anders test dit niets)
    afwijkend = [b for b in range(1440) if round(uren_uit_minuten(b, (b + 495) % 1440, 1.5) * 4) != 46]
    assert 0 < len(afwijkend) < 1440


@pytest.mark.parametrize("factor, staffel", [(1.5, ((5.5, 0.5),)), (2.0, ((4, 0.25), (9, 0.75))), (1.25, ())])
def test_snelle_berekening_gelijk_aan_uren_uit_minuten(factor, staffel):
    for minuten in (0, 75, 330, 495, 540, 545, 1439):
        assert ex.kwartieren_vba(minuten, factor, staffel) == [
            round(uren_uit_minuten(b, (b + minuten) % 1440, factor, staffel) * 4) for b in range(1440)]


def test_correcties_alleen_voor_halve_kwartieren_en_grenzen():
    for sleutel, tekst in ex.correcties(1.5, ((5.5, 0.5),)):
        minuten = sleutel % 1440
        assert sleutel // 1440 == 1500 and len(tekst) == 1440
        assert minuten % 10 == 5 or minuten == 330
    assert [s % 1440 for s, _ in ex.correcties(1.0, ((5.5, 0.5),))] == [330]
    assert ex.correcties(1.0, ()) == ()  # geen pauze, factor 1: niets te corrigeren


def test_correcties_gelden_niet_meer_na_een_andere_pauze_in_excel(app, codes):
    """Wijzig je de pauzeregels in Lijsten, dan rekent Excel exact (zonder correcties)."""
    inhoud = _proefwerkboek([(MAANDAG, "13:00", "18:30", None, None)])
    assert bereken(inhoud)["Proef!K1"] == 5.0  # zoals de app (VBA-kommagetal)
    gewijzigd = bereken(inhoud, {"Lijsten!O9": 0.25})
    assert gewijzigd["Rekenhulp!D1"] is False and gewijzigd["Proef!K1"] == 5.5  # exact: niet boven 5,5


def test_week_53_en_jaargrens_in_het_proefblad(app, codes):
    _vergelijk([(dag, "07:15", "15:45", None, None)
                for dag in (date(2026, 12, 31), date(2027, 1, 1), date(2025, 12, 29) + timedelta(days=5))])
