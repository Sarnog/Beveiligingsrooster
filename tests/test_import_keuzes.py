"""Excel-import: jaar kiezen (B2) en kiezen wat overschreven wordt (B3), versie 1.5.0.

De bestanden worden hier met fictieve gegevens opgebouwd in de opbouw van het oude
Excel-rooster (zie de docstring van app/services/excel_import.py).
"""

from datetime import date, datetime, time, timedelta

import openpyxl
import pytest

from app.extensions import db
from app.models import Contracturen, Dagopmerking, Dienst, Dienstcode, Logboek, Medewerker, SyncTaak, Vakantie
from app.services import instellingen
from app.services.excel_import import (
    MODUS_AANVULLEN,
    MODUS_GEDEELTELIJK,
    ImportFout,
    ImportKeuzes,
    effect,
    importeer,
    jaar_uit_naam,
    lees_bestand,
)
from app.services.kalender import maandag_van_week
from app.services.voorbeeldpakket import laad_voorbeeldpakket

from .test_import_backup import DAG_KOLOMMEN, keuzeformulier, upload

MEDEWERKERS = [("MXA", "Medewerker X", 1500), ("MYB", "Medewerker Y", 1200)]


def maak_rooster(pad: str, jaar: int | None, weken: dict, e2: int | None = None, toeslagen=(1.75, 2.5),
                 dagopmerkingen: dict | None = None, vakanties=(),
                 codes=((4, "VW Vroeg", "07:15", "15:45"),)):
    """Een fictief rooster. weken: {week: {naam: {dagindex: (code, begin, eind)}}}."""
    boek = openpyxl.Workbook()
    lijsten = boek.active
    lijsten.title = "Lijsten"
    for rij, (init, naam, uren) in enumerate(MEDEWERKERS, start=2):
        lijsten.cell(rij, 2, init)
        lijsten.cell(rij, 3, naam)
        lijsten.cell(rij, 4, uren)
    for rij, (nummer, oms, van, tot) in enumerate(codes, start=2):
        lijsten.cell(rij, 6, nummer)
        lijsten.cell(rij, 7, oms)
        lijsten.cell(rij, 8, time.fromisoformat(van) if van else None)
        lijsten.cell(rij, 9, time.fromisoformat(tot) if tot else None)
    lijsten["N2"], lijsten["N3"] = toeslagen
    blad = boek.create_sheet("Vakanties")
    blad.append(["Vakantie", "Van", "Tot"])
    for naam, van, tot in vakanties:
        blad.append([naam, datetime.combine(van, time()), datetime.combine(tot, time())])
    if e2 is not None or jaar is not None:
        boek.create_sheet("Kalender")["E2"] = e2 if e2 is not None else jaar
    for week, per_naam in weken.items():
        blad = boek.create_sheet(f"W{week}")
        maandag = maandag_van_week(jaar, week)
        for i, kolom in enumerate(DAG_KOLOMMEN):
            blad.cell(2, kolom, datetime.combine(maandag + timedelta(days=i), time()))
            tekst = (dagopmerkingen or {}).get(maandag + timedelta(days=i))
            if tekst:
                blad.cell(3, kolom, tekst)
        for n, (init, naam, _uren) in enumerate(MEDEWERKERS):
            basis = 4 + 4 * n
            blad.cell(basis, 2, naam)
            blad.cell(6 + 2 * n, 28, init)
            for dag, (code, begin, eind) in per_naam.get(naam, {}).items():
                kolom = DAG_KOLOMMEN[dag]
                blad.cell(6 + 2 * n, 29 + dag, code)
                blad.cell(basis + 2, kolom, "VW Vroeg")
                blad.cell(basis + 3, kolom, time.fromisoformat(begin))
                blad.cell(basis + 3, kolom + 1, time.fromisoformat(eind))
    boek.save(pad)


@pytest.fixture
def basis(klaar):
    """Bestaande medewerkers X en Y met diensten in 2026 (week 10 en 12)."""
    laad_voorbeeldpakket()
    x = Medewerker(naam="Medewerker X", initialen="MXA", volgorde=1)
    y = Medewerker(naam="Medewerker Y", initialen="MYB", volgorde=2)
    db.session.add_all([x, y])
    db.session.flush()
    for medewerker in (x, y):
        for dag in (date(2026, 3, 2), date(2026, 3, 3), date(2026, 3, 16)):  # W10 ma/di, W12 ma
            db.session.add(Dienst(medewerker_id=medewerker.id, datum=dag, dienstnaam_override="Oud",
                                  begin="08:00", eind="16:00", opmerking_tekst=""))
    db.session.commit()
    return {"x": x.id, "y": y.id}


def _diensten(medewerker_id, van=None, tot=None):
    db.session.expire_all()
    query = Dienst.query.filter_by(medewerker_id=medewerker_id)
    if van:
        query = query.filter(Dienst.datum >= van)
    if tot:
        query = query.filter(Dienst.datum <= tot)
    return {(d.datum, d.volgnummer): (d.dienstnaam, d.begin, d.eind) for d in query.all()}


def _bestand_2026(tmp_path, **extra):
    """W10 van 2026: X werkt ma en wo (code 4), Y werkt ma (code 4, eigen eindtijd)."""
    pad = str(tmp_path / "rooster-2026.xlsx")
    maak_rooster(pad, 2026, {10: {"Medewerker X": {0: (4, "07:15", "15:45"), 2: (4, "07:15", "15:45")},
                                  "Medewerker Y": {0: (4, "07:15", "18:00")}}}, **extra)
    return pad


# ---------------------------------------------------------------------------
# B2 · Jaar kiezen
# ---------------------------------------------------------------------------

def test_b2_import_2027_laat_2026_volledig_intact(app, basis, tmp_path):
    voor = {k: _diensten(v) for k, v in basis.items()}
    pad = str(tmp_path / "volgend.xlsx")
    maak_rooster(pad, 2027, {10: {"Medewerker X": {0: (4, "07:15", "15:45")}},
                             1: {"Medewerker Y": {0: (4, "07:15", "15:45")}}})
    importeer(lees_bestand(pad, 2027))
    for sleutel, mid in basis.items():
        assert _diensten(mid, tot=date(2027, 1, 3)) == voor[sleutel]  # 2026 (ISO) precies gelijk
    assert (date(2027, 3, 8), 1) in _diensten(basis["x"])
    assert (date(2027, 1, 4), 1) in _diensten(basis["y"])


def test_b2_gekozen_jaar_wijkt_af_van_e2_en_bladen_worden_overgeslagen(app, basis, tmp_path):
    pad = _bestand_2026(tmp_path)
    plan = lees_bestand(pad, 2027)  # bestand is van 2026
    assert plan.jaar == 2027 and plan.jaar_in_bestand == 2026
    assert any("Kalender!E2" in w and "2026" in w for w in plan.waarschuwingen)
    assert any("W10 overgeslagen" in w for w in plan.waarschuwingen)
    assert plan.diensten == [] and plan.weken == []


def test_b2_zonder_jaar_en_zonder_e2_geen_stille_terugval(app, klaar, tmp_path):
    pad = str(tmp_path / "zonder-jaar.xlsx")
    boek = openpyxl.Workbook()
    boek.active.title = "Lijsten"
    boek.save(pad)
    with pytest.raises(ImportFout, match="Kies voor welk jaar"):
        lees_bestand(pad)
    assert lees_bestand(pad, 2027).jaar == 2027


def test_b2_jaar_uit_bestandsnaam():
    assert jaar_uit_naam("Rooster 2027.xlsm") == 2027
    assert jaar_uit_naam("rooster-2027-v2.xlsx") == 2027
    assert jaar_uit_naam("rooster12345.xlsx") is None and jaar_uit_naam("Rooster.xlsm") is None


def test_b2_jaarstap_voorstel_uit_e2_of_bestandsnaam(app, als_beheerder, tmp_path):
    pad = _bestand_2026(tmp_path)
    upload(als_beheerder, pad, jaar=None)
    assert 'name="jaar" value="2026"' in als_beheerder.get("/beheer/importeren").data.decode()
    # Zonder E2: het jaar uit de bestandsnaam, anders leeg (nooit stil het huidige jaar)
    zonder = str(tmp_path / "zonder.xlsx")
    maak_rooster(zonder, None, {})
    upload(als_beheerder, zonder, jaar=None, naam="Rooster 2028.xlsx")
    assert 'name="jaar" value="2028"' in als_beheerder.get("/beheer/importeren").data.decode()
    upload(als_beheerder, zonder, jaar=None, naam="Rooster.xlsx")
    pagina = als_beheerder.get("/beheer/importeren").data.decode()
    assert 'name="jaar" value=""' in pagina and "required" in pagina.split('name="jaar"')[1][:200]
    # Zonder jaar kan het voorbeeld niet
    assert als_beheerder.post("/beheer/importeren/jaar", data={"jaar": ""}).status_code == 302
    antwoord = als_beheerder.get("/beheer/importeren/voorbeeld")
    assert antwoord.headers["Location"].endswith("/beheer/importeren")


def test_b2_globale_gegevens_niet_stil_overschreven_voor_ander_jaar(app, basis, tmp_path):
    db.session.add(Contracturen(medewerker_id=basis["x"], jaar=2026, uren=1000))
    instellingen.schrijf("toeslag_zaterdag", "1.5")
    db.session.commit()
    omschrijving = Dienstcode.query.filter_by(nummer=4).one().omschrijving
    pad = str(tmp_path / "volgend.xlsx")
    maak_rooster(pad, 2027, {10: {"Medewerker X": {0: (4, "07:15", "15:45")}}},
                 vakanties=[("Kerst 2026", date(2026, 12, 19), date(2027, 1, 3)),
                            ("Zomer 2025", date(2025, 7, 1), date(2025, 8, 1)),
                            ("Mei 2027", date(2027, 4, 24), date(2027, 5, 2))],
                 codes=((4, "Andere naam", "06:00", "14:00"), (77, "Nieuwe code", "09:00", "17:00")))
    plan = lees_bestand(pad, 2027)
    keuzes = ImportKeuzes.standaard(plan)
    assert keuzes.toeslagen is False  # toeslagen gelden voor alle jaren
    importeer(plan, keuzes)
    assert instellingen.lees("toeslag_zaterdag") == "1.5"
    x = db.session.get(Medewerker, basis["x"])
    assert {c.jaar: c.uren for c in x.contracturen} == {2026: 1000, 2027: 1500}  # alleen 2027
    assert sorted(v.naam for v in Vakantie.query.all()) == ["Kerst 2026", "Mei 2027"]  # niet 2025
    assert Dienstcode.query.filter_by(nummer=4).one().omschrijving == omschrijving  # nooit gewijzigd
    assert Dienstcode.query.filter_by(nummer=77).one().omschrijving == "Nieuwe code"  # wel erbij


def test_b2_onderdelen_niet_overnemen(app, basis, tmp_path):
    pad = _bestand_2026(tmp_path, codes=((77, "Nieuwe code", "09:00", "17:00"),),
                        vakanties=[("Mei", date(2026, 4, 25), date(2026, 5, 3))])
    plan = lees_bestand(pad, 2026)
    importeer(plan, ImportKeuzes(contracturen=False, toeslagen=False, vakanties=False, codes=False))
    assert Contracturen.query.count() == 0 and Vakantie.query.count() == 0
    assert Dienstcode.query.filter_by(nummer=77).first() is None


# ---------------------------------------------------------------------------
# B3 · Overschrijfmodus
# ---------------------------------------------------------------------------

MA, DI, WO = date(2026, 3, 2), date(2026, 3, 3), date(2026, 3, 4)
W12 = date(2026, 3, 16)


def test_b3_alles_vervangt_alleen_weken_uit_het_bestand(app, basis, tmp_path):
    importeer(lees_bestand(_bestand_2026(tmp_path), 2026))
    x = _diensten(basis["x"])
    assert set(x) == {(MA, 1), (WO, 1), (W12, 1)}  # di weg, wo nieuw, week 12 (geen blad) blijft
    assert x[(MA, 1)][0] == "VW Vroeg" and x[(W12, 1)][0] == "Oud"
    assert _diensten(basis["y"])[(MA, 1)][2] == "18:00"


def test_b3_alleen_persoon_x_laat_y_ongemoeid(app, basis, tmp_path):
    voor_y = _diensten(basis["y"])
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    importeer(plan, ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",)))
    assert _diensten(basis["y"]) == voor_y
    assert set(_diensten(basis["x"])) == {(MA, 1), (WO, 1), (W12, 1)}


def test_b3_alleen_periode_laat_de_rest_staan(app, basis, tmp_path):
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    importeer(plan, ImportKeuzes(modus=MODUS_GEDEELTELIJK, van=DI, tot=WO))
    for mid in basis.values():
        diensten = _diensten(mid)
        assert diensten[(MA, 1)][0] == "Oud"  # maandag valt buiten de periode
        assert (DI, 1) not in diensten  # dinsdag binnen de periode: niet in het bestand -> weg
    assert _diensten(basis["x"])[(WO, 1)][0] == "VW Vroeg"


def test_b3_persoon_en_periode(app, basis, tmp_path):
    voor_y = _diensten(basis["y"])
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    importeer(plan, ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",), van=MA, tot=MA))
    x = _diensten(basis["x"])
    assert x[(MA, 1)][0] == "VW Vroeg" and x[(DI, 1)][0] == "Oud" and (WO, 1) not in x
    assert _diensten(basis["y"]) == voor_y


def test_b3_alleen_lege_dagen_aanvullen(app, basis, tmp_path):
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    uitkomst = effect(plan, ImportKeuzes(modus=MODUS_AANVULLEN))
    assert uitkomst.totaal.overgeslagen == 2 and uitkomst.totaal.nieuw == 1
    assert uitkomst.totaal.vervangen == uitkomst.totaal.verwijderd == 0
    importeer(plan, ImportKeuzes(modus=MODUS_AANVULLEN))
    x = _diensten(basis["x"])
    assert x[(MA, 1)][0] == "Oud" and x[(DI, 1)][0] == "Oud" and x[(WO, 1)][0] == "VW Vroeg"
    assert _diensten(basis["y"])[(MA, 1)][0] == "Oud"


def test_b3_herhaalde_import_is_idempotent(app, basis, tmp_path):
    pad = _bestand_2026(tmp_path, dagopmerkingen={WO: "Teamoverleg"})
    keuzes = ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",), van=MA, tot=WO)
    importeer(lees_bestand(pad, 2026), keuzes)
    na_eerste = {mid: _diensten(mid) for mid in basis.values()}
    versies = {d.id: d.versie for d in Dienst.query.all()}
    uitkomst = effect(lees_bestand(pad, 2026), keuzes)
    assert (uitkomst.totaal.nieuw, uitkomst.totaal.vervangen, uitkomst.totaal.verwijderd) == (0, 0, 0)
    resultaat = importeer(lees_bestand(pad, 2026), keuzes)
    assert resultaat["diensten"] == resultaat["vervangen"] == resultaat["verwijderd"] == 0
    assert {mid: _diensten(mid) for mid in basis.values()} == na_eerste
    assert {d.id: d.versie for d in Dienst.query.all()} == versies  # niets aangeraakt


def test_b3_dagopmerkingen_volgen_periode_en_persoonsfilter(app, basis, tmp_path):
    pad = _bestand_2026(tmp_path, dagopmerkingen={MA: "Opmerking ma", WO: "Opmerking wo"})
    plan = lees_bestand(pad, 2026)
    alleen_x = ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",))
    assert effect(plan, alleen_x).dagopmerkingen == []  # bij een persoonsfilter standaard ongemoeid
    met = ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",), van=WO,
                       dagopmerkingen_bij_selectie=True)
    assert [d for d, _, _ in effect(plan, met).dagopmerkingen] == [WO]  # alleen binnen de periode
    importeer(plan, met)
    assert [(d.datum, d.tekst) for d in Dagopmerking.query.all()] == [(WO, "Opmerking wo")]


def test_b3_ongeldige_keuzes_geweigerd(app, basis, tmp_path):
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    for keuzes in (ImportKeuzes(modus="onzin"),
                   ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Bestaat niet",)),
                   ImportKeuzes(modus=MODUS_GEDEELTELIJK, van=date(2027, 6, 1)),
                   ImportKeuzes(modus=MODUS_GEDEELTELIJK, van=WO, tot=MA)):
        with pytest.raises(ImportFout):
            importeer(plan, keuzes)
    assert _diensten(basis["x"])[(MA, 1)][0] == "Oud"


def test_b3_logboek_met_modus_jaar_medewerkers_periode_en_aantallen(app, basis, tmp_path):
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    importeer(plan, ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",), van=MA, tot=WO))
    regel = Logboek.query.filter_by(actie="Excel-import", veld="").one()
    for stukje in ("jaar 2026", "modus gedeeltelijk", "Medewerker X", "02-03-2026 t/m 04-03-2026",
                   "diensten: 1", "vervangen: 1", "verwijderd: 1"):
        assert stukje in regel.details, stukje
    # En per dienst een regel met oud en nieuw
    per_dienst = Logboek.query.filter(Logboek.actie == "Excel-import", Logboek.veld == "dienst").all()
    assert len(per_dienst) == 3


def test_b3_google_alleen_geraakte_medewerkers_en_afspraken_blijven_tot_de_worker(app, basis, tmp_path):
    for mid in basis.values():
        medewerker = db.session.get(Medewerker, mid)
        medewerker.agenda_modus, medewerker.agenda_id = "B", "agenda"
    dinsdag_x = Dienst.query.filter_by(medewerker_id=basis["x"], datum=DI).one()
    dinsdag_x.google_event_id = "event-di"
    db.session.commit()
    plan = lees_bestand(_bestand_2026(tmp_path), 2026)
    importeer(plan, ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",)))
    taken = SyncTaak.query.all()
    assert {t.medewerker_id for t in taken} == {basis["x"]}  # Y niet geraakt: geen sync
    # De verwijderde dinsdag met een afspraak blijft (leeg) staan tot de worker hem opruimt
    dinsdag = Dienst.query.filter_by(medewerker_id=basis["x"], datum=DI).one()
    assert dinsdag.is_leeg and dinsdag.google_event_id == "event-di"
    assert any(t.soort == "dag" and t.datum == DI for t in taken)  # buiten de sync-periode: eigen taak


# ---------------------------------------------------------------------------
# B3 · Het scherm: zelfde keuzes bevestigen, opnieuw controleren bij opslaan
# ---------------------------------------------------------------------------

def test_b3_scherm_toont_effect_per_modus(app, als_beheerder, basis, tmp_path):
    upload(als_beheerder, _bestand_2026(tmp_path))
    pagina = als_beheerder.get("/beheer/importeren/voorbeeld").data.decode()
    assert "Overschrijfmodus" in pagina and "rooster voor 2026" in pagina
    keuzes = ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker X",), toeslagen=True)
    als_beheerder.post("/beheer/importeren/voorbeeld", data=keuzeformulier(keuzes, actie="voorbeeld"))
    pagina = als_beheerder.get("/beheer/importeren/voorbeeld").data.decode()
    tabel = pagina.split("data-effect")[1].split("</table>")[0]
    assert "Medewerker X" in tabel and "Medewerker Y" not in tabel
    assert 'value="gedeeltelijk" checked' in pagina


def test_b3_bevestigen_alleen_met_ongewijzigde_keuzes(app, als_beheerder, basis, tmp_path):
    upload(als_beheerder, _bestand_2026(tmp_path))
    als_beheerder.get("/beheer/importeren/voorbeeld")  # voorbeeld met de standaardkeuzes
    anders = ImportKeuzes(modus=MODUS_GEDEELTELIJK, medewerkers=("Medewerker Y",), toeslagen=True)
    antwoord = als_beheerder.post("/beheer/importeren/voorbeeld", data=keuzeformulier(anders),
                                  follow_redirects=True)
    assert "keuzes zijn gewijzigd" in antwoord.data.decode()
    assert _diensten(basis["x"])[(DI, 1)][0] == "Oud"  # niets geïmporteerd
    # Nu met precies de getoonde keuzes: wel importeren
    antwoord = als_beheerder.post("/beheer/importeren/voorbeeld", data=keuzeformulier(anders))
    assert antwoord.status_code == 302 and "/kalender" in antwoord.headers["Location"]
    assert _diensten(basis["x"])[(DI, 1)][0] == "Oud" and (DI, 1) not in _diensten(basis["y"])


def test_b3_ongeldige_periode_via_scherm(app, als_beheerder, basis, tmp_path):
    upload(als_beheerder, _bestand_2026(tmp_path))
    als_beheerder.get("/beheer/importeren/voorbeeld")
    data = keuzeformulier(ImportKeuzes(modus=MODUS_GEDEELTELIJK), actie="voorbeeld")
    data["van"] = "2025-01-01"
    antwoord = als_beheerder.post("/beheer/importeren/voorbeeld", data=data, follow_redirects=True)
    assert "binnen 2026" in antwoord.data.decode()
