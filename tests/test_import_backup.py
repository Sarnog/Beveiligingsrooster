"""Excel-import, back-ups terugzetten en handmatige uren (fase 4).

Het testbestand wordt hier met fictieve gegevens opgebouwd in de opbouw van het
oude Excel-rooster. Er zit dus geen echte roosterdata in de repository.
"""

import io
import os
import shutil
from datetime import date, datetime, time

import openpyxl
import pytest

from app.extensions import db
from app.models import Dagopmerking, Dienst, Dienstcode, Medewerker, Vakantie
from app.services import backup, instellingen
from app.services.excel_import import importeer, lees_bestand
from app.services.voorbeeldpakket import laad_voorbeeldpakket
from app.services.weekrooster import Wijziging, wijzig_cellen

DAG_KOLOMMEN = [4, 7, 10, 13, 16, 19, 22]


def maak_testbestand(pad: str) -> None:
    """Een klein, fictief 'oud rooster' met twee medewerkers en week 10 van 2026."""
    boek = openpyxl.Workbook()
    lijsten = boek.active
    lijsten.title = "Lijsten"
    lijsten["B1"], lijsten["C1"], lijsten["D1"] = "Initialen", "Personeel", "Contract"
    for rij, (init, naam, uren) in enumerate([("MVA", "Medewerker Vijf A", 1659),
                                               ("MVB", "Medewerker Vijf B", 1295)], start=2):
        lijsten.cell(rij, 2, init)
        lijsten.cell(rij, 3, naam)
        lijsten.cell(rij, 4, uren)
    codes = [(4, "VW Vroeg", time(7, 15), time(15, 45), 8), (10, "Bapo", None, None, None),
             (15, 15, None, None, None), (42, "Nieuwe dienst", time(9, 0), time(17, 0), 7.5)]
    for rij, (nummer, oms, van, tot, uren) in enumerate(codes, start=2):
        lijsten.cell(rij, 6, nummer)
        lijsten.cell(rij, 7, oms)
        lijsten.cell(rij, 8, van)
        lijsten.cell(rij, 9, tot)
        lijsten.cell(rij, 10, uren)
    lijsten["N2"], lijsten["N3"] = 1.5, 2

    vakanties = boek.create_sheet("Vakanties")
    vakanties.append(["Vakantie", "Datum van", "Datum tot", "Aantal werkdagen"])
    vakanties.append(["Voorjaarsvakantie", datetime(2026, 2, 16), datetime(2026, 2, 22), 5])

    kalender = boek.create_sheet("Kalender")
    kalender["E2"] = 2026

    week = boek.create_sheet("W10")
    maandag = date(2026, 3, 2)
    for i, kolom in enumerate(DAG_KOLOMMEN):
        week.cell(2, kolom, datetime(2026, 3, 2 + i))
    week.cell(3, DAG_KOLOMMEN[4], "Eigen dagtekst")
    # Blok 0: medewerker A
    week["B4"], week["Z4"], week["AB6"] = "Medewerker Vijf A", 1659, "MVA"
    # ma: code 4 met standaardtijden, en een opmerking met tijden
    week.cell(6, 29, 4)
    week.cell(6, 4, "VW Vroeg")
    week.cell(7, 4, time(7, 15))
    week.cell(7, 5, time(15, 45))
    week.cell(7, 6, 8)
    week.cell(4, 4, "BV")
    week.cell(5, 4, time(13, 30))
    week.cell(5, 5, time(15, 45))
    # di: code 4 met afwijkende eindtijd (handmatig)
    week.cell(6, 30, 4)
    week.cell(6, 7, "VW Vroeg")
    week.cell(7, 7, time(7, 15))
    week.cell(7, 8, time(16, 0))
    week.cell(7, 9, 8.25)
    # wo: vrije dienst zonder code en tijden, met zelf getypte uren
    week.cell(6, 10, "Cursus extern")
    week.cell(7, 12, 8)
    # do: dienst tot 13:00, daarna training 13:00-17:00; uren van de hele dag met de hand getypt
    week.cell(6, 32, 4)
    week.cell(6, 13, "VW Vroeg")
    week.cell(7, 13, time(7, 15))
    week.cell(7, 14, time(13, 0))
    week.cell(7, 15, 9.25)
    week.cell(4, 13, "Training")
    week.cell(5, 13, time(13, 0))
    week.cell(5, 14, time(17, 0))
    # za: blanco-code 15 (als datum 1900-01-15) = leeg
    week.cell(6, 34, datetime(1900, 1, 15))
    week.cell(6, 19, datetime(1900, 1, 15))
    week["Z6"] = 33.5
    # Blok 1: medewerker B met een code die nog niet bestaat
    week["B8"], week["AB8"] = "Medewerker Vijf B", "MVB"
    week.cell(8, 29, 42)
    week.cell(10, 4, "Nieuwe dienst")
    week.cell(11, 4, time(9, 0))
    week.cell(11, 5, time(17, 0))
    week.cell(11, 6, 7.5)
    week["Z10"] = 99  # bewust fout weektotaal: moet gemeld worden

    # Bladen met wachtwoorden worden genegeerd
    beveiliging = boek.create_sheet("Beveiliging")
    beveiliging.append(["Gebruiker", "Wachtwoord"])
    beveiliging.append(["iemand", "geheim"])
    del maandag
    boek.save(pad)


def upload(client, pad: str, jaar: int | None = 2026, naam: str = "Rooster.xlsm"):
    """Bestand uploaden via het scherm (stap 1), eventueel meteen met 'Rooster voor jaar'."""
    with open(pad, "rb") as f:
        return client.post("/beheer/importeren", data={
            "bestand": (io.BytesIO(f.read()), naam), "jaar": "" if jaar is None else str(jaar)},
            content_type="multipart/form-data")


def keuzeformulier(keuzes=None, jaar: int = 2026, actie: str = "importeren", bevestig: bool = True) -> dict:
    """De velden van het droogloopformulier voor deze keuzes (standaard: de standaardkeuzes)."""
    from app.services import klok
    from app.services.excel_import import ImportKeuzes

    keuzes = keuzes or ImportKeuzes(toeslagen=jaar == klok.vandaag().year)
    data = {"modus": keuzes.modus, "medewerkers": list(keuzes.medewerkers), "actie": actie,
            "van": keuzes.van.isoformat() if keuzes.van else "",
            "tot": keuzes.tot.isoformat() if keuzes.tot else ""}
    if keuzes.dagopmerkingen_bij_selectie:
        data["dagopmerkingen_bij_selectie"] = "1"
    for onderdeel in ("contracturen", "toeslagen", "vakanties", "codes"):
        data[onderdeel] = "ja" if getattr(keuzes, onderdeel) else "nee"
    if bevestig:
        data["bevestig"] = "1"
    return data


@pytest.fixture
def bestand(tmp_path):
    pad = str(tmp_path / "oud_rooster.xlsx")
    maak_testbestand(pad)
    return pad


def test_droogloop(app, klaar, bestand):
    plan = lees_bestand(bestand)
    assert plan.jaar == 2026 and plan.weken == [10]
    assert [m.naam for m in plan.medewerkers] == ["Medewerker Vijf A", "Medewerker Vijf B"]
    assert [c.nummer for c in plan.codes] == [4, 10, 42]  # 15 = blanco
    assert len(plan.diensten) == 5 and plan.afwijkend == 2
    # Alleen het bewust foute weektotaal wijkt af; de met de hand getypte 9,25 telt mee
    assert plan.weektotaal_verschillen() == ["W10 Medewerker Vijf B: Excel 99.00, nieuw 7.50"]
    assert len(plan.handmatige_uren()) == 1 and "9.25 (overgenomen)" in plan.handmatige_uren()[0]
    assert Dienst.query.count() == 0  # droogloop schrijft niets


def test_definitief_importeren(app, klaar, bestand):
    laad_voorbeeldpakket()  # code 4 en 10 bestaan al; 42 komt erbij
    resultaat = importeer(lees_bestand(bestand))
    assert resultaat["medewerkers"] == 2 and resultaat["codes"] == 1 and resultaat["diensten"] == 5
    a = Medewerker.query.filter_by(initialen="MVA").one()
    assert a.contracturen_voor(2026) == 1659

    def dag(d):
        return Dienst.query.filter_by(medewerker_id=a.id, datum=date(2026, 3, d)).one()

    ma, di, wo = dag(2), dag(3), dag(4)
    assert (ma.dienstnaam, ma.uren_berekend, ma.tijden_handmatig) == ("VW Vroeg", 8.0, False)
    assert (ma.opmerking_tekst, ma.opmerking_begin, ma.opmerking_eind) == ("BV", "13:30", "15:45")
    assert (di.eind, di.tijden_handmatig, di.uren_berekend) == ("16:00", True, 8.25)
    assert (wo.dienstnaam, wo.uren_handmatig, wo.uren_berekend) == ("Cursus extern", 8.0, 8.0)
    do = dag(5)
    assert (do.eind, do.uren_handmatig, do.uren_berekend) == ("13:00", 9.25, 9.25)
    assert ma.uren_handmatig is None and di.uren_handmatig is None  # gewoon berekend
    assert Dienst.query.filter_by(medewerker_id=a.id, datum=date(2026, 3, 7)).first() is None
    assert Dienstcode.query.filter_by(nummer=42).one().std_begin == "09:00"
    assert Vakantie.query.count() == 1
    assert Dagopmerking.query.one().tekst == "Eigen dagtekst"
    assert instellingen.lees_float("toeslag_zaterdag") == 1.5
    # Nogmaals importeren overschrijft de week (geen dubbele diensten)
    importeer(lees_bestand(bestand))
    assert Dienst.query.count() == 5


def test_import_via_scherm(app, als_beheerder, bestand):
    # Sinds 1.5.0: na het uploaden eerst 'Rooster voor jaar' kiezen (voorstel uit Kalender!E2)
    antwoord = upload(als_beheerder, bestand, jaar=None)
    assert antwoord.headers["Location"].endswith("/beheer/importeren")
    pagina = als_beheerder.get("/beheer/importeren").data.decode()
    assert 'name="jaar" value="2026"' in pagina
    antwoord = als_beheerder.post("/beheer/importeren/jaar", data={"jaar": "2026"})
    assert antwoord.headers["Location"].endswith("/beheer/importeren/voorbeeld")
    pagina = als_beheerder.get("/beheer/importeren/voorbeeld").data.decode()
    assert "Voorbeeld (droogloop)" in pagina and "Excel 99.00" in pagina
    als_beheerder.post("/beheer/importeren/voorbeeld", data=keuzeformulier())
    assert Dienst.query.count() == 5
    # Het geüploade bestand is weer weg, en er is vooraf een back-up gemaakt
    assert os.listdir(os.path.join(app.config["DATA_MAP"], "import")) == []
    assert any("voor-import" in b["naam"] for b in backup.lijst_backups())


def test_verkeerd_bestand(app, als_beheerder, tmp_path):
    antwoord = als_beheerder.post("/beheer/importeren", data={
        "bestand": (io.BytesIO(b"geen excel"), "iets.xlsm")}, content_type="multipart/form-data",
        follow_redirects=True)
    assert "kan niet gelezen worden" in antwoord.data.decode()


# ---------- Handmatige uren en vrije dienstnaam in het raster ----------

def test_handmatige_uren_en_vrije_dienstnaam(app, klaar):
    laad_voorbeeldpakket()
    medewerker = Medewerker(naam="Medewerker A", initialen="TSA")
    db.session.add(medewerker)
    db.session.commit()
    zaterdag = date(2026, 3, 7)
    bijgewerkt, fouten = wijzig_cellen([
        Wijziging(medewerker.id, zaterdag, "dienstnaam", "Cursus extern"),
        Wijziging(medewerker.id, zaterdag, "uren", "7,5")])
    assert fouten == []
    gegevens = bijgewerkt[f"{medewerker.id}|{zaterdag.isoformat()}"]
    assert gegevens["dienstnaam"] == "Cursus extern" and gegevens["uren"] == "7,50"
    assert gegevens["uren_handmatig"] is True  # geen zaterdagtoeslag op zelf getypte uren
    _, fouten = wijzig_cellen([Wijziging(medewerker.id, zaterdag, "uren", "abc")])
    assert "geen aantal uren" in fouten[0]["melding"]
    # Een code invoeren zet alles weer automatisch
    wijzig_cellen([Wijziging(medewerker.id, zaterdag, "code", "4")])
    dienst = Dienst.query.one()
    assert dienst.uren_handmatig is None and dienst.uren_berekend == 12.0


# ---------- Back-ups ----------

@pytest.fixture
def gemigreerd(app):
    """Database via de echte migraties (nodig voor een geldige back-up)."""
    from flask_migrate import upgrade

    db.drop_all()
    upgrade(directory=os.path.join(os.path.dirname(__file__), "..", "migrations"))
    return app


def test_backup_terugzetten(gemigreerd, client):
    from .conftest import login, maak_gebruiker

    instellingen.schrijf("setup_voltooid", "1")
    db.session.commit()
    maak_gebruiker("beheerder", "beheerder")
    login(client, "beheerder")
    db.session.add(Medewerker(naam="Medewerker Oud", initialen="OUD"))
    db.session.commit()
    client.post("/beheer/backups/maken")
    naam = backup.lijst_backups()[0]["naam"]
    assert client.get(f"/beheer/backups/download/{naam}").status_code == 200
    assert client.get("/beheer/backups/download/..%2Frooster.db").status_code == 404

    # Na de back-up een medewerker toevoegen; terugzetten haalt die weer weg
    db.session.add(Medewerker(naam="Medewerker Nieuw", initialen="NW"))
    db.session.commit()
    antwoord = client.post("/beheer/backups/terugzetten", data={"naam": naam, "bevestig": "1"})
    assert antwoord.status_code == 302
    db.session.remove()
    assert {m.initialen for m in Medewerker.query.all()} == {"OUD"}
    assert any("voor-terugzetten" in b["naam"] for b in backup.lijst_backups())


def test_backup_verwijderen(gemigreerd, client):
    from app.models import Logboek

    from .conftest import login, maak_gebruiker

    instellingen.schrijf("setup_voltooid", "1")
    db.session.commit()
    maak_gebruiker("beheerder", "beheerder")
    login(client, "beheerder")
    client.post("/beheer/backups/maken")
    naam = backup.lijst_backups()[0]["naam"]
    assert "Verwijderen" in client.get("/beheer/backups").data.decode()
    antwoord = client.post("/beheer/backups/verwijderen", data={"naam": naam}, follow_redirects=True)
    assert f"Back-up {naam} is verwijderd" in antwoord.data.decode()
    assert backup.lijst_backups() == []
    assert Logboek.query.filter_by(actie="Back-up verwijderd").count() == 1
    # Alleen eigen back-upnamen: nooit de database zelf of iets buiten de back-upmap
    for onzin in ("../rooster.db", "rooster.db", naam):
        antwoord = client.post("/beheer/backups/verwijderen", data={"naam": onzin}, follow_redirects=True)
        assert "Onbekende back-up" in antwoord.data.decode()
    assert os.path.exists(backup.database_pad())


def test_ongeldige_backup_upload(gemigreerd, client):
    from .conftest import login, maak_gebruiker

    instellingen.schrijf("setup_voltooid", "1")
    db.session.commit()
    maak_gebruiker("beheerder", "beheerder")
    login(client, "beheerder")
    antwoord = client.post("/beheer/backups/terugzetten", data={
        "bevestig": "1", "bestand": (io.BytesIO(b"onzin"), "x.db")},
        content_type="multipart/form-data", follow_redirects=True)
    assert "geen geldige" in antwoord.data.decode() or "geen back-up" in antwoord.data.decode()


def test_oude_backup_zonder_volgnummer_terugzetten(gemigreerd, tmp_path):
    """Een back-up van 1.3.0 (databaseversie 0005, zonder volgnummer) kan nog terug."""
    from flask_migrate import upgrade

    from app import create_app

    from .conftest import TestConfig

    # Een database zoals 1.3.0 hem had, in een eigen map
    oud_map = tmp_path / "oud"
    oud_map.mkdir()
    oude_app = create_app(TestConfig(str(oud_map)))
    with oude_app.app_context():
        upgrade(directory=os.path.join(os.path.dirname(__file__), "..", "migrations"), revision="0005")
        with db.engine.begin() as verbinding:
            verbinding.exec_driver_sql(
                "INSERT INTO medewerker (id, naam, initialen, functie_opmerking, volgorde, email, "
                "agenda_modus, agenda_id, agenda_laatste_fout, ics_token) "
                "VALUES (1, 'Medewerker Oud', 'OUD', '', 1, '', '', '', '', '')")
            verbinding.exec_driver_sql(
                "INSERT INTO dienst (datum, medewerker_id, dienstnaam_override, begin, eind, "
                "tijden_handmatig, uren_berekend, opmerking_tekst, google_event_id, versie, gewijzigd_op) "
                "VALUES ('2026-03-02', 1, 'Cursus', '09:00', '17:30', 1, 8.0, 'Locatie A', '', 2, "
                "'2026-03-01 10:00:00')")
        db.engine.dispose()
    # De TestConfig van de oude app zette DATA_MAP om; terugzetten naar de back-upmap van deze app
    os.environ["DATA_MAP"] = gemigreerd.config["DATA_MAP"]

    naam = "rooster-20260301-100000-handmatig.db"
    shutil.copy(oud_map / "test.db", os.path.join(backup.backup_map(), naam))
    backup.zet_terug(backup.pad_van(naam))
    db.session.remove()
    dienst = Dienst.query.one()
    assert (dienst.volgnummer, dienst.dienstnaam, dienst.opmerking_tekst, dienst.versie) == \
        (1, "Cursus", "Locatie A", 2)


def test_backup_met_tweede_dienst_terugzetten(gemigreerd):
    db.session.add(Medewerker(id=1, naam="Medewerker A", initialen="TSA"))
    db.session.add_all([Dienst(medewerker_id=1, datum=date(2026, 3, 2), volgnummer=v, versie=1,
                               dienstnaam_override=f"Dienst {v}") for v in (1, 2)])
    db.session.commit()
    pad = backup.maak_backup("handmatig")
    Dienst.query.filter_by(volgnummer=2).delete()
    db.session.commit()
    backup.zet_terug(pad)
    db.session.remove()
    assert sorted((d.volgnummer, d.dienstnaam) for d in Dienst.query.all()) == \
        [(1, "Dienst 1"), (2, "Dienst 2")]
