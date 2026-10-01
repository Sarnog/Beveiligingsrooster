"""Rooster exporteren naar MS Excel (.xlsx) en weer importeren.

1.5.0: export met waarden, voor iedereen onder /export/.
1.6.0 (bewust anders): alleen de beheerder, onder Beheer → Excel import/export, met vier
keuzes (week, periode, jaar, één persoon) en met formules die rekenen zoals de app.
De formules worden uitgerekend met de bibliotheek 'formulas' (zie tests/excel_formules.py).
"""

import io
from datetime import date, time, timedelta

import openpyxl
import pytest

from app.extensions import db
from app.models import Contracturen, Dienst, Dienstcode, Logboek, Medewerker, Vakantie
from app.services import instellingen
from app.services.voorbeeldpakket import laad_voorbeeldpakket
from app.services.weekrooster import Wijziging, wijzig_cellen

from .excel_formules import bereken

MAANDAG = date(2026, 3, 2)  # week 10
URL = "/beheer/exporteren/rooster.xlsx"


@pytest.fixture
def rooster(klaar):
    laad_voorbeeldpakket()
    a = Medewerker(naam="Medewerker A", initialen="MA", volgorde=1)
    b = Medewerker(naam="=HYPERLINK(\"http://x\")", initialen="MB", volgorde=2)  # formule-injectie
    db.session.add_all([a, b])
    db.session.flush()
    db.session.add(Contracturen(medewerker_id=a.id, jaar=2026, uren=1659))
    db.session.add(Vakantie(naam="Voorjaarsvakantie", datum_van=date(2026, 2, 16),
                            datum_tot=date(2026, 2, 22)))
    db.session.commit()
    dinsdag, woensdag, donderdag = (MAANDAG + timedelta(days=i) for i in (1, 2, 3))
    _, fouten = wijzig_cellen([
        Wijziging(a.id, MAANDAG, "code", "4"), Wijziging(a.id, MAANDAG, "opmerking", "Locatie A"),
        Wijziging(a.id, MAANDAG, "opm_begin", "13:30"), Wijziging(a.id, MAANDAG, "opm_eind", "15:45"),
        Wijziging(a.id, dinsdag, "code", "17/3"),  # twee diensten
        Wijziging(a.id, dinsdag, "eind", "22:30", volgnummer=2),  # eigen tijd bij dienst 2
        Wijziging(a.id, woensdag, "dienstnaam", "=SOM(A1:A9)"),  # vrije dienst die op een formule lijkt
        Wijziging(a.id, woensdag, "uren", "6"),
        Wijziging(a.id, donderdag, "code", "4"),
        Wijziging(a.id, donderdag, "dienstnaam", "VW Vroeg tot 12:00"),
        Wijziging(a.id, donderdag, "eind", "12:00"),
        Wijziging(b.id, MAANDAG, "code", "/3"),  # alleen een tweede dienst
        Wijziging(b.id, date(2026, 12, 30), "code", "5"),  # week 53
    ])
    assert fouten == []
    from app.services.weekrooster import pas_dagopmerking_toe

    pas_dagopmerking_toe(MAANDAG, "Eigen dagtekst")
    db.session.commit()
    return {"a": a.id, "b": b.id}


def _download(client, **params):
    antwoord = client.get(URL, query_string=params)
    assert antwoord.status_code == 200, antwoord.data[:300]
    return antwoord, openpyxl.load_workbook(io.BytesIO(antwoord.data))


# ---------------------------------------------------------------------------
# 1a/1b: plek, toegang en keuzes
# ---------------------------------------------------------------------------

def test_download_jaar_met_bladen_en_bestandsnaam(app, als_beheerder, rooster):
    antwoord, boek = _download(als_beheerder, soort="jaar", jaar=2026)
    assert antwoord.headers["Content-Disposition"] == "attachment; filename=rooster-2026.xlsx"
    assert antwoord.mimetype.endswith("spreadsheetml.sheet")
    assert [f"W{w}" for w in range(1, 54)] == [n for n in boek.sheetnames if n.startswith("W")]
    bladen = {"Lijsten", "Feestdagen", "Urenoverzicht", "Vakanties", "Kalender", "Rekenhulp"}
    assert bladen <= set(boek.sheetnames)
    assert boek["Kalender"]["E2"].value == 2026
    assert boek["Vakanties"]["A2"].value == "Voorjaarsvakantie"
    assert boek["Vakanties"]["D2"].value == "=NETWORKDAYS(B2,C2)"
    lijsten = boek["Lijsten"]
    assert (lijsten["B2"].value, lijsten["C2"].value, lijsten["D2"].value) == ("MA", "Medewerker A", 1659)
    regel = Logboek.query.filter_by(actie="Rooster geëxporteerd").one()
    assert "29-12-2025 t/m 03-01-2027" in regel.details and regel.nieuwe_waarde == "rooster-2026.xlsx"


@pytest.mark.parametrize("params, naam, bladen", [
    ({"soort": "week", "jaar": "2026", "week": "10"}, "rooster-2026-W10.xlsx", ["W10"]),
    ({"soort": "week", "jaar": "2026", "week": "53"}, "rooster-2026-W53.xlsx", ["W53"]),
    ({"soort": "periode", "van": "2026-03-02", "tot": "2026-03-15"}, "rooster-20260302-20260315.xlsx",
     ["W10", "W11"]),
])
def test_keuzes_week_en_periode(app, als_beheerder, rooster, params, naam, bladen):
    antwoord, boek = _download(als_beheerder, **params)
    assert antwoord.headers["Content-Disposition"] == f"attachment; filename={naam}"
    assert [n for n in boek.sheetnames if n.startswith("W")] == bladen


def test_week_of_periode_met_een_medewerker(app, als_beheerder, rooster):
    antwoord, boek = _download(als_beheerder, soort="week", jaar=2026, week=10, medewerker=rooster["a"])
    assert antwoord.headers["Content-Disposition"] == "attachment; filename=rooster-2026-W10-ma.xlsx"
    assert boek["W10"]["B4"].value == "Medewerker A" and boek["W10"]["B8"].value is None
    assert [r[0].value for r in boek["Urenoverzicht"].iter_rows(min_row=2)] == ["Medewerker A"]
    antwoord, _ = _download(als_beheerder, soort="periode", van="2026-03-02", tot="2026-03-08",
                            medewerker=rooster["a"])
    assert antwoord.headers["Content-Disposition"].endswith("rooster-20260302-20260308-ma.xlsx")


def test_jaarrooster_van_een_persoon(app, als_beheerder, rooster):
    antwoord, boek = _download(als_beheerder, soort="persoon", jaar=2026, medewerker=rooster["a"])
    assert antwoord.headers["Content-Disposition"] == "attachment; filename=rooster-2026-ma.xlsx"
    assert len([n for n in boek.sheetnames if n.startswith("W")]) == 53
    assert boek["W10"]["B4"].value == "Medewerker A" and boek["W10"]["B8"].value is None


@pytest.mark.parametrize("params, melding", [
    ({}, "Kies wat je wilt exporteren"),
    ({"soort": "onzin"}, "Kies wat je wilt exporteren"),
    ({"soort": "week", "jaar": "2026", "week": "54"}, "weken 1 t/m 53"),
    ({"soort": "week", "jaar": "2025", "week": "53"}, "weken 1 t/m 52"),
    ({"soort": "week", "jaar": "2026", "week": "0"}, "weken 1 t/m 53"),
    ({"soort": "week", "jaar": "2026", "week": "abc"}, "weeknummer"),
    ({"soort": "week", "jaar": "99999999999", "week": "1"}, "jaar"),
    ({"soort": "jaar", "jaar": "1800"}, "Kies een jaar tussen"),
    ({"soort": "jaar", "jaar": ""}, "jaar"),
    ({"soort": "periode", "van": "2026-12-28", "tot": "2027-01-10"}, "binnen één roosterjaar"),
    ({"soort": "periode", "van": "2026-03-08", "tot": "2026-03-02"}, "vóór de begindatum"),
    ({"soort": "periode", "van": "onzin", "tot": "2026-03-02"}, "geldige periode"),
    ({"soort": "periode", "van": "1800-01-01", "tot": "1800-01-02"}, "periode tussen"),
    ({"soort": "persoon", "jaar": "2026"}, "Kies de medewerker"),
    ({"soort": "jaar", "jaar": "2026", "medewerker": "999"}, "Onbekende medewerker"),
    ({"soort": "jaar", "jaar": "2026", "medewerker": "99999999999999"}, "medewerker"),
])
def test_ongeldige_keuzes_geven_een_melding(app, als_beheerder, rooster, params, melding):
    antwoord = als_beheerder.get(URL, query_string=params)
    assert antwoord.status_code == 302 and "/beheer/importeren" in antwoord.headers["Location"], params
    pagina = als_beheerder.get(antwoord.headers["Location"]).data.decode()
    assert melding in pagina, (params, melding)
    assert 'data-export-formulier' in pagina


def test_alleen_de_beheerder(app, client, klaar, rooster):
    from .conftest import login

    assert client.get(URL + "?soort=jaar&jaar=2026").status_code == 302  # niet ingelogd: naar login
    login(client, "collega")
    assert client.get(URL + "?soort=jaar&jaar=2026").status_code == 403
    assert client.get("/export/").status_code == 404  # de oude route bestaat niet meer
    assert client.get("/export/rooster.xlsx?jaar=2026").status_code == 404


def test_geen_excelknoppen_meer_buiten_beheer(app, als_beheerder, rooster):
    for pagina in ("/week/2026/10", "/kalender/?jaar=2026", "/overzicht/uren?jaar=2026",
                   "/zoeken/?van=2026-03-02&tot=2026-03-08&code=4"):
        tekst = als_beheerder.get(pagina).data.decode()
        assert "Exporteren (Excel)" not in tekst and "data-export-excel" not in tekst, pagina
        assert "Rooster exporteren (Excel)" not in tekst
    assert "Exporteren (CSV)" in als_beheerder.get("/overzicht/uren?jaar=2026").data.decode()
    scherm = als_beheerder.get("/beheer/importeren").data.decode()
    assert "Excel import/export" in scherm and 'id="importeren"' in scherm and 'id="exporteren"' in scherm
    assert "Excel import/export" in als_beheerder.get("/beheer/").data.decode()


def test_formule_injectie_wordt_tekst(app, als_beheerder, rooster):
    _, boek = _download(als_beheerder, soort="jaar", jaar=2026)
    blad = boek["W10"]
    for cel in (blad["B8"], blad["J6"], boek["Lijsten"]["C3"]):
        assert cel.data_type == "s" and str(cel.value).startswith("=")
    waarden = bereken(_download(als_beheerder, soort="week", jaar=2026, week=10)[0].data)
    assert waarden["W10!J6"] == "=SOM(A1:A9)"  # de vrije dienstnaam blijft tekst, ook na uitrekenen


# ---------------------------------------------------------------------------
# 1c: formules (uitgerekend zonder Excel)
# ---------------------------------------------------------------------------

def _week10(client, **wijzig):
    antwoord, boek = _download(client, soort="week", jaar=2026, week=10)
    return boek, bereken(antwoord.data, {f"W10!{k}": v for k, v in wijzig.items()})


def test_weekblad_rekent_zoals_de_app(app, als_beheerder, rooster):
    boek, w = _week10(als_beheerder)
    blad = boek["W10"]
    assert blad["B4"].value == "Medewerker A" and blad["D3"].value == "Eigen dagtekst"
    assert blad["D4"].value == "Locatie A" and str(blad["D5"].value) == "13:30:00"
    # Standaarddienst: naam en tijden zoeken de code op in Lijsten
    assert blad["D6"].value.startswith("=IFERROR(VLOOKUP(") and blad["F7"].value.startswith("=")
    assert w["W10!D6"] == "VW Vroeg" and w["W10!F7"] == 8.0
    code4 = Dienstcode.query.filter_by(nummer=4).one()
    assert blad["D6"].fill.fgColor.rgb.endswith(code4.kleur_achtergrond.lstrip("#").upper())
    # Twee diensten: code-raster '17/3', dienst 2 rechts met eigen eindtijd (een waarde)
    assert blad["AD6"].value == "17/3" and blad["AN6"].value == "VW Avond"
    assert str(blad["AO7"].value) == "22:30:00" and w["W10!I7"] == 4.0 and w["W10!AP7"] == 7.5
    # Zelf ingevulde uren: vaste waarde, rood, met een opmerking
    assert blad["L7"].value == 6 and blad["L7"].comment is not None
    assert blad["L7"].font.color.rgb.endswith("C00000")
    # Afwijkende eindtijd en een aanvulling achter de naam: waarden
    assert blad["M6"].value == "VW Vroeg tot 12:00" and str(blad["N7"].value) == "12:00:00"
    assert w["W10!O7"] == 4.75
    # Weektotaal: SUM van dienst 1 én dienst 2; contracturen blijven een waarde
    assert blad["Z6"].value.startswith("=SUM(F7,AM7,I7,AP7")
    assert w["W10!Z6"] == 30.25 and blad["Z4"].value == 1659
    assert blad["AB6"].value == "MA" and blad["AC6"].value == 4
    # Alleen een tweede dienst (code '/3'): dienst 1 leeg, dienst 2 met uren
    assert blad["AC8"].value == "/3" and w["W10!F11"] == "" and w["W10!AM11"] == 8.0 and w["W10!Z10"] == 8.0


def test_tijd_of_code_wijzigen_in_excel_rekent_opnieuw(app, als_beheerder, rooster):
    # Eindtijd maandag 15:45 -> 17:15 (10 uur - 0,5 pauze = 9,5) en de code van vrijdag erbij
    _, w = _week10(als_beheerder, E7=time(17, 15), AG6=5)
    assert w["W10!F7"] == 9.5
    assert (w["W10!P6"], w["W10!P7"], w["W10!R7"]) == ("VW Dag", 0.302083333333333, 9.0)  # 07:15-16:45
    assert w["W10!Z6"] == 30.25 + 1.5 + 9
    # Een code zonder standaardtijden geeft geen 00:00 maar een lege tijd en geen uren
    _, w = _week10(als_beheerder, AG6=10)
    assert (w["W10!P6"], w["W10!P7"], w["W10!R7"]) == ("Bapo", "", "")
    # Twee codes in één cel
    _, w = _week10(als_beheerder, AG6="4/17")
    assert (w["W10!P6"], w["W10!R7"], w["W10!AW6"], w["W10!AY7"]) == ("VW Vroeg", 8.0, "BHV", 4.0)


def test_urenoverzicht_en_vakanties(app, als_beheerder, rooster):
    antwoord, boek = _download(als_beheerder, soort="jaar", jaar=2026)
    overzicht = boek["Urenoverzicht"]
    assert overzicht["L2"].value == "='W10'!$Z$6"  # kolom L = W10
    assert overzicht["BD2"].value == "=SUM(C2:BC2)" and overzicht["BF2"].value == '=IF(BE2="","",BD2-BE2)'
    # Een jaar uitrekenen is te zwaar voor de test: één persoon in één week is genoeg
    antwoord, _ = _download(als_beheerder, soort="week", jaar=2026, week=10, medewerker=rooster["a"])
    w = bereken(antwoord.data, {"W10!E7": time(17, 15)})
    assert w["Urenoverzicht!L2"] == 31.75
    assert w["Urenoverzicht!BD2"] == 31.75  # totaal (andere weken: waarden uit de app, hier leeg)
    assert w["Urenoverzicht!BF2"] == 31.75 - 1659  # verschil met de contracturen
    antwoord, _ = _download(als_beheerder, soort="periode", van="2026-02-16", tot="2026-02-22")
    assert bereken(antwoord.data)["Vakanties!D2"] == 5.0  # NETWORKDAYS, zoals in het oude bestand


# ---------------------------------------------------------------------------
# Herimport
# ---------------------------------------------------------------------------

def _rooster_als_tuples():
    """Alles wat bij een herimport terug moet komen (op naam, zonder database-ID's)."""
    db.session.expire_all()
    diensten = sorted(
        (d.medewerker.naam, d.datum, d.volgnummer, d.dienstcode.nummer if d.dienstcode else None,
         d.dienstnaam, d.begin, d.eind, d.opmerking_tekst, d.opmerking_begin, d.opmerking_eind,
         d.uren_handmatig, d.uren_berekend)
        for d in Dienst.query.all() if not d.is_leeg or d.volgnummer == 1)
    codes = sorted((c.nummer, c.omschrijving, c.std_begin, c.std_eind) for c in Dienstcode.query.all())
    contract = sorted((c.medewerker.naam, c.jaar, c.uren) for c in Contracturen.query.all())
    return diensten, codes, contract


def _herimport(inhoud: bytes, tmp_path, jaar=2026, voorbereiden=None):
    from app import create_app
    from app.services.excel_import import importeer, lees_bestand

    from .conftest import TestConfig

    pad = tmp_path / "rooster.xlsx"
    pad.write_bytes(inhoud)
    leeg = create_app(TestConfig(str(tmp_path / "leeg")))
    with leeg.app_context():
        db.create_all()
        if voorbereiden:
            voorbereiden()
        plan = lees_bestand(str(pad), jaar)
        assert plan.waarschuwingen == []
        assert plan.weektotaal_verschillen() == []
        assert plan.handmatige_uren() == []  # uren uit een formule zijn niet 'zelf ingevuld'
        importeer(plan)
        na = _rooster_als_tuples()
        db.session.remove()
    return na


def test_export_en_weer_importeren_in_een_lege_database(app, als_beheerder, rooster, tmp_path):
    voor = _rooster_als_tuples()
    assert len(voor[0]) == 8 and any(d[2] == 2 and d[6] == "22:30" for d in voor[0])  # o.a. dienst 2
    assert ("Medewerker A", date(2026, 3, 5), 1, 4, "VW Vroeg tot 12:00") in [d[:5] for d in voor[0]]
    na = _herimport(als_beheerder.get(URL + "?soort=jaar&jaar=2026").data, tmp_path)
    assert na[0] == voor[0]  # diensten, tijden, opmerkingen, uren, twee diensten
    assert na[1] == voor[1]  # dienstcodes
    assert na[2] == voor[2]  # contracturen


def test_herimport_van_een_in_excel_opgeslagen_bestand(app, als_beheerder, rooster, tmp_path):
    """Formules mét opgeslagen waarden (zoals Excel ze bewaart): die waarden gelden."""
    inhoud = als_beheerder.get(URL + "?soort=week&jaar=2026&week=10").data
    voor = [d for d in _rooster_als_tuples()[0] if d[1].isocalendar()[1] == 10]
    na = _herimport(_als_excel_opgeslagen(inhoud), tmp_path)
    assert [d for d in na[0] if d[1].isocalendar()[1] == 10] == voor


def _als_excel_opgeslagen(inhoud: bytes) -> bytes:
    """Bootst 'opslaan in Excel' na: elke formulecel krijgt zijn uitgerekende waarde erbij.

    openpyxl kan geen formule én waarde schrijven; we zetten de waarde in de XML (<v>).
    """
    import re
    import zipfile

    waarden = bereken(inhoud)
    bron = zipfile.ZipFile(io.BytesIO(inhoud))
    uit = io.BytesIO()
    namen = {}
    werkboek = bron.read("xl/workbook.xml").decode()
    relaties = bron.read("xl/_rels/workbook.xml.rels").decode()
    bladpatroon = r'<sheet name="([^"]+)" sheetId="\d+"(?: state="\w+")? r:id="(rId\d+)"'
    for naam, rid in re.findall(bladpatroon, werkboek):
        doel = re.search(rf'Id="{rid}"[^>]*Target="([^"]+)"|Target="([^"]+)"[^>]*Id="{rid}"', relaties)
        namen["xl/" + (doel.group(1) or doel.group(2)).removeprefix("/xl/")] = naam
    with zipfile.ZipFile(uit, "w", zipfile.ZIP_DEFLATED) as doel_zip:
        for onderdeel in bron.infolist():
            data = bron.read(onderdeel.filename)
            if onderdeel.filename in namen:
                blad = namen[onderdeel.filename]

                def met_waarde(m, blad=blad):
                    adres, attrs, formule = m.group(1), m.group(2), m.group(3)
                    waarde = waarden.get(f"{blad}!{adres}", "")
                    if isinstance(waarde, bool) or isinstance(waarde, float):
                        return f'<c r="{adres}"{attrs}><f>{formule}</f><v>{float(waarde)!r}</v></c>'
                    tekst = str(waarde).replace("&", "&amp;").replace("<", "&lt;")
                    attrs = re.sub(r' t="\w+"', "", attrs)
                    return f'<c r="{adres}"{attrs} t="str"><f>{formule}</f><v>{tekst}</v></c>'

                data = re.sub(r'<c r="([A-Z]+\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>',
                              met_waarde, data.decode()).encode()
            doel_zip.writestr(onderdeel, data)
    return uit.getvalue()


def test_herimport_met_andere_instellingen_in_het_bestand(app, als_beheerder, rooster, tmp_path):
    """Pauze, feestdagtoeslag en opmerkingtijden uit Lijsten: de controle rekent ermee."""
    instellingen.schrijf("pauze", instellingen.pauze_json(True, [(5.5, 0.5), (9, 0.75)]))
    instellingen.schrijf("opmerkingtijden_meetellen", "1")
    db.session.commit()
    from app.services.rooster import herbereken_alle

    herbereken_alle()
    voor = _rooster_als_tuples()

    def zelfde_instellingen():
        instellingen.schrijf("pauze", instellingen.pauze_json(True, [(5.5, 0.5), (9, 0.75)]))
        instellingen.schrijf("opmerkingtijden_meetellen", "1")
        db.session.commit()

    na = _herimport(als_beheerder.get(URL + "?soort=jaar&jaar=2026").data, tmp_path,
                    voorbereiden=zelfde_instellingen)
    assert na[0] == voor[0]
