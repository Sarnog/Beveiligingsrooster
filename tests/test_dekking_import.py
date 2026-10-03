"""Excel-import: hulpfuncties, onleesbare bestanden en zeldzame situaties. Alle data is fictief."""

import zipfile
from datetime import date, datetime, time

import openpyxl
import pytest
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models import Contracturen, Dagopmerking, Dienst, Medewerker
from app.services import excel_import
from app.services.excel_import import (
    MODUS_AANVULLEN,
    ImportFout,
    ImportKeuzes,
    controleer_bestand,
    importeer,
    jaar_uit_bestand,
    jaar_uit_naam,
    lees_bestand,
)
from app.services.feestdagen import feestdagen_in_periode

from .test_import_keuzes import basis, maak_rooster  # noqa: F401  (basis is een fixture)

KONINGSDAG = date(2026, 4, 27)  # maandag, week 18


def test_hulpfuncties_voor_celwaarden():
    assert excel_import._tijd(datetime(2026, 1, 1, 7, 15)) == "07:15"
    assert excel_import._tijd(0.5) == "12:00"  # Excel: tijd als deel van een dag
    assert excel_import._tijd("geen tijd") is None and excel_import._tijd(1.5) is None
    # Een tijd die als tekst in de cel staat (zo stond 15:45 in W34 van het rooster van 2026)
    assert excel_import._tijd("15:45") == "15:45" and excel_import._tijd(" 7.15 ") == "07:15"
    assert excel_import._tijd("07:15:00") == "07:15" and excel_import._tijd("24:00") == "00:00"
    assert excel_import._tijd("25:00") is None and excel_import._tijd("7:75") is None
    assert excel_import._tekst(4.0) == "4" and excel_import._tekst(4.5) == "4.5"
    assert excel_import._getal("abc") is None and excel_import._getal(10 ** 400) is None
    assert excel_import._getal("1,5") == 1.5
    assert excel_import._datum(date(2026, 1, 2)) == date(2026, 1, 2) and excel_import._datum("x") is None


def test_eindtijd_als_tekst_telt_mee(app, basis, tmp_path):  # noqa: F811
    """Stond de eindtijd als tekst in Excel ('15:45'), dan had de dienst geen uren (8 uur te weinig)."""
    pad = str(tmp_path / "tekst.xlsx")
    maak_rooster(pad, 2026, {10: {"Medewerker X": {0: (4, "07:15", "15:45")}}})
    boek = openpyxl.load_workbook(pad)
    boek["W10"]["E7"] = "15:45"  # eindtijd van Medewerker X op maandag, als tekst
    boek["W10"]["F7"] = 8  # de uren zoals de oude macro ze zette
    boek.save(pad)
    importeer(lees_bestand(pad, 2026))
    dienst = Dienst.query.filter_by(medewerker_id=basis["x"], datum=date(2026, 3, 2)).one()
    assert (dienst.eind, dienst.uren_handmatig, dienst.uren_berekend) == ("15:45", None, 8)


def test_jaar_uit_naam_en_onleesbaar_bestand(tmp_path):
    assert jaar_uit_naam("Rooster 1800 en 2027.xlsm") == 2027  # 1800 valt buiten de grenzen
    assert jaar_uit_naam("") is None
    rommel = tmp_path / "rommel.xlsx"
    rommel.write_bytes(b"geen excel")
    assert jaar_uit_bestand(str(rommel)) is None


def test_zip_zonder_excel_wordt_geweigerd(tmp_path):
    pad = tmp_path / "zip.xlsx"
    with zipfile.ZipFile(pad, "w") as archief:
        archief.writestr("leesmij.txt", "geen werkboek")
    with pytest.raises(ImportFout, match="kan niet gelezen worden"):
        controleer_bestand(str(pad))
    with pytest.raises(ImportFout, match="kan niet gelezen worden"):
        lees_bestand(str(pad), 2026)


def test_bestand_zonder_lijsten_en_ongeldig_jaar(tmp_path):
    pad = str(tmp_path / "leeg.xlsx")
    openpyxl.Workbook().save(pad)
    with pytest.raises(ImportFout, match="Lijsten"):
        lees_bestand(pad, 2026)
    with pytest.raises(ImportFout, match="Lijsten"):
        controleer_bestand(pad)
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {})
    with pytest.raises(ImportFout, match="Kies een jaar tussen"):
        lees_bestand(pad, 1800)


def test_weekblad_dat_niet_in_het_jaar_bestaat(app, klaar, tmp_path):
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {53: {"Medewerker X": {0: (4, "07:15", "15:45")}}})
    boek = openpyxl.load_workbook(pad)
    for kolom in (4, 7, 10, 13, 16, 19, 22):
        boek["W53"].cell(2, kolom).value = None  # geen datums: het weeknummer telt
    boek.create_sheet("W60")["B4"] = "Medewerker X"  # geen diensten: zonder melding overgeslagen
    boek.save(pad)
    plan = lees_bestand(pad, 2027)  # 2027 heeft 52 weken
    assert any("W53 overgeslagen: 2027 heeft geen week 53" in w for w in plan.waarschuwingen)
    assert not any("W60" in w for w in plan.waarschuwingen)


def test_ongeldige_pauzeregels_en_rare_bladinhoud(app, klaar, tmp_path):
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {10: {"Medewerker X": {0: (4, "07:15", "15:45")}}})
    boek = openpyxl.load_workbook(pad)
    lijsten = boek["Lijsten"]
    lijsten["M6"], lijsten["N6"] = "Pauzeaftrek aan", 1
    lijsten.cell(9, 14, 5.5)
    lijsten.cell(9, 15, 9)  # meer pauze dan de dienst lang is: ongeldig
    week = boek["W10"]
    week.cell(2, 4, datetime(2026, 5, 4))  # datum van een andere week
    week.cell(4 + 4 * 2, 2, 0)  # '0' uit een formule: geen naam
    week.cell(4 + 4 * 3, 2, "Iemand Anders")  # staat niet in Lijsten
    for n in range(4, 30):  # alle 30 blokken met iets erin: de lus loopt helemaal door
        week.cell(4 + 4 * n + 1, 2, "x")
    boek.save(pad)
    plan = lees_bestand(pad, 2026)
    tekst = " ".join(plan.waarschuwingen)
    assert "pauzeregels in Lijsten zijn ongeldig" in tekst
    assert "past niet bij week 10" in tekst
    assert "'Iemand Anders' staat niet in Lijsten" in tekst
    assert plan.pauze == ((5.5, 0.5),) or plan.pauze is not None


def test_import_botsing_wordt_importfout(app, basis, tmp_path, monkeypatch):  # noqa: F811
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {10: {"Medewerker X": {0: (4, "07:15", "15:45")}}})
    plan = lees_bestand(pad, 2026)

    def botsing(*_args):
        raise IntegrityError("INSERT", {}, Exception("UNIQUE constraint failed"))

    monkeypatch.setattr(excel_import, "_importeer", botsing)
    with pytest.raises(ImportFout, match="gegevens die botsen"):
        importeer(plan)


def test_import_bestaande_contracturen_en_dagopmerkingen(app, basis, tmp_path):  # noqa: F811
    x = db.session.get(Medewerker, basis["x"])
    x.contracturen.append(Contracturen(jaar=2026, uren=1000))
    feestdag = feestdagen_in_periode(KONINGSDAG, KONINGSDAG)[KONINGSDAG]
    tweede_paasdag = date(2026, 4, 6)
    db.session.add_all([Dagopmerking(datum=KONINGSDAG, tekst="Eigen tekst", handmatig=True),
                        Dagopmerking(datum=tweede_paasdag, tekst="Oud", handmatig=True)])
    dienst = Dienst(medewerker_id=x.id, datum=KONINGSDAG, dienstnaam_override="Oud", begin="08:00",
                    eind="16:00", google_event_id="afspraak-1")
    db.session.add(dienst)
    x.agenda_modus, x.agenda_id = "B", "agenda-x"
    db.session.commit()
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {15: {}, 18: {}},
                 dagopmerkingen={KONINGSDAG: feestdag, tweede_paasdag: "Nieuwe tekst"})
    plan = lees_bestand(pad, 2026)
    resultaat = importeer(plan, ImportKeuzes(contracturen=True))
    db.session.expire_all()
    assert db.session.get(Medewerker, basis["x"]).contracturen_voor(2026) == 1500
    assert Dagopmerking.query.filter_by(datum=KONINGSDAG).first() is None  # weer automatisch
    assert Dagopmerking.query.filter_by(datum=tweede_paasdag).one().tekst == "Nieuwe tekst"
    assert resultaat["dagopmerkingen"] == 2


def test_aanvullen_laat_handmatige_dagopmerking_staan(app, basis, tmp_path):  # noqa: F811
    db.session.add(Dagopmerking(datum=date(2026, 3, 4), tekst="Eigen tekst", handmatig=True))
    db.session.commit()
    pad = str(tmp_path / "rooster.xlsx")
    maak_rooster(pad, 2026, {10: {}}, dagopmerkingen={date(2026, 3, 4): "Uit Excel"})
    plan = lees_bestand(pad, 2026)
    uitkomst = excel_import.effect(plan, ImportKeuzes(modus=MODUS_AANVULLEN))
    assert uitkomst.dagopmerkingen == []
    assert time  # noqa: B018
