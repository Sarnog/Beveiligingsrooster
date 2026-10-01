"""Regressietests voor de bevindingen uit de audit van versie 1.1.2.

Elke test hoort bij één bevinding (C = critical, H = high, M = medium, L = low,
S = security). De test faalde op de oude code en slaagt na de reparatie.
"""

import os
import sqlite3
from datetime import datetime, timedelta

import pytest

from app import worker
from app.extensions import db
from app.models import Logboek, LoginPoging
from app.services import backup, instellingen, klok

# ---------------------------------------------------------------------------
# Hulpjes
# ---------------------------------------------------------------------------


class Klok:
    """Een klok die bij elke aanroep één seconde verder loopt (unieke back-upnamen)."""

    def __init__(self, start: datetime) -> None:
        self.tijd = start

    def __call__(self) -> datetime:
        self.tijd += timedelta(seconds=1)
        return self.tijd


class KapotteVerbinding:
    """Een echte SQLite-verbinding, behalve dat backup() faalt (schijf vol)."""

    def __init__(self, echt):
        self._echt = echt

    def backup(self, *args, **kwargs):
        raise sqlite3.OperationalError("database or disk is full")

    def __getattr__(self, naam):
        return getattr(self._echt, naam)

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


class KapotSqlite:
    """Nep-module voor backup.sqlite3: connect() werkt, backup() mislukt."""

    OperationalError = sqlite3.OperationalError
    DatabaseError = sqlite3.DatabaseError

    def connect(self, *args, **kwargs):
        return KapotteVerbinding(sqlite3.connect(*args, **kwargs))


def _is_geldige_db(pad: str) -> bool:
    if os.path.getsize(pad) == 0:
        return False
    verbinding = sqlite3.connect(pad)
    try:
        return verbinding.execute("PRAGMA integrity_check").fetchone()[0] == "ok" and bool(
            verbinding.execute("SELECT count(*) FROM sqlite_master").fetchone()[0])
    finally:
        verbinding.close()


# ---------------------------------------------------------------------------
# C1 · Back-uprotatie gooit alle goede back-ups weg na een tijdelijke fout
# ---------------------------------------------------------------------------

def test_c1_mislukte_backups_verdringen_goede_niet(app, klaar, monkeypatch):
    instellingen.schrijf("backup_bewaren", "6")
    db.session.commit()
    monkeypatch.setattr(klok, "nu", Klok(datetime(2026, 3, 1, 1, 0)))
    goed = [backup.maak_backup() for _ in range(5)]

    planning = worker.Planning()
    moment = datetime(2026, 3, 2, 2, 0)
    monkeypatch.setattr(backup, "sqlite3", KapotSqlite())
    for _ in range(10):
        try:
            worker.een_ronde(planning, moment)
        except Exception:  # oude code liet de fout door
            pass
        moment += timedelta(minutes=31)  # na de backoff mag het opnieuw
    monkeypatch.setattr(backup, "sqlite3", sqlite3)
    worker.een_ronde(planning, moment)

    bestanden = os.listdir(backup.backup_map())
    assert not [n for n in bestanden if not n.endswith(".db")]  # geen .tmp-resten
    paden = [os.path.join(backup.backup_map(), n) for n in bestanden]
    assert all(_is_geldige_db(p) for p in paden), "lege of halve back-up gevonden"
    assert all(os.path.exists(p) for p in goed)  # de oude goede back-ups staan er nog
    assert len(paden) == 6
    assert Logboek.query.filter_by(actie="Back-up mislukt").count() >= 1


def test_c1_backoff_na_mislukte_backup(app, klaar, monkeypatch):
    """Na een mislukte back-up niet elke ronde (5 s) opnieuw proberen."""
    pogingen = []

    def mislukt(label=""):
        pogingen.append(label)
        raise OSError("schijf vol")

    monkeypatch.setattr(backup, "maak_backup", mislukt)
    planning = worker.Planning()
    moment = datetime(2026, 3, 2, 3, 0)
    for _ in range(20):
        worker.een_ronde(planning, moment)
        moment += timedelta(seconds=5)
    assert len(pogingen) == 1
    worker.een_ronde(planning, moment + timedelta(minutes=31))
    assert len(pogingen) == 2


def test_c1_ruim_oude_op_telt_alleen_geldige_automatische(app, klaar):
    instellingen.schrijf("backup_bewaren", "2")
    db.session.commit()
    map_ = backup.backup_map()
    for naam in ("rooster-20260101-010000.db", "rooster-20260102-010000.db",
                 "rooster-20260103-010000.db"):
        with open(os.path.join(map_, naam), "wb") as f:
            f.write(b"x")
    open(os.path.join(map_, "rooster-20260104-010000.db"), "wb").close()  # leeg
    with open(os.path.join(map_, "rooster-20260105-010000-handmatig.db"), "wb") as f:
        f.write(b"x")
    backup.ruim_oude_op()
    over = sorted(os.listdir(map_))
    assert "rooster-20260102-010000.db" in over and "rooster-20260103-010000.db" in over
    assert "rooster-20260101-010000.db" not in over
    assert "rooster-20260105-010000-handmatig.db" in over


# ---------------------------------------------------------------------------
# C2 · Anonieme bezoeker kan de database onbeperkt vullen via /login
# ---------------------------------------------------------------------------

def _db_grootte(app) -> int:
    with db.engine.connect() as verbinding:
        verbinding.exec_driver_sql("PRAGMA wal_checkpoint(TRUNCATE)")
    pad = app.config["SQLALCHEMY_DATABASE_URI"].removeprefix("sqlite:///")
    return os.path.getsize(pad)


def test_c2_lange_gebruikersnaam_vult_database_niet(app, client, klaar):
    voor = _db_grootte(app)
    naam = "a" * 400_000
    for _ in range(40):
        antwoord = client.post("/login", data={"gebruikersnaam": naam, "wachtwoord": "x"})
        assert antwoord.status_code in (401, 429)
    db.session.remove()
    groei = _db_grootte(app) - voor
    assert groei < 32 * 1024, f"database groeide {groei} bytes"
    assert max(len(p.gebruikersnaam) for p in LoginPoging.query.all()) <= 64
    assert Logboek.query.filter_by(actie="Login geblokkeerd").count() <= 1


def test_c2_logboekvelden_centraal_begrensd(app, klaar):
    from app.services import logboek

    logboek.log("Test", "d" * 5000, oud="o" * 5000, nieuw="n" * 5000, gebruiker="g" * 500,
                medewerker="m" * 500, veld="v" * 500)
    db.session.commit()
    regel = Logboek.query.filter_by(actie="Test").one()
    assert len(regel.details) <= 1000 and len(regel.oude_waarde) <= 1000
    assert len(regel.nieuwe_waarde) <= 1000 and len(regel.gebruiker) <= 64
    assert len(regel.medewerker) <= 120 and len(regel.veld) <= 40


def test_c2_oude_loginpogingen_worden_opgeruimd(app, klaar):
    nu = klok.utc_nu()
    db.session.add(LoginPoging(gebruikersnaam="x", ip="1.2.3.4", tijdstip=nu - timedelta(days=2)))
    db.session.add(LoginPoging(gebruikersnaam="y", ip="1.2.3.4", tijdstip=nu))
    db.session.commit()
    worker.een_ronde(worker.Planning(), datetime(2026, 3, 2, 1, 0))
    assert [p.gebruikersnaam for p in LoginPoging.query.all()] == ["y"]


def test_c2_geblokkeerd_maar_een_keer_gelogd(app, client, klaar):
    for _ in range(12):
        client.post("/login", data={"gebruikersnaam": "beheerder", "wachtwoord": "fout"})
    assert Logboek.query.filter_by(actie="Login geblokkeerd").count() == 1
    assert LoginPoging.query.count() == 5  # geblokkeerde pogingen schrijven niets meer


@pytest.fixture(autouse=True)
def _geen_echte_google(monkeypatch):
    """De worker mag in deze tests nooit echt met Google praten."""
    from app.services import google_agenda
    from app.services.google_agenda import AgendaFout

    def geen():
        raise AgendaFout("geen Google in tests")

    monkeypatch.setattr(google_agenda, "klant", geen)


# ---------------------------------------------------------------------------
# H1 · Zelfde dienstcode opnieuw invoeren reset handmatige tijden stil
# ---------------------------------------------------------------------------

MAANDAG = datetime(2026, 3, 2).date()  # week 10 van 2026


@pytest.fixture
def mw(klaar):
    from app.models import Medewerker
    from app.services.voorbeeldpakket import laad_voorbeeldpakket

    laad_voorbeeldpakket()
    medewerker = Medewerker(naam="Medewerker A", initialen="MA", volgorde=1)
    db.session.add(medewerker)
    db.session.commit()
    return medewerker


def test_h1_zelfde_code_opnieuw_laat_handmatige_tijden_staan(app, mw):
    from app.models import Dienst
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4"), Wijziging(mw.id, MAANDAG, "eind", "18:00")])
    dienst = Dienst.query.one()
    voor = (dienst.eind, dienst.tijden_handmatig, dienst.uren_berekend, dienst.versie)
    regels = Logboek.query.count()
    assert voor[0] == "18:00" and voor[1] is True

    _, fouten = wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4", versie=dienst.versie)])
    assert fouten == []
    db.session.expire_all()
    dienst = Dienst.query.one()
    assert (dienst.eind, dienst.tijden_handmatig, dienst.uren_berekend, dienst.versie) == voor
    assert Logboek.query.count() == regels


# ---------------------------------------------------------------------------
# H2 · Excel-import is niet atomair
# ---------------------------------------------------------------------------

def _import_dienst(naam, dag, **extra):
    from app.services.excel_import import ImportDienst

    waarden = dict(naam=naam, datum=dag, code=None, dienstnaam="Cursus", begin="08:00",
                   eind="16:00", opmerking="", opm_begin=None, opm_eind=None, excel_uren=None)
    waarden.update(extra)
    return ImportDienst(**waarden)


def _bestaande_dienst(mw):
    from app.models import Dienst

    db.session.add(Dienst(medewerker_id=mw.id, datum=MAANDAG, dienstnaam_override="Oud",
                          begin="07:00", eind="15:00", opmerking_tekst=""))
    db.session.commit()


def test_h2_dubbele_dienst_laat_oud_rooster_heel(app, mw):
    from app.models import Dienst, Feestdag
    from app.services.excel_import import ImportMedewerker, ImportPlan, importeer

    _bestaande_dienst(mw)
    Feestdag.query.delete()  # verse installatie: feestdagen bestaan nog niet
    db.session.commit()
    dinsdag = MAANDAG + timedelta(days=1)
    plan = ImportPlan(jaar=2026, weken=[10],
                      medewerkers=[ImportMedewerker("Medewerker A", "MA", None)],
                      diensten=[_import_dienst("Medewerker A", dinsdag),
                                _import_dienst("Medewerker A", dinsdag)])
    with pytest.raises(Exception):  # noqa: B017 - oude code: IntegrityError, nieuw: ImportFout
        importeer(plan)
    db.session.rollback()
    assert Dienst.query.filter_by(datum=MAANDAG).one().dienstnaam_override == "Oud"


def test_h2_fout_halverwege_draait_alles_terug(app, mw, monkeypatch):
    from app.models import Dienst, Feestdag
    from app.services import excel_import
    from app.services.excel_import import ImportMedewerker, ImportPlan

    _bestaande_dienst(mw)
    Feestdag.query.delete()
    db.session.commit()

    def kapot():
        raise RuntimeError("onverwacht")

    monkeypatch.setattr(excel_import, "markeer_bijgewerkt", kapot)
    plan = ImportPlan(jaar=2026, weken=[10],
                      medewerkers=[ImportMedewerker("Medewerker A", "MA", None)],
                      diensten=[_import_dienst("Medewerker A", MAANDAG + timedelta(days=2))])
    with pytest.raises(Exception):  # noqa: B017
        excel_import.importeer(plan)
    db.session.rollback()
    assert Dienst.query.filter_by(datum=MAANDAG).one().dienstnaam_override == "Oud"
    assert Dienst.query.count() == 1


def test_h2_droogloop_meldt_dubbele_diensten_en_scherm_geeft_nette_fout(app, als_beheerder, tmp_path):
    import io

    import openpyxl

    boek = openpyxl.Workbook()
    lijsten = boek.active
    lijsten.title = "Lijsten"
    lijsten.cell(2, 2, "MD")
    lijsten.cell(2, 3, "Medewerker Dubbel")
    boek.create_sheet("Kalender")["E2"] = 2026
    week = boek.create_sheet("W10")
    for blok in (0, 1):  # dezelfde naam twee keer in één week
        week.cell(4 + 4 * blok, 2, "Medewerker Dubbel")
        week.cell(6 + 4 * blok, 4, "Cursus")
        week.cell(7 + 4 * blok, 4, datetime(1900, 1, 1, 8, 0).time())
        week.cell(7 + 4 * blok, 5, datetime(1900, 1, 1, 16, 0).time())
    pad = str(tmp_path / "dubbel.xlsx")
    boek.save(pad)

    from app.services.excel_import import lees_bestand

    plan = lees_bestand(pad)
    assert any("dubbel" in w.lower() for w in plan.waarschuwingen)

    with open(pad, "rb") as f:
        als_beheerder.post("/beheer/importeren", data={"bestand": (io.BytesIO(f.read()), "x.xlsx")},
                           content_type="multipart/form-data")
    antwoord = als_beheerder.post("/beheer/importeren/voorbeeld", data={"bevestig": "1"},
                                  follow_redirects=True)
    assert antwoord.status_code == 200
    assert "niet geïmporteerd" in antwoord.data.decode()


# ---------------------------------------------------------------------------
# H3 · Wijziging tijdens een lopende agenda-sync gaat verloren
# ---------------------------------------------------------------------------

def _plan_dag_in_ander_proces(medewerker_id, datum, wijzig_dienst):
    """Doet wat een webverzoek doet (dienst wijzigen + plan_dag), in een eigen sessie."""
    from sqlalchemy.orm import Session

    from app.models import Dienst, SyncTaak

    with Session(db.engine) as sessie:
        dienst = sessie.query(Dienst).filter_by(medewerker_id=medewerker_id, datum=datum).one()
        wijzig_dienst(dienst)
        dienst.versie += 1
        straks = klok.utc_nu() + timedelta(seconds=10)
        bestaand = sessie.query(SyncTaak).filter_by(
            medewerker_id=medewerker_id, datum=datum, soort="dag", status="wacht").first()
        if bestaand:
            bestaand.niet_voor = straks
        else:
            sessie.add(SyncTaak(medewerker_id=medewerker_id, datum=datum, soort="dag",
                                niet_voor=straks))
        sessie.commit()


@pytest.fixture
def gekoppeld(mw, monkeypatch):
    mw.agenda_modus, mw.agenda_id = "B", "agenda-a"
    db.session.commit()
    return mw


def _wachtrij_nu():
    from app.models import SyncTaak
    from app.services import sync

    SyncTaak.query.update({"niet_voor": datetime(2000, 1, 1)})
    db.session.commit()
    return sync.verwerk_wachtrij()


def test_h3_wijziging_tijdens_sync_blijft_in_wachtrij(app, gekoppeld, monkeypatch):
    from app.models import SyncTaak
    from app.services import google_agenda
    from app.services.weekrooster import Wijziging, wijzig_cellen

    from .test_agenda import NepKlant

    class GelijktijdigeKlant(NepKlant):
        def maak_afspraak(self, agenda_id, body):
            def later(dienst):
                dienst.eind = "18:00"

            _plan_dag_in_ander_proces(gekoppeld.id, MAANDAG, later)
            return super().maak_afspraak(agenda_id, body)

    klant = GelijktijdigeKlant()
    monkeypatch.setattr(google_agenda, "klant", lambda: klant)
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    _wachtrij_nu()
    db.session.expire_all()
    wachtend = SyncTaak.query.filter_by(status="wacht").count()
    afspraak = next(iter(klant.agendas["agenda-a"].values()))
    assert wachtend == 1 or afspraak["end"]["dateTime"].endswith("18:00:00")

    # En de volgende ronde zet de nieuwe tijd in Google
    monkeypatch.setattr(google_agenda, "klant", lambda: NepKlant.__new__(NepKlant))
    klant2 = NepKlant()
    klant2.agendas = klant.agendas
    monkeypatch.setattr(google_agenda, "klant", lambda: klant2)
    _wachtrij_nu()
    afspraak = next(iter(klant.agendas["agenda-a"].values()))
    assert afspraak["end"]["dateTime"].endswith("18:00:00")
    assert len(klant.agendas["agenda-a"]) == 1


def test_h3_blijven_hangen_op_bezig_wordt_hersteld(app, gekoppeld, monkeypatch):
    from app.models import SyncTaak
    from app.services import google_agenda, sync

    from .test_agenda import NepKlant

    monkeypatch.setattr(google_agenda, "klant", NepKlant)
    db.session.add(SyncTaak(medewerker_id=gekoppeld.id, datum=MAANDAG, soort="dag",
                            status="bezig", niet_voor=klok.utc_nu() - timedelta(minutes=11)))
    db.session.add(SyncTaak(medewerker_id=gekoppeld.id, datum=MAANDAG + timedelta(days=1),
                            soort="dag", status="bezig", niet_voor=klok.utc_nu()))
    db.session.commit()
    sync.verwerk_wachtrij()
    db.session.expire_all()
    # De oude 'bezig'-taak is opnieuw opgepakt en klaar; de verse loopt nog (ander proces)
    assert [t.status for t in SyncTaak.query.all()] == ["bezig"]


# ---------------------------------------------------------------------------
# H4 · Gedeelde agenda (modus B): sync verwijdert afspraken van collega's
# ---------------------------------------------------------------------------

def _team(monkeypatch):
    from app.models import Medewerker
    from app.services import google_agenda, sync
    from app.services.weekrooster import Wijziging, wijzig_cellen

    from .test_agenda import NepKlant

    klant = NepKlant()
    monkeypatch.setattr(google_agenda, "klant", lambda: klant)
    a = Medewerker(naam="Collega A", initialen="CA", agenda_modus="B", agenda_id="team")
    b = Medewerker(naam="Collega B", initialen="CB", agenda_modus="B", agenda_id="team")
    db.session.add_all([a, b])
    db.session.commit()
    dag = sync.sync_periode()[0] + timedelta(days=14)  # binnen de sync-periode
    wijzig_cellen([Wijziging(a.id, dag, "code", "4"), Wijziging(b.id, dag, "code", "5")])
    _wachtrij_nu()
    assert len(klant.agendas["team"]) == 2
    return klant, a, b


def _van(klant, medewerker):
    return [e for e in klant.agendas["team"].values()
            if e["extendedProperties"]["private"]["medewerker_id"] == str(medewerker.id)]


def test_h4_volledige_sync_laat_collega_met_rust(app, mw, monkeypatch):
    from app.services import sync_planning

    klant, a, b = _team(monkeypatch)
    sync_planning.plan_volledig(a)
    _wachtrij_nu()
    assert len(_van(klant, b)) == 1 and len(_van(klant, a)) == 1


def test_h4_ontkoppelen_modus_b_wist_alleen_eigen_afspraken(app, mw, monkeypatch):
    from app.services import sync_planning

    klant, a, b = _team(monkeypatch)
    sync_planning.plan_ontkoppel(a, verwijder=True)
    db.session.commit()
    _wachtrij_nu()
    assert _van(klant, a) == [] and len(_van(klant, b)) == 1


def test_h4_oude_afspraak_zonder_medewerker_id(app, mw, monkeypatch):
    """Afspraken van vóór de medewerker-markering: alleen weg als de dienst van A is."""
    from app.models import Dienst
    from app.services import sync

    klant, a, b = _team(monkeypatch)
    dienst_b = Dienst.query.filter_by(medewerker_id=b.id).one()
    klant.agendas["team"]["oud-b"] = {"id": "oud-b", "extendedProperties": {"private": {
        "bron": "beveiligingsrooster", "dienst_id": str(dienst_b.id)}}}
    klant.agendas["team"]["oud-a"] = {"id": "oud-a", "extendedProperties": {"private": {
        "bron": "beveiligingsrooster", "dienst_id": "999999"}}}
    dienst_a = Dienst.query.filter_by(medewerker_id=a.id).one()
    klant.agendas["team"]["oud-a2"] = {"id": "oud-a2", "extendedProperties": {"private": {
        "bron": "beveiligingsrooster", "dienst_id": str(dienst_a.id)}}}
    sync.sync_volledig(klant, a)
    assert "oud-b" in klant.agendas["team"] and "oud-a" in klant.agendas["team"]
    assert "oud-a2" not in klant.agendas["team"]


# ---------------------------------------------------------------------------
# H5 · Excel-import koppelt op initialen aan de verkeerde medewerker
# ---------------------------------------------------------------------------

def test_h5_alleen_initialen_gelijk_geeft_nieuwe_medewerker(app, klaar, tmp_path):
    from app.models import Dienst, Medewerker
    from app.services.excel_import import importeer, lees_bestand

    from .test_import_backup import maak_testbestand

    bestaand = Medewerker(naam="Iemand Anders", initialen="MVA")  # zelfde initialen
    zelfde_naam = Medewerker(naam="Medewerker Vijf B", initialen="XYZ")
    db.session.add_all([bestaand, zelfde_naam])
    db.session.commit()
    pad = str(tmp_path / "oud.xlsx")
    maak_testbestand(pad)

    plan = lees_bestand(pad)
    koppeling = {m.naam: (m.koppeling, m.bestaand_id) for m in plan.medewerkers}
    assert koppeling["Medewerker Vijf A"] == ("initialen", None)
    assert koppeling["Medewerker Vijf B"] == ("naam", zelfde_naam.id)
    assert any("MVA" in w and "Iemand Anders" in w for w in plan.waarschuwingen)

    importeer(plan)
    nieuw = Medewerker.query.filter_by(naam="Medewerker Vijf A").one()
    assert nieuw.id != bestaand.id and nieuw.initialen not in ("MVA", "")
    assert Dienst.query.filter_by(medewerker_id=bestaand.id).count() == 0
    assert Dienst.query.filter_by(medewerker_id=zelfde_naam.id).count() == 1


# ---------------------------------------------------------------------------
# H6 · Back-up van een nieuwere versie terugzetten legt de app plat
# M6 · Uitkomst van integrity_check werd genegeerd
# ---------------------------------------------------------------------------

@pytest.fixture
def gemigreerd(app):
    from flask_migrate import upgrade

    db.drop_all()
    upgrade(directory=os.path.join(os.path.dirname(__file__), "..", "migrations"))
    instellingen.schrijf("setup_voltooid", "1")
    db.session.commit()
    return app


def _revisie() -> str:
    with db.engine.connect() as verbinding:
        return verbinding.exec_driver_sql("SELECT version_num FROM alembic_version").scalar()


def test_h6_backup_van_nieuwere_versie_wordt_geweigerd(gemigreerd):
    from app.models import Medewerker

    pad = backup.maak_backup("test")
    verbinding = sqlite3.connect(pad)
    verbinding.execute("UPDATE alembic_version SET version_num = '0009_toekomst'")
    verbinding.commit()
    verbinding.close()
    db.session.add(Medewerker(naam="Na de back-up", initialen="NB"))
    db.session.commit()
    revisie = _revisie()

    with pytest.raises(ValueError, match="nieuwere versie"):
        backup.zet_terug(pad)
    db.session.remove()
    assert Medewerker.query.filter_by(initialen="NB").count() == 1
    assert _revisie() == revisie


def test_h6_mislukte_upgrade_zet_veiligheidsbackup_terug(gemigreerd, monkeypatch):
    import flask_migrate

    from app.models import Medewerker

    pad = backup.maak_backup("test")
    db.session.add(Medewerker(naam="Na de back-up", initialen="NB"))
    db.session.commit()

    def kapotte_upgrade(*args, **kwargs):
        raise SystemExit(1)  # zo stopt Flask-Migrate bij een fout

    monkeypatch.setattr(flask_migrate, "upgrade", kapotte_upgrade)
    with pytest.raises(ValueError, match="teruggezet"):
        backup.zet_terug(pad)
    db.session.remove()
    assert Medewerker.query.filter_by(initialen="NB").count() == 1


def test_m6_beschadigde_backup_wordt_geweigerd(gemigreerd):
    from app.models import Medewerker

    db.session.add_all([Medewerker(naam=f"Medewerker {i}", initialen=f"M{i}", ics_token=f"t{i}")
                        for i in range(20)])
    db.session.commit()
    pad = backup.maak_backup("test")
    verbinding = sqlite3.connect(pad)
    verbinding.execute("PRAGMA writable_schema = ON")
    verbinding.execute("UPDATE sqlite_master SET sql = 'CREATE INDEX ix_medewerker_ics_token "
                       "ON medewerker (naam)' WHERE name = 'ix_medewerker_ics_token'")
    verbinding.commit()
    verbinding.close()
    with pytest.raises(ValueError, match="beschadigd"):
        backup.controleer_backupbestand(pad)


def test_h6_scherm_ruimt_upload_op_bij_elke_fout(gemigreerd, client, monkeypatch):
    import io

    from .conftest import login, maak_gebruiker

    maak_gebruiker("beheerder", "beheerder")
    login(client, "beheerder")

    def kapot(pad):
        raise RuntimeError("onverwacht")

    monkeypatch.setattr(backup, "zet_terug", kapot)
    antwoord = client.post("/beheer/backups/terugzetten", data={
        "bevestig": "1", "bestand": (io.BytesIO(b"x"), "x.db")},
        content_type="multipart/form-data", follow_redirects=True)
    assert antwoord.status_code == 200 and "Terugzetten mislukt" in antwoord.data.decode()
    assert not [b for b in backup.lijst_backups() if b["naam"].endswith("-upload.db")]


# ---------------------------------------------------------------------------
# M1 · inf/nan als toeslagfactor
# M2 · tijdzone als vrije tekst, drie tijdzonebronnen
# M3 · blanco-code gelijk aan een bestaande dienstcode
# ---------------------------------------------------------------------------

def _instellingen_formulier(**anders) -> dict:
    formulier = {s: instellingen.lees(s) for s in instellingen.STANDAARD}
    formulier.update({"eerste_jaar": "2026", "tijdzone": "Europe/Amsterdam"})
    formulier.pop("deellink_actief")
    formulier.pop("opmerkingtijden_meetellen")
    formulier.update(anders)
    return formulier


@pytest.mark.parametrize("waarde", ["inf", "nan", "-inf", "1e400", "11", "0"])
def test_m1_toeslagfactor_moet_eindig_en_redelijk_zijn(app, als_beheerder, waarde):
    antwoord = als_beheerder.post("/beheer/instellingen",
                                  data=_instellingen_formulier(toeslag_zaterdag=waarde))
    assert antwoord.status_code == 400
    assert instellingen.lees("toeslag_zaterdag") == "1.5"


def test_m1_feestdagfactor_inf_geweigerd(app, als_beheerder):
    antwoord = als_beheerder.post("/beheer/instellingen", data=_instellingen_formulier(
        feestdagtoeslag_aan="1", toeslag_feestdag="inf"))
    assert antwoord.status_code == 400


def test_m1_getal_en_lees_float_weigeren_inf_nan(app, klaar):
    from app.blueprints.hulp import getal

    assert getal("inf") is None and getal("nan") is None and getal("1,5") == 1.5
    instellingen.schrijf("toeslag_zaterdag", "inf")
    assert instellingen.lees_float("toeslag_zaterdag") is None
    assert instellingen.toeslagen()["factor_zaterdag"] == 1.5


def test_m1_setup_weigert_inf(app, client):
    from app.services import setup_code

    from .conftest import WACHTWOORD

    client.get("/setup/")
    client.post("/setup/", data={"code": setup_code.lees_code()})
    client.post("/setup/stap/1", data={"gebruikersnaam": "planner", "weergavenaam": "P",
                                       "wachtwoord": WACHTWOORD, "herhaling": WACHTWOORD})
    formulier = {"teamnaam": "T", "tijdzone": "Europe/Amsterdam", "eerste_jaar": "2026",
                 "toeslag_zaterdag": "nan", "toeslag_zondag": "2", "logboek_dagen": "31",
                 "logboek_uren": "0"}
    assert client.post("/setup/stap/2", data=formulier).status_code == 400
    formulier.update(toeslag_zaterdag="1,5", tijdzone="Mars/Olympus")
    assert client.post("/setup/stap/2", data=formulier).status_code == 400
    assert instellingen.lees("tijdzone") != "Mars/Olympus"


def test_m2_ongeldige_tijdzone_geweigerd(app, als_beheerder):
    antwoord = als_beheerder.post("/beheer/instellingen",
                                  data=_instellingen_formulier(tijdzone="Mars/Olympus"))
    assert antwoord.status_code == 400 and "tijdzone" in antwoord.data.decode().lower()
    assert instellingen.lees("tijdzone") != "Mars/Olympus"


def test_m2_een_tijdzonebron_voor_klok_ics_en_agenda(app, als_beheerder, mw):
    from app.models import Dienst
    from app.services import ics
    from app.services.google_agenda import afspraak_voor

    antwoord = als_beheerder.post("/beheer/instellingen",
                                  data=_instellingen_formulier(tijdzone="Europe/London"))
    assert antwoord.status_code == 302
    assert klok.tijdzone_naam() == "Europe/London"
    db.session.add(Dienst(medewerker_id=mw.id, datum=klok.vandaag(), dienstnaam_override="X",
                          begin="07:00", eind="15:00", opmerking_tekst=""))
    db.session.commit()
    feed = ics.maak_feed(mw)
    assert "X-WR-TIMEZONE:Europe/London" in feed
    assert ics._utc("2026-07-01T07:15:00") == "20260701T061500Z"  # Londen: UTC+1 in de zomer
    assert afspraak_voor(Dienst.query.one()).body["start"]["timeZone"] == "Europe/London"


def test_m3_blanco_code_mag_geen_bestaande_dienstcode_zijn(app, als_beheerder, mw):
    antwoord = als_beheerder.post("/beheer/instellingen", data=_instellingen_formulier(blanco_code="4"))
    assert antwoord.status_code == 400 and "blanco" in antwoord.data.decode().lower()
    assert instellingen.lees("blanco_code") == "15"


# ---------------------------------------------------------------------------
# M4 · ImportError slikte echte fouten; een sync-fout blokkeerde de back-up
# ---------------------------------------------------------------------------

def test_m4_sync_fout_blokkeert_backup_niet(app, klaar, monkeypatch, caplog):
    from app.services import sync

    def kapot(*args, **kwargs):
        raise RuntimeError("sync stuk")

    monkeypatch.setattr(sync, "verwerk_wachtrij", kapot)
    with caplog.at_level("ERROR"):
        worker.een_ronde(worker.Planning(), datetime(2026, 3, 2, 3, 0))
    assert len(backup.lijst_backups()) == 1
    assert "agenda-synchronisatie" in caplog.text


def test_m4_importerror_wordt_niet_ingeslikt(app, klaar, monkeypatch, caplog):
    from app.services import sync

    def kapot(*args, **kwargs):
        raise ImportError("google_auth_httplib2 ontbreekt")

    monkeypatch.setattr(sync, "verwerk_wachtrij", kapot)
    with caplog.at_level("ERROR"):
        worker.een_ronde(worker.Planning(), datetime(2026, 3, 2, 1, 0))
    assert "google_auth_httplib2" in caplog.text


def test_m4_google_afhankelijkheden_in_requirements():
    import pathlib

    tekst = (pathlib.Path(__file__).parent.parent / "requirements.txt").read_text().lower()
    assert "google-auth-httplib2" in tekst and "httplib2" in tekst.replace("google-auth-httplib2", "")


# ---------------------------------------------------------------------------
# M5 · Per-IP-limiet blokkeert iedereen achter een proxy
# ---------------------------------------------------------------------------

def test_m5_juist_wachtwoord_komt_door_een_ip_blokkade(app, client, klaar):
    from .conftest import login

    for i in range(20):  # veel verschillende (fictieve) namen vanaf hetzelfde IP
        client.post("/login", data={"gebruikersnaam": f"onbekend{i}", "wachtwoord": "fout"})
    # Een nieuwe naam met een fout wachtwoord: geblokkeerd
    assert client.post("/login", data={"gebruikersnaam": "nogeen", "wachtwoord": "fout"}).status_code == 429
    # Een collega met het juiste wachtwoord komt er wel in
    assert login(client, "collega").status_code == 302


def test_m5_ip_blokkade_geeft_geen_extra_raadpogingen(app, client, klaar):
    for i in range(20):
        client.post("/login", data={"gebruikersnaam": f"onbekend{i}", "wachtwoord": "fout"})
    client.post("/login", data={"gebruikersnaam": "collega", "wachtwoord": "fout1"})
    # Na één fout tijdens een IP-blokkade helpt zelfs het juiste wachtwoord niet meer
    from .conftest import login

    assert login(client, "collega").status_code == 429


def test_m5_waarschuwing_bij_forwarded_for_zonder_proxy(app, client, klaar, caplog):
    from app.blueprints import auth

    auth._proxy_gewaarschuwd = False
    with caplog.at_level("WARNING"):
        client.post("/login", data={"gebruikersnaam": "x", "wachtwoord": "y"},
                    headers={"X-Forwarded-For": "203.0.113.9"})
    assert "PROXY_VERTROUWEN" in caplog.text


# ---------------------------------------------------------------------------
# M7 · Wachtwoord wijzigen/resetten en deactiveren logt andere sessies niet uit
# M10 · Na terugzetten blijven sessies geldig op gebruikers-ID
# ---------------------------------------------------------------------------

def _uitgelogd(client) -> bool:
    """True als deze sessie naar het inlogscherm wordt gestuurd."""
    antwoord = client.get("/kalender/")
    return antwoord.status_code == 302 and "/login" in antwoord.headers["Location"]


def _twee_sessies(app, naam="beheerder"):
    from .conftest import login

    eerste, tweede = app.test_client(), app.test_client()
    login(eerste, naam)
    login(tweede, naam)
    assert eerste.get("/kalender/").status_code == 200 and tweede.get("/kalender/").status_code == 200
    return eerste, tweede


def test_m7_wachtwoord_wijzigen_logt_andere_sessie_uit(app, klaar):
    from .conftest import WACHTWOORD

    eerste, tweede = _twee_sessies(app)
    antwoord = eerste.post("/account/wachtwoord", data={
        "huidig": WACHTWOORD, "nieuw": "nieuwwachtwoord1", "herhaling": "nieuwwachtwoord1"})
    assert antwoord.status_code == 302
    assert eerste.get("/kalender/").status_code == 200  # wie wijzigde blijft ingelogd
    assert _uitgelogd(tweede)  # de andere sessie niet


def test_m7_reset_en_deactiveren_loggen_gebruiker_uit(app, klaar):
    from app.models import Gebruiker

    from .conftest import login

    beheerder = app.test_client()
    login(beheerder, "beheerder")
    collega, _ = _twee_sessies(app, "collega")
    gid = Gebruiker.query.filter_by(gebruikersnaam="collega").one().id
    beheerder.post(f"/beheer/gebruikers/{gid}/reset")
    assert _uitgelogd(collega)

    collega2 = app.test_client()
    db.session.get(Gebruiker, gid).moet_wachtwoord_wijzigen = False
    db.session.commit()
    login(collega2, "collega", "x")  # wachtwoord is nu onbekend: direct via de sessie testen
    gebruiker = db.session.get(Gebruiker, gid)
    with collega2.session_transaction() as sessie:
        sessie["_user_id"] = gebruiker.get_id()
        sessie["_fresh"] = True
    assert collega2.get("/kalender/").status_code == 200
    beheerder.post(f"/beheer/gebruikers/{gid}", data={
        "gebruikersnaam": "collega", "weergavenaam": "Collega", "rol": "gebruiker"})  # actief uit
    assert _uitgelogd(collega2)


def test_m7_oud_sessieformaat_is_ongeldig(app, klaar):
    client = app.test_client()
    with client.session_transaction() as sessie:
        sessie["_user_id"] = str(klaar["beheerder"].id)  # zoals vóór 1.3.0
        sessie["_fresh"] = True
    assert _uitgelogd(client)


def test_m10_terugzetten_logt_iedereen_uit(gemigreerd):
    from .conftest import login, maak_gebruiker

    maak_gebruiker("beheerder", "beheerder")
    maak_gebruiker("collega", "gebruiker")
    beheerder, collega = gemigreerd.test_client(), gemigreerd.test_client()
    login(beheerder, "beheerder")
    login(collega, "collega")
    beheerder.post("/beheer/backups/maken")
    naam = backup.lijst_backups()[0]["naam"]
    assert collega.get("/kalender/").status_code == 200
    antwoord = beheerder.post("/beheer/backups/terugzetten", data={"naam": naam, "bevestig": "1"})
    assert antwoord.status_code == 302
    assert _uitgelogd(collega) and _uitgelogd(beheerder)


# ---------------------------------------------------------------------------
# M8 · Optimistic locking was check-then-write (TOCTOU)
# ---------------------------------------------------------------------------

def test_m8_gelijktijdige_wijziging_geeft_conflict(app, mw, monkeypatch):
    from app.models import Dienst
    from app.services import weekrooster
    from app.services.weekrooster import VersieConflict, Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4")])
    dienst = Dienst.query.one()
    versie = dienst.versie
    echt = weekrooster._pas_veld_toe

    def met_tussenkomst(dienst, veld, waarde):
        # Precies tussen lezen en schrijven wijzigt een ander proces dezelfde dienst
        with db.engine.begin() as verbinding:
            verbinding.exec_driver_sql(
                "UPDATE dienst SET eind = '20:00', versie = versie + 1 WHERE id = ?", (dienst.id,))
        return echt(dienst, veld, waarde)

    monkeypatch.setattr(weekrooster, "_pas_veld_toe", met_tussenkomst)
    with pytest.raises(VersieConflict):
        wijzig_cellen([Wijziging(mw.id, MAANDAG, "eind", "17:00", versie=versie)])
    db.session.expire_all()
    assert Dienst.query.one().eind == "20:00"  # de wijziging van de ander blijft staan


def test_m8_conflict_via_api_geeft_409(app, als_beheerder, mw, monkeypatch):
    from app.models import Dienst
    from app.services import weekrooster
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4")])
    dienst_id = Dienst.query.one().id
    echt = weekrooster._pas_veld_toe

    def met_tussenkomst(dienst, veld, waarde):
        with db.engine.begin() as verbinding:
            verbinding.exec_driver_sql("UPDATE dienst SET versie = versie + 1 WHERE id = ?", (dienst_id,))
        return echt(dienst, veld, waarde)

    monkeypatch.setattr(weekrooster, "_pas_veld_toe", met_tussenkomst)
    antwoord = als_beheerder.post("/api/cellen", json={"opslaan": True, "wijzigingen": [
        {"mw": mw.id, "datum": MAANDAG.isoformat(), "veld": "eind", "waarde": "17:00", "versie": 1}]})
    assert antwoord.status_code == 409


# ---------------------------------------------------------------------------
# M9 · De Excel-import plande geen agenda-sync
# ---------------------------------------------------------------------------

def test_m9_import_plant_agenda_sync(app, klaar, tmp_path):
    from app.models import Medewerker, SyncTaak
    from app.services.excel_import import importeer, lees_bestand

    from .test_import_backup import maak_testbestand

    db.session.add(Medewerker(naam="Medewerker Vijf A", initialen="MVA", agenda_modus="B",
                              agenda_id="agenda-a"))
    db.session.commit()
    pad = str(tmp_path / "oud.xlsx")
    maak_testbestand(pad)
    importeer(lees_bestand(pad))
    taken = SyncTaak.query.all()
    assert [(t.soort, db.session.get(Medewerker, t.medewerker_id).naam) for t in taken] \
        == [("volledig", "Medewerker Vijf A")]


# ---------------------------------------------------------------------------
# L1 · Taak verwijderd of onverwachte fout tijdens de sync
# ---------------------------------------------------------------------------

def test_l1_taak_intussen_verwijderd_geeft_geen_crash(app, gekoppeld, monkeypatch):
    from sqlalchemy.orm import Session

    from app.models import SyncTaak
    from app.services import google_agenda
    from app.services.google_agenda import AgendaFout
    from app.services.weekrooster import Wijziging, wijzig_cellen

    from .test_agenda import NepKlant

    class Klant(NepKlant):
        def maak_afspraak(self, agenda_id, body):
            with Session(db.engine) as sessie:  # een ander proces ruimt de taak op
                sessie.query(SyncTaak).delete()
                sessie.commit()
            raise AgendaFout("Google even weg", tijdelijk=True, status=503)

    monkeypatch.setattr(google_agenda, "klant", Klant)
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    _wachtrij_nu()  # mag niet crashen
    assert SyncTaak.query.count() == 0


def test_l1_onverwachte_fout_zet_taak_terug_met_backoff(app, gekoppeld, monkeypatch):
    from app.models import SyncTaak
    from app.services import google_agenda
    from app.services.weekrooster import Wijziging, wijzig_cellen

    from .test_agenda import NepKlant

    class Klant(NepKlant):
        def maak_afspraak(self, agenda_id, body):
            raise RuntimeError("onverwacht")

    monkeypatch.setattr(google_agenda, "klant", Klant)
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4"),
                   Wijziging(gekoppeld.id, MAANDAG + timedelta(days=1), "code", "4")])
    _wachtrij_nu()
    taken = SyncTaak.query.all()
    assert [(t.status, t.pogingen) for t in taken] == [("wacht", 1), ("wacht", 1)]
    assert all("onverwacht" in t.laatste_fout for t in taken)


# ---------------------------------------------------------------------------
# L2 · '²' en andere 'cijfers' gaven een 500
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("invoer", ["²", "7:²", "²:30", "١٢", "7:٣"])
def test_l2_vreemde_cijfers_zijn_ongeldige_tijd(invoer):
    from app.services.tijden import OngeldigeTijd, normaliseer_tijd

    with pytest.raises(OngeldigeTijd):
        normaliseer_tijd(invoer)


def test_l2_vreemde_cijfers_als_code_geven_celfout(app, als_beheerder, mw):
    for waarde in ("²", "١٢"):
        antwoord = als_beheerder.post("/api/cellen", json={"opslaan": True, "wijzigingen": [
            {"mw": mw.id, "datum": MAANDAG.isoformat(), "veld": "code", "waarde": waarde}]})
        assert antwoord.status_code == 200 and antwoord.json["fouten"]
    antwoord = als_beheerder.post("/api/cellen", json={"opslaan": True, "wijzigingen": [
        {"mw": mw.id, "datum": MAANDAG.isoformat(), "veld": "begin", "waarde": "²"}]})
    assert antwoord.status_code == 200 and antwoord.json["fouten"]


# ---------------------------------------------------------------------------
# L3 · Ongeldige body en datums in /api/cellen; jaar in /beheer/feestdagen
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("body", [[1, 2], "tekst", 5])
def test_l3_body_moet_een_object_zijn(app, als_beheerder, mw, body):
    assert als_beheerder.post("/api/cellen", json=body).status_code == 400


@pytest.mark.parametrize("datum", ["0001-01-01", "9999-12-31", "1949-12-31", "2151-01-01"])
def test_l3_datum_buiten_bereik(app, als_beheerder, mw, datum):
    antwoord = als_beheerder.post("/api/cellen", json={"opslaan": True, "wijzigingen": [
        {"mw": mw.id, "datum": datum, "veld": "code", "waarde": "4"}]})
    assert antwoord.status_code == 400
    antwoord = als_beheerder.post("/api/cellen", json={"dagopmerkingen": [{"datum": datum, "tekst": "x"}]})
    assert antwoord.status_code == 400


@pytest.mark.parametrize("jaar", ["99999", "-5", "1"])
def test_l3_feestdagen_jaar_begrensd(app, als_beheerder, jaar):
    assert als_beheerder.get(f"/beheer/feestdagen?jaar={jaar}").status_code == 200


# ---------------------------------------------------------------------------
# L4 · Agenda bijwerken na wijziging van codenummer, vakanties en feestdagen
# ---------------------------------------------------------------------------

def _gekoppeld_met_dienst(mw, dag):
    from app.models import SyncTaak
    from app.services.weekrooster import Wijziging, wijzig_cellen

    mw.agenda_modus, mw.agenda_id = "B", "agenda-a"
    db.session.commit()
    wijzig_cellen([Wijziging(mw.id, dag, "code", "4")])
    SyncTaak.query.delete()
    db.session.commit()


def test_l4_nieuw_codenummer_plant_agenda(app, als_beheerder, mw):
    from app.models import Dienstcode, SyncTaak

    dag = klok.vandaag() + timedelta(days=3)
    _gekoppeld_met_dienst(mw, dag)
    code = Dienstcode.query.filter_by(nummer=4).one()
    als_beheerder.post(f"/beheer/dienstcodes/{code.id}", data={
        "nummer": "40", "omschrijving": code.omschrijving, "std_begin": "07:15", "std_eind": "15:45",
        "vet": "1", "actief": "1", "in_agenda": "1"})
    assert Dienstcode.query.filter_by(nummer=40).count() == 1
    assert SyncTaak.query.filter_by(datum=dag).count() == 1


def test_l4_vakantie_en_feestdag_plannen_agenda(app, als_beheerder, mw):
    from app.models import Feestdag, SyncTaak, Vakantie

    dag = klok.vandaag() + timedelta(days=3)
    _gekoppeld_met_dienst(mw, dag)
    als_beheerder.post("/beheer/vakanties/opslaan", data={
        "naam": "Herfstvakantie", "datum_van": dag.isoformat(), "datum_tot": dag.isoformat()})
    assert SyncTaak.query.filter_by(datum=dag).count() == 1
    SyncTaak.query.delete()
    db.session.commit()
    vakantie = Vakantie.query.one()
    als_beheerder.post(f"/beheer/vakanties/{vakantie.id}/verwijder")
    assert SyncTaak.query.filter_by(datum=dag).count() == 1
    SyncTaak.query.delete()
    db.session.commit()
    als_beheerder.post("/beheer/feestdagen/nieuw", data={"naam": "Teamdag", "datum": dag.isoformat()})
    assert SyncTaak.query.filter_by(datum=dag).count() == 1
    SyncTaak.query.delete()
    db.session.commit()
    feestdag = Feestdag.query.filter_by(naam="Teamdag").one()
    als_beheerder.post(f"/beheer/feestdagen/{feestdag.id}/wissel")
    assert SyncTaak.query.filter_by(datum=dag).count() == 1


# ---------------------------------------------------------------------------
# L5 · Week kopiëren nam gearchiveerde medewerkers mee
# ---------------------------------------------------------------------------

def test_l5_kopieer_week_respecteert_archiefdatum(app, mw):
    from app.models import Dienst
    from app.services.weekrooster import Wijziging, kopieer_week, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG + timedelta(days=i), "code", "4") for i in range(7)])
    doel = MAANDAG + timedelta(days=7)
    mw.gearchiveerd_vanaf = doel + timedelta(days=2)  # vanaf woensdag gearchiveerd
    db.session.commit()
    kopieer_week(MAANDAG, doel)
    gekopieerd = Dienst.query.filter(Dienst.datum >= doel).all()
    assert sorted(d.datum for d in gekopieerd) == [doel, doel + timedelta(days=1)]


# ---------------------------------------------------------------------------
# L6 · Feestdagen dubbel aanmaken (race)
# ---------------------------------------------------------------------------

def test_l6_feestdag_maar_een_keer_per_jaar(app, klaar):
    from sqlalchemy.exc import IntegrityError

    from app.models import Feestdag
    from app.services import feestdagen

    feestdagen.zorg_voor_jaar(2026)
    db.session.add(Feestdag(jaar=2026, datum=datetime(2026, 4, 27).date(), naam="Dubbel",
                            sleutel="koningsdag"))
    with pytest.raises(IntegrityError):
        db.session.commit()
    db.session.rollback()
    # Eigen dagen (zonder sleutel) mogen wel meerdere keren
    db.session.add_all([Feestdag(jaar=2026, datum=MAANDAG, naam="Eigen", sleutel=""),
                        Feestdag(jaar=2026, datum=MAANDAG, naam="Eigen 2", sleutel="")])
    db.session.commit()


def test_l6_gelijktijdig_aanmaken_geeft_geen_fout(app, klaar, monkeypatch):
    from app.models import Feestdag
    from app.services import feestdagen

    feestdagen.zorg_voor_jaar(2026)
    # Alsof een ander proces ze net tegelijk aanmaakte: wij 'zien' ze nog niet
    monkeypatch.setattr(feestdagen, "_bestaande_sleutels", lambda jaar: set(), raising=False)
    feestdagen.zorg_voor_jaar(2026)
    assert Feestdag.query.filter_by(jaar=2026).count() == 11


def test_l6_migratie_ruimt_dubbele_feestdagen_op(tmp_path):
    from flask_migrate import upgrade

    from app import create_app

    from .conftest import TestConfig

    app = create_app(TestConfig(str(tmp_path)))
    map_ = os.path.join(os.path.dirname(__file__), "..", "migrations")
    with app.app_context():
        upgrade(directory=map_, revision="0003")
        with db.engine.begin() as verbinding:
            for _ in range(3):
                verbinding.exec_driver_sql(
                    "INSERT INTO feestdag (jaar, datum, naam, sleutel, actief) "
                    "VALUES (2026, '2026-04-27', 'Koningsdag', 'koningsdag', 1)")
            verbinding.exec_driver_sql(
                "INSERT INTO feestdag (jaar, datum, naam, sleutel, actief) "
                "VALUES (2026, '2026-03-02', 'Eigen', '', 1)")
        upgrade(directory=map_)
        with db.engine.connect() as verbinding:
            assert verbinding.exec_driver_sql("SELECT count(*) FROM feestdag").scalar() == 2
        db.session.remove()


# ---------------------------------------------------------------------------
# L7 · Bewaartermijn 0 dagen + 0 uren wiste het hele logboek
# ---------------------------------------------------------------------------

def test_l7_bewaartermijn_nul_betekent_nooit_opschonen(app, klaar):
    from app.services import logboek

    instellingen.schrijf("logboek_dagen", "0")
    instellingen.schrijf("logboek_uren", "0")
    db.session.add(Logboek(actie="Oud", tijdstempel=datetime(2000, 1, 1)))
    db.session.commit()
    assert logboek.opschonen() == 0
    assert Logboek.query.filter_by(actie="Oud").count() == 1
    instellingen.schrijf("logboek_uren", "1")
    db.session.commit()
    assert logboek.opschonen() >= 1


# ---------------------------------------------------------------------------
# L8 · Het dubbele uur bij de wintertijd
# ---------------------------------------------------------------------------

def test_l8_loginblokkade_rekent_in_utc(app, client, klaar, monkeypatch):
    from .conftest import login

    lokaal, utc = datetime(2026, 10, 25, 2, 50), datetime(2026, 10, 25, 0, 50)
    monkeypatch.setattr(klok, "nu", lambda: lokaal)
    monkeypatch.setattr(klok, "utc_nu", lambda: utc, raising=False)
    for _ in range(5):
        login(client, "collega", "fout")
    # De klok gaat om 03:00 terug naar 02:00; 16 minuten later is het lokaal 02:06
    monkeypatch.setattr(klok, "nu", lambda: datetime(2026, 10, 25, 2, 6))
    monkeypatch.setattr(klok, "utc_nu", lambda: datetime(2026, 10, 25, 1, 6), raising=False)
    assert login(client, "collega").status_code == 302


def test_l8_wachtrij_rekent_in_utc(app, gekoppeld, monkeypatch):
    from app.services import google_agenda, sync
    from app.services.weekrooster import Wijziging, wijzig_cellen

    from .test_agenda import NepKlant

    klant = NepKlant()
    monkeypatch.setattr(google_agenda, "klant", lambda: klant)
    monkeypatch.setattr(klok, "nu", lambda: datetime(2026, 10, 25, 2, 55))
    monkeypatch.setattr(klok, "utc_nu", lambda: datetime(2026, 10, 25, 0, 55), raising=False)
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    # 10 minuten later, ná het terugzetten van de klok
    monkeypatch.setattr(klok, "nu", lambda: datetime(2026, 10, 25, 2, 5))
    monkeypatch.setattr(klok, "utc_nu", lambda: datetime(2026, 10, 25, 1, 5), raising=False)
    assert sync.verwerk_wachtrij() == 1


# ---------------------------------------------------------------------------
# S1 · Timing: onbekende gebruiker zonder wachtwoordcontrole
# ---------------------------------------------------------------------------

def test_s1_onbekende_gebruiker_kost_evenveel_werk(app, client, klaar, monkeypatch):
    from app.services import wachtwoorden

    echt = wachtwoorden._hasher
    teller = []

    class Teller:
        def __getattr__(self, naam):
            return getattr(echt, naam)

        def verify(self, *args):
            teller.append(1)
            return echt.verify(*args)

    monkeypatch.setattr(wachtwoorden, "_hasher", Teller())
    client.post("/login", data={"gebruikersnaam": "bestaat-niet", "wachtwoord": "x"})
    assert len(teller) == 1


# ---------------------------------------------------------------------------
# S2 · CSV-export: formules in cellen
# ---------------------------------------------------------------------------

def test_s2_csv_zonder_formules(app, als_beheerder, klaar):
    from app.models import Dienst, Medewerker

    m = Medewerker(naam="=HYPERLINK(\"http://x\")", initialen="FX")
    db.session.add(m)
    db.session.commit()
    db.session.add(Dienst(medewerker_id=m.id, datum=MAANDAG, dienstnaam_override="@SUM(A1)",
                          begin="07:00", eind="15:00", uren_berekend=7.5, opmerking_tekst=""))
    db.session.commit()
    uren = als_beheerder.get("/overzicht/uren.csv?jaar=2026").data.decode()
    assert "'=HYPERLINK" in uren and ";=HYPERLINK" not in uren
    zoek = als_beheerder.get("/zoeken/export.csv?naam=HYPER").data.decode()
    assert "'=HYPERLINK" in zoek and "'@SUM" in zoek


def test_s2_negatieve_getallen_blijven_getallen():
    from app.blueprints.hulp import csv_cel

    assert csv_cel("-3,50") == "-3,50" and csv_cel("-tekst") == "'-tekst" and csv_cel("+31") == "'+31"
    assert csv_cel("Gewoon") == "Gewoon" and csv_cel(5) == 5


# ---------------------------------------------------------------------------
# S3 · Tokens in de toegangslog
# ---------------------------------------------------------------------------

def test_s3_tokens_gemaskeerd_in_toegangslog():
    import runpy

    from app.toegangslog import maskeer_tokens

    assert maskeer_tokens("GET /ics/abcdefghijklmnopqrstuvwx.ics HTTP/1.1") == "GET /ics/***.ics HTTP/1.1"
    assert maskeer_tokens("/deel/geheim123/week/2026/10?dag=x") == "/deel/***/week/2026/10?dag=x"
    assert maskeer_tokens("/kalender/") == "/kalender/"
    conf = runpy.run_path(os.path.join(os.path.dirname(__file__), "..", "docker", "gunicorn.conf.py"))
    assert conf["logger_class"] == "app.toegangslog.ToegangsLogger"


def test_s3_logger_maskeert_atomen():
    from types import SimpleNamespace

    from gunicorn.config import Config

    from app.toegangslog import ToegangsLogger

    logger = ToegangsLogger(Config())
    verzoek = SimpleNamespace(headers=[], method="GET", path="/ics/geheimtoken.ics", query="",
                              version=(1, 1))
    antwoord = SimpleNamespace(status="200 OK", headers=[], sent=10, response_length=10)
    omgeving = {"RAW_URI": "/ics/geheimtoken.ics", "REQUEST_METHOD": "GET", "QUERY_STRING": "",
                "SERVER_PROTOCOL": "HTTP/1.1", "REMOTE_ADDR": "1.2.3.4"}
    atomen = logger.atoms(antwoord, verzoek, omgeving, timedelta(seconds=0.1))
    assert "geheimtoken" not in " ".join(str(v) for v in atomen.values())


# ---------------------------------------------------------------------------
# S5 · Duidelijke melding bij verkeerde bestandsrechten
# ---------------------------------------------------------------------------

def test_s5_geheime_sleutel_zonder_schrijfrechten(tmp_path, monkeypatch):
    from app import _lees_of_maak_geheime_sleutel

    def geen_rechten(*args, **kwargs):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(os, "open", geen_rechten)
    with pytest.raises(RuntimeError, match="-u rooster"):
        _lees_of_maak_geheime_sleutel(str(tmp_path))


def test_s5_setup_code_zonder_schrijfrechten(app, monkeypatch):
    from app.services import setup_code

    def geen_rechten(*args, **kwargs):
        raise PermissionError(13, "Permission denied")

    monkeypatch.setattr(os, "open", geen_rechten)
    with pytest.raises(RuntimeError, match="-u rooster"):
        setup_code.haal_of_maak_code()


def test_s5_cli_docstring_gebruikt_u_rooster():
    from app import cli

    regels = [r for r in cli.__doc__.splitlines() if "docker compose exec" in r]
    assert regels and all("-u rooster" in r for r in regels)


# ---------------------------------------------------------------------------
# S6 · Achtergebleven importbestanden
# ---------------------------------------------------------------------------

def test_s6_worker_ruimt_oude_importbestanden_op(app, klaar):
    import time

    map_ = os.path.join(app.config["DATA_MAP"], "import")
    os.makedirs(map_, exist_ok=True)
    oud, nieuw = os.path.join(map_, "oud.xlsm"), os.path.join(map_, "nieuw.xlsm")
    for pad in (oud, nieuw):
        open(pad, "wb").close()
    twee_dagen = time.time() - 2 * 24 * 3600
    os.utime(oud, (twee_dagen, twee_dagen))
    worker.een_ronde(worker.Planning(), datetime(2026, 3, 2, 1, 0))
    assert not os.path.exists(oud) and os.path.exists(nieuw)


# ---------------------------------------------------------------------------
# Overig: zoeken, BASE_URL, gelabelde back-ups, kleurregels, import
# ---------------------------------------------------------------------------

def test_zoeken_wildcards_worden_letterlijk_gezocht(app, als_beheerder, mw):
    from app.services.overzichten import zoek_diensten
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4")])
    assert zoek_diensten(naam="___") == [] and zoek_diensten(naam="%%%") == []
    assert len(zoek_diensten(naam="ewerk")) == 1


def test_base_url_voor_externe_links(app, als_beheerder, klaar, mw):
    from app.models import Gebruiker

    app.config["BASE_URL"] = "https://rooster.voorbeeld.nl"
    mw.ics_token = "t" * 30
    gebruiker = db.session.get(Gebruiker, klaar["beheerder"].id)
    gebruiker.medewerker_id = mw.id
    db.session.commit()
    pagina = als_beheerder.get("/mijn/agenda").data.decode()
    assert f"https://rooster.voorbeeld.nl/ics/{'t' * 30}.ics" in pagina
    pagina = als_beheerder.get("/beheer/agenda").data.decode()
    assert f"https://rooster.voorbeeld.nl/ics/{'t' * 30}.ics" in pagina


def test_gelabelde_backups_hebben_een_bewaartermijn(app, klaar):
    map_ = backup.backup_map()
    for dag in range(1, 16):  # 15 oude handmatige back-ups uit januari 2025
        with open(os.path.join(map_, f"rooster-202501{dag:02d}-120000-handmatig.db"), "wb") as f:
            f.write(b"x")
    with open(os.path.join(map_, "rooster-20250101-020000.db"), "wb") as f:
        f.write(b"x")
    backup.ruim_gelabelde_op()
    namen = [b["naam"] for b in backup.lijst_backups()]
    assert len([n for n in namen if "handmatig" in n]) == backup.GELABELD_MINIMAAL
    assert "rooster-20250115-120000-handmatig.db" in namen  # de nieuwste blijven
    assert "rooster-20250101-020000.db" in namen  # automatische vallen hierbuiten


def test_kleurregel_uniek_zonder_hoofdletters_en_max_60(app, als_beheerder, klaar):
    from app.models import OpmerkingKleurregel

    als_beheerder.post("/beheer/kleurregels/nieuw", data={"tekst": "Locatie A"})
    als_beheerder.post("/beheer/kleurregels/nieuw", data={"tekst": "locatie a"})
    als_beheerder.post("/beheer/kleurregels/nieuw", data={"tekst": "x" * 80})
    teksten = [r.tekst for r in OpmerkingKleurregel.query.all()]
    assert teksten.count("Locatie A") == 1 and "locatie a" not in teksten
    assert all(len(t) <= 60 for t in teksten)


def test_import_jaar_buiten_bereik(app, klaar, tmp_path):
    import openpyxl

    from app.services.excel_import import ImportFout, lees_bestand

    boek = openpyxl.Workbook()
    boek.active.title = "Lijsten"
    boek.create_sheet("Kalender")["E2"] = 1900
    pad = str(tmp_path / "oud.xlsx")
    boek.save(pad)
    with pytest.raises(ImportFout, match="jaar"):
        lees_bestand(pad)


def test_import_lange_dienstnaam_en_oude_toeslag_gelogd(app, mw):
    from app.models import Dienst
    from app.services.excel_import import ImportMedewerker, ImportPlan, importeer

    instellingen.schrijf("toeslag_zaterdag", "1.25")
    db.session.commit()
    plan = ImportPlan(jaar=2026, weken=[10], toeslag_zaterdag=1.5,
                      medewerkers=[ImportMedewerker("Medewerker A", "MA", None)],
                      diensten=[_import_dienst("Medewerker A", MAANDAG, dienstnaam="d" * 80)])
    importeer(plan)
    assert len(Dienst.query.one().dienstnaam_override) == 60
    regel = Logboek.query.filter_by(actie="Instelling gewijzigd", veld="toeslag_zaterdag").one()
    assert (regel.oude_waarde, regel.nieuwe_waarde) == ("1.25", "1.5")


def test_voorbeeldpakket_zonder_echte_plaatsnamen(app, klaar):
    from app.models import OpmerkingKleurregel
    from app.services.voorbeeldpakket import laad_voorbeeldpakket

    laad_voorbeeldpakket()
    assert sorted(r.tekst for r in OpmerkingKleurregel.query.all()) == ["Locatie A", "Locatie B"]


def test_dagnamen_staan_op_een_plek(app, als_beheerder, mw):
    from app.services.kalender import DAGNAMEN_KORT

    pagina = als_beheerder.get("/week/2026/10").data.decode()
    assert f"{DAGNAMEN_KORT[0]} 02-03-26" in pagina
