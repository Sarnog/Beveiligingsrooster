"""Browsertests (Playwright): Excel-import met jaarkeuze en overschrijfmodus, en de Excel-export (1.5.0).

- uploaden zonder jaar -> stap 'Rooster voor jaar' met het voorstel uit Kalender!E2 -> voorbeeld;
- overschrijfmodus 'gedeeltelijk' met één medewerker -> voorbeeld bijwerken -> alleen die
  medewerker in de tabel -> definitief importeren -> alleen die medewerker gewijzigd;
- sinds 1.6.0 alleen in Beheer → Excel import/export: een week exporteren geeft een
  .xlsx-download met formules; een ongeldige week geeft een melding op hetzelfde scherm.
"""

import io
from datetime import date

import openpyxl
import pytest

from .browser_hulp import nieuwe_pagina

sync_api = pytest.importorskip("playwright.sync_api")
pytestmark = pytest.mark.browser


def _diensten_van(server, naam: str, jaar: int) -> dict:
    from app.extensions import db
    from app.models import Dienst, Medewerker

    with server.app.app_context():
        medewerker = Medewerker.query.filter_by(naam=naam).first()
        if medewerker is None:
            db.session.remove()
            return {}
        resultaat = {(d.datum, d.volgnummer): d.dienstnaam for d in Dienst.query.filter(
            Dienst.medewerker_id == medewerker.id, Dienst.datum >= date(jaar, 1, 1))}
        db.session.remove()
        return resultaat


def test_import_met_jaarkeuze_en_overschrijfmodus(server, browser, sessies, tmp_path):
    from app.services import klok

    from .test_import_keuzes import maak_rooster

    jaar = klok.vandaag().year + 1  # volgend jaar naast het huidige
    pad = str(tmp_path / f"rooster-{jaar}.xlsx")
    maak_rooster(pad, jaar, {10: {"Medewerker X": {0: (4, "07:15", "15:45")},
                                  "Medewerker Y": {1: (4, "07:15", "15:45")}}})

    context, pagina = nieuwe_pagina(browser, "1280x800", sessies["beheerder"])
    pagina.goto(server.url + "/beheer/importeren")
    pagina.set_input_files("input[name=bestand]", pad)
    pagina.click("button:has-text('Uploaden')")
    # Stap 2: verplicht jaarveld, vooringevuld met Kalender!E2
    jaarveld = pagina.locator("[data-stap-jaar] input[name=jaar]")
    sync_api.expect(jaarveld).to_have_value(str(jaar))
    assert jaarveld.get_attribute("required") is not None
    pagina.click("button:has-text('Verder naar het voorbeeld')")
    sync_api.expect(pagina.locator("section.kaart h2").first).to_contain_text(f"rooster voor {jaar}")

    # Overschrijfmodus: gedeeltelijk, alleen Medewerker X
    pagina.check("input[name=modus][value=gedeeltelijk]")
    pagina.check("input[name=medewerkers][value='Medewerker X']")
    pagina.click("button[name=actie][value=voorbeeld]")
    tabel = pagina.locator("table[data-effect]")
    sync_api.expect(tabel).to_contain_text("Medewerker X")
    sync_api.expect(tabel).not_to_contain_text("Medewerker Y")
    sync_api.expect(pagina.locator("input[name=modus][value=gedeeltelijk]")).to_be_checked()

    # Een keuze wijzigen zonder nieuw voorbeeld: niet importeren, eerst opnieuw bekijken
    pagina.check("input[name=medewerkers][value='Medewerker Y']")
    pagina.check("input[name=bevestig]")
    pagina.click("button[name=actie][value=importeren]")
    sync_api.expect(pagina.locator(".melding").first).to_contain_text("gewijzigd")
    assert _diensten_van(server, "Medewerker X", jaar) == {}
    # Terug naar alleen X en bevestigen
    pagina.uncheck("input[name=medewerkers][value='Medewerker Y']")
    pagina.click("button[name=actie][value=voorbeeld]")
    pagina.check("input[name=bevestig]")
    pagina.click("button[name=actie][value=importeren]")
    pagina.wait_for_url(lambda url: "/kalender" in url)
    sync_api.expect(pagina.locator("body")).to_contain_text("Import klaar")
    assert list(_diensten_van(server, "Medewerker X", jaar).values()) == ["VW Vroeg"]
    assert _diensten_van(server, "Medewerker Y", jaar) == {}  # niet gekozen: niet geïmporteerd
    context.close()


def test_export_in_beheer_geeft_xlsx(server, browser, sessies):
    context, pagina = nieuwe_pagina(browser, "1280x800", sessies["beheerder"])
    context.set_default_timeout(20000)
    pagina.goto(server.url + "/beheer/importeren")
    formulier = pagina.locator("[data-export-formulier]")
    formulier.locator("input[name=soort][value=week]").check()
    formulier.locator("input[name=jaar]").fill("2026")
    formulier.locator("input[name=week]").fill("10")
    with pagina.expect_download() as wacht:
        formulier.locator("button[type=submit]").click()
    download = wacht.value
    assert download.suggested_filename == "rooster-2026-W10.xlsx"
    boek = openpyxl.load_workbook(io.BytesIO(open(download.path(), "rb").read()))
    assert "Lijsten" in boek.sheetnames and [n for n in boek.sheetnames if n.startswith("W")] == ["W10"]
    assert str(boek["W10"]["Z6"].value).startswith("=SUM(")  # weektotaal als formule
    # Ongeldige week (2025 heeft er 52): een melding op het scherm, geen foutpagina
    formulier.locator("input[name=jaar]").fill("2025")
    formulier.locator("input[name=week]").fill("53")
    formulier.locator("button[type=submit]").click()
    sync_api.expect(pagina.locator(".melding").first).to_contain_text("weken 1 t/m 52")
    sync_api.expect(pagina.locator("[data-export-formulier] input[name=week]")).to_have_value("53")
    context.close()


def test_collega_ziet_geen_excelexport(server, browser, sessies):
    context, pagina = nieuwe_pagina(browser, "1280x800", sessies["collega"])
    for pad in ("/week", "/kalender/", "/overzicht/uren"):
        pagina.goto(server.url + pad)
        assert pagina.locator("[data-export-excel]").count() == 0
        sync_api.expect(pagina.locator("body")).not_to_contain_text("Exporteren (Excel)")
    context.close()
