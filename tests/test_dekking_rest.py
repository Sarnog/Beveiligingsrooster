"""Kleine services, Excel-export en de worker: laatste randgevallen. Alle data is fictief."""

import io
import logging
from datetime import date, datetime, timedelta

import openpyxl
import pytest

from app.extensions import db
from app.models import Dienst, Dienstcode, Feestdag, Medewerker, OpmerkingKleurregel, SyncTaak
from app.services import (
    excel_export,
    feestdagen,
    ics,
    instellingen,
    klok,
    logboek,
    rooster,
    roosteracties,
    setup_code,
    statistieken,
    wachtwoorden,
)
from app.services.google_agenda import afspraak_voor
from app.services.voorbeeldpakket import laad_voorbeeldpakket

MAANDAG = date(2026, 3, 2)


@pytest.fixture
def mw(klaar):
    laad_voorbeeldpakket()
    medewerker = Medewerker(naam="Medewerker A", initialen="MA")
    db.session.add(medewerker)
    db.session.commit()
    return medewerker


# ---------------------------------------------------------------------------
# Excel-export
# ---------------------------------------------------------------------------

def test_formules_zonder_pauze_correcties_of_feestdagen():
    formules = excel_export._Formules(0, 0)
    assert formules.pauze("A1") == "0"
    assert "Rekenhulp" not in formules.uren("X", "B1", "C1", "1")
    assert formules.factor("A1", 0).startswith("IF(WEEKDAY(")


def test_correctie_buiten_bereik_wordt_overgeslagen(caplog):
    excel_export.correcties.cache_clear()
    with caplog.at_level(logging.WARNING, logger="app.services.excel_export"):
        excel_export.correcties(20.0, ((1.0, 0.75),))
    assert "correctie buiten bereik" in caplog.text


def test_export_periode_midden_in_de_week_met_lege_regels(app, mw):
    db.session.add(OpmerkingKleurregel(tekst="Effen", kleur_achtergrond="#00FF00", kleur_tekst="#000000"))
    code = Dienstcode.query.filter_by(nummer=4).one()
    db.session.add_all([
        Dienst(medewerker_id=mw.id, datum=MAANDAG + timedelta(days=2), volgnummer=1, versie=1,
               dienstcode_id=code.id, begin="07:15", eind="15:45"),
        Dienst(medewerker_id=mw.id, datum=MAANDAG + timedelta(days=2), volgnummer=2, versie=1,
               google_event_id="leeg-2"),  # lege tweede regel die nog op de agenda wacht
        Dienst(medewerker_id=mw.id, datum=MAANDAG + timedelta(days=3), volgnummer=1, versie=1,
               google_event_id="leeg-1"),  # lege eerste regel zonder dienst 2
    ])
    db.session.commit()
    keuze = excel_export.ExportKeuze("periode", 2026, MAANDAG + timedelta(days=2),
                                     MAANDAG + timedelta(days=3))
    boek = openpyxl.load_workbook(io.BytesIO(excel_export.maak_export(keuze)))
    blad = boek["W10"]
    assert blad.cell(6, 29 + 2).value == 4  # woensdag: code 4 in het raster
    assert blad.cell(6, 29).value is None and blad.cell(6, 29 + 3).value is None  # ma buiten, do leeg


# ---------------------------------------------------------------------------
# Feestdagen, agenda-afspraak en ICS
# ---------------------------------------------------------------------------

def test_twee_feestdagen_op_een_dag(app, klaar):
    koningsdag = date(2026, 4, 27)
    db.session.add(Feestdag(jaar=2026, datum=koningsdag, naam="Teamdag", sleutel=""))
    db.session.commit()
    namen = feestdagen.feestdagen_in_periode(koningsdag, koningsdag)[koningsdag]
    assert "Teamdag" in namen and " / " in namen


def test_afspraak_voor_randgevallen_en_ics_hele_dag(app, mw):
    assert afspraak_voor(None) is None
    met_opmerking = Dienst(medewerker_id=mw.id, datum=MAANDAG, dienstnaam_override="Cursus",
                           begin="08:00", eind="12:00", opmerking_tekst="Lokaal 3",
                           opmerking_begin="13:00", opmerking_eind="")
    assert "(13:00–)" in afspraak_voor(met_opmerking).body["description"]
    zonder_tijden = Dienst(medewerker_id=mw.id, datum=MAANDAG, dienstnaam_override="Cursus")
    assert afspraak_voor(zonder_tijden) is None
    hele_dag = Dienstcode.query.filter_by(nummer=4).one()
    hele_dag.hele_dag_zonder_tijden = True
    mw.ics_token = "t" * 30
    db.session.add_all([
        Dienst(medewerker_id=mw.id, datum=MAANDAG, volgnummer=1, versie=1, dienstcode_id=hele_dag.id),
        Dienst(medewerker_id=mw.id, datum=MAANDAG + timedelta(days=1), volgnummer=1, versie=1,
               dienstnaam_override="Zonder tijden"),
    ])
    db.session.commit()
    tekst = ics.maak_feed(mw, dagen_terug=3650)
    assert "DTSTART;VALUE=DATE:20260302" in tekst and "Zonder tijden" not in tekst


# ---------------------------------------------------------------------------
# Instellingen, klok, logboek, rooster, roosteracties
# ---------------------------------------------------------------------------

def test_instellingen_ongeldige_waarden(app, klaar):
    instellingen.schrijf("toeslag_feestdag", "veel")
    assert instellingen.lees_float("toeslag_feestdag") is None
    instellingen.schrijf("pauze", '{"aan": true, "regels": [[30, 1]]}')  # grens boven 24 uur
    assert instellingen.pauze_instelling() == {"aan": True, "regels": list(instellingen.STANDAARD_PAUZE)}


def test_klok_zonder_app_of_database(app, monkeypatch):
    klok.wis_cache()
    assert klok._lees_instelling() == ""  # geen tijdzone ingesteld

    def kapot(*_args, **_kwargs):
        raise RuntimeError("database weg")

    klok.wis_cache()
    monkeypatch.setattr(db.engine, "connect", kapot)
    assert klok._lees_instelling() == ""
    klok.wis_cache()


def test_klok_buiten_een_app_context():
    assert klok._lees_instelling() == ""


def test_logboek_opschonen_met_ongeldige_uren(app, klaar):
    instellingen.schrijf("logboek_dagen", "0")
    instellingen.schrijf("logboek_uren", "99")
    db.session.commit()
    assert logboek.opschonen(klok.nu()) == 0


def test_herberekenen_binnen_een_lege_periode_en_geen_afwijkende_tijden(app, mw):
    assert rooster.herbereken_alle(date(2030, 1, 1), date(2030, 1, 31)) == 0
    code = Dienstcode.query.filter_by(nummer=4).one()
    assert rooster.pas_std_tijden_toe(code, MAANDAG) == []


def test_plan_agenda_met_dag_binnen_de_syncperiode(app, mw):
    mw.agenda_modus, mw.agenda_id = "B", "agenda-a"
    db.session.commit()
    roosteracties.plan_agenda({mw: {klok.vandaag()}}, "test")
    assert SyncTaak.query.filter_by(soort="volledig").count() == 1
    assert SyncTaak.query.filter_by(soort="dag").count() == 0  # zit al in de volledige synchronisatie


# ---------------------------------------------------------------------------
# Setup-code, statistieken, voorbeeldpakket, wachtwoorden
# ---------------------------------------------------------------------------

def test_setup_code_randgevallen(app, monkeypatch):
    assert setup_code.controleer_code("") is False
    setup_code.verwijder_code()  # er is geen code: geen fout
    setup_code.verwijder_code()

    def geen_rechten(*_args, **_kwargs):
        raise PermissionError("geen toegang")

    monkeypatch.setattr("builtins.open", geen_rechten)
    with pytest.raises(RuntimeError):
        setup_code.lees_code()


def test_statistieken_met_gelabelde_backup(app, als_beheerder):
    from app.services import backup

    backup.maak_backup("handmatig")
    pagina = als_beheerder.get("/beheer/statistieken").data.decode()
    assert "Laatste back-up met label" in pagina and "-handmatig.db" in pagina
    assert statistieken.grootte_tekst(2048) == "2,0 kB" and statistieken.grootte_tekst(10) == "10 B"
    assert isinstance(statistieken.meldingen(aantal_dubbel=0), list)


def test_voorbeeldpakket_twee_keer_laden(app, klaar):
    eerste = laad_voorbeeldpakket()
    assert laad_voorbeeldpakket() == (0, 0) and eerste[0] > 0


def test_onleesbare_wachtwoordhash_moet_opnieuw():
    assert wachtwoorden.moet_opnieuw_hashen("geen-argon2-hash") is True


# ---------------------------------------------------------------------------
# Debuglog en worker
# ---------------------------------------------------------------------------

def test_debuglog_zonder_bestand_en_zonder_rechten(app, monkeypatch, tmp_path):
    from app import debuglog

    assert debuglog.laatste_regels(str(tmp_path / "leeg")) == []

    def geen_rechten(*_args, **_kwargs):
        raise PermissionError("geen toegang")

    monkeypatch.setattr(debuglog, "WatchedFileHandler", geen_rechten)
    app.config["DEBUG_LOG"] = True
    instellingen.schrijf("debug_log", "1")
    db.session.commit()
    with pytest.raises(RuntimeError):
        debuglog.ververs(app, direct=True)


def test_extensions_alleen_pragmas_voor_sqlite():
    from app.extensions import _sqlite_instellingen

    class Andere:
        def cursor(self):  # pragma: no cover - mag niet aangeroepen worden
            raise AssertionError("geen SQLite: geen PRAGMA's")

    _sqlite_instellingen(Andere(), None)


def test_api_token_zonder_voorvoegsel(app, klaar):
    antwoord = app.test_client().get("/api/v1/ik", headers={"Authorization": "Bearer geen-geldig-token"})
    assert antwoord.status_code == 401


def test_worker_ronde_debuglog_roteren_en_fouten(app, klaar, monkeypatch, caplog):
    from app import debuglog, worker

    monkeypatch.setattr(debuglog, "roteer", lambda _map: True)
    with caplog.at_level(logging.INFO):
        worker.een_ronde(worker.Planning(), klok.nu())
    assert "Debuglog geroteerd" in caplog.text
    monkeypatch.setattr(debuglog, "roteer", lambda _map: (_ for _ in ()).throw(OSError("weg")))
    worker.een_ronde(worker.Planning(), klok.nu())
    assert "Debuglog roteren mislukt" in caplog.text


def test_worker_backup_mislukt_en_logboek_ook(app, klaar, monkeypatch):
    from app import worker
    from app.services import backup

    monkeypatch.setattr(backup, "maak_backup", lambda *_: (_ for _ in ()).throw(OSError("schijf vol")))
    monkeypatch.setattr(logboek, "log", lambda *_a, **_k: (_ for _ in ()).throw(OSError("ook vol")))
    planning = worker.Planning()
    worker._stap_backup(planning, datetime(2026, 3, 2, 3, 0))
    assert planning.backup_niet_voor is not None and planning.backup_gedaan is None


def test_worker_main_stopt_netjes_na_een_fout(app, monkeypatch, caplog):
    from app import worker

    class EenRonde:
        """Event dat na één ronde 'gezet' is."""

        def __init__(self):
            self.rondes = 0

        def is_set(self):
            self.rondes += 1
            return self.rondes > 1

        def set(self):
            pass

        def wait(self, _seconden):
            pass

    monkeypatch.setattr(worker, "create_app", lambda: app)
    monkeypatch.setattr(worker.threading, "Event", EenRonde)
    monkeypatch.setattr(worker.signal, "signal", lambda *_: None)
    monkeypatch.setattr(worker, "een_ronde", lambda _planning: (_ for _ in ()).throw(RuntimeError("stuk")))
    with caplog.at_level(logging.INFO):
        worker.main()
    assert "Fout in worker-ronde" in caplog.text and "Worker gestopt" in caplog.text
