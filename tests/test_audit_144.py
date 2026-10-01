"""Regressietests voor de bevindingen uit de audit van versie 1.4.4.

Elke test hoort bij één bevinding (H = high, M = medium, L = low, S = security).
De test faalde op de oude code en slaagt na de reparatie.
"""

from datetime import date, datetime, timedelta

import pytest

from app.extensions import db
from app.models import Dienst, Medewerker, SyncTaak

MAANDAG = date(2026, 3, 2)  # week 10 van 2026


@pytest.fixture
def mw(klaar):
    from app.services.voorbeeldpakket import laad_voorbeeldpakket

    laad_voorbeeldpakket()
    medewerker = Medewerker(naam="Medewerker A", initialen="MA", volgorde=1)
    db.session.add(medewerker)
    db.session.commit()
    return medewerker


def _import_dienst(naam, dag, **extra):
    from app.services.excel_import import ImportDienst

    waarden = dict(naam=naam, datum=dag, code=None, dienstnaam="Cursus", begin="08:00",
                   eind="16:00", opmerking="", opm_begin=None, opm_eind=None, excel_uren=None)
    waarden.update(extra)
    return ImportDienst(**waarden)


def _wachtrij_nu():
    from app.services import sync

    SyncTaak.query.update({"niet_voor": datetime(2000, 1, 1)})
    db.session.commit()
    return sync.verwerk_wachtrij()


# ---------------------------------------------------------------------------
# H1 · De Excel-import wist diensten in weken die niet in het bestand staan
# ---------------------------------------------------------------------------

def test_h1_import_laat_weken_zonder_blad_ongemoeid(app, mw):
    from app.services.excel_import import ImportMedewerker, ImportPlan, importeer

    week12 = MAANDAG + timedelta(weeks=2)
    db.session.add(Dienst(medewerker_id=mw.id, datum=week12, dienstnaam_override="Blijft",
                          begin="07:00", eind="15:00", opmerking_tekst=""))
    db.session.commit()
    plan = ImportPlan(jaar=2026, weken=[10, 14],
                      medewerkers=[ImportMedewerker("Medewerker A", "MA", None)],
                      diensten=[_import_dienst("Medewerker A", MAANDAG),
                                _import_dienst("Medewerker A", MAANDAG + timedelta(weeks=4))])
    importeer(plan)
    assert Dienst.query.filter_by(datum=week12).one().dienstnaam_override == "Blijft"
    assert Dienst.query.count() == 3


def test_h1_droogloop_toont_te_verwijderen_bestaande_diensten(app, mw):
    from app.services.excel_import import ImportMedewerker, ImportPlan, effect

    for dag in (MAANDAG, MAANDAG + timedelta(days=1), MAANDAG + timedelta(weeks=2)):
        db.session.add(Dienst(medewerker_id=mw.id, datum=dag, dienstnaam_override="Oud",
                              begin="07:00", eind="15:00", opmerking_tekst=""))
    db.session.commit()
    plan = ImportPlan(jaar=2026, weken=[10],
                      medewerkers=[ImportMedewerker("Medewerker A", "MA", None)],
                      diensten=[_import_dienst("Medewerker A", MAANDAG)])
    per_mw = effect(plan).per_medewerker["Medewerker A"]
    # Maandag wordt vervangen, dinsdag verwijderd; week 12 (geen blad) telt niet mee
    assert (per_mw.nieuw, per_mw.vervangen, per_mw.verwijderd) == (0, 1, 1)


def test_h1_droogloopscherm_toont_verwijderde_diensten_per_medewerker(app, als_beheerder, tmp_path):
    import io

    from .test_import_backup import maak_testbestand

    medewerker = Medewerker(naam="Medewerker Vijf B", initialen="MVB")
    db.session.add(medewerker)
    db.session.commit()
    for dag in (MAANDAG + timedelta(days=1), MAANDAG + timedelta(days=2)):  # niet in het bestand
        db.session.add(Dienst(medewerker_id=medewerker.id, datum=dag, dienstnaam_override="Oud",
                              opmerking_tekst=""))
    db.session.commit()
    pad = str(tmp_path / "oud.xlsx")
    maak_testbestand(pad)
    with open(pad, "rb") as f:
        als_beheerder.post("/beheer/importeren", data={"bestand": (io.BytesIO(f.read()), "x.xlsx")},
                           content_type="multipart/form-data")
    pagina = als_beheerder.get("/beheer/importeren/voorbeeld").data.decode()
    rij = pagina.split("data-effect")[1].split("Medewerker Vijf B")[1].split("</tr>")[0]
    assert "<strong>2</strong>" in rij


# ---------------------------------------------------------------------------
# H2 · Dubbele Google-afspraken na een gedeeltelijk mislukte dag-sync
# ---------------------------------------------------------------------------

@pytest.fixture
def gekoppeld(mw):
    from app.services import instellingen

    mw.agenda_modus, mw.agenda_id = "B", "agenda-a"
    instellingen.schrijf("agenda_sync_dagen_terug", "3650")  # vaste datums: zie test_agenda.gekoppeld
    db.session.commit()
    return mw


@pytest.fixture
def nep(app, monkeypatch):
    from app.services import google_agenda

    from .test_agenda import NepKlant

    klant = NepKlant()
    klant.agendas["agenda-a"] = {}
    monkeypatch.setattr(google_agenda, "klant", lambda: klant)
    return klant


def test_h2_tijdelijke_fout_bij_tweede_afspraak_geeft_geen_dubbele(app, gekoppeld, nep):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "17/3")])  # twee diensten
    nep.faal_bij = {"maak_afspraak": {2}}  # de tweede insert geeft een tijdelijke 503
    _wachtrij_nu()
    assert SyncTaak.query.one().pogingen == 1  # opnieuw proberen
    _wachtrij_nu()
    assert SyncTaak.query.count() == 0
    assert sorted(a["summary"] for a in nep.agendas["agenda-a"].values()) == ["BHV", "VW Avond"]
    ids = {d.google_event_id for d in Dienst.query.all()}
    assert len(ids) == 2 and ids == set(nep.agendas["agenda-a"])


def test_h2_mislukte_commit_na_insert_geeft_geen_dubbele(app, gekoppeld, nep, monkeypatch):
    """Google maakte de afspraak wel, maar het antwoord (of de commit) ging verloren."""
    from app.services import sync
    from app.services.google_agenda import AgendaFout
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    echt = nep.maak_afspraak

    def antwoord_kwijt(agenda_id, body):
        echt(agenda_id, body)
        raise AgendaFout("time-out", tijdelijk=True)

    monkeypatch.setattr(nep, "maak_afspraak", antwoord_kwijt)
    _wachtrij_nu()
    monkeypatch.setattr(nep, "maak_afspraak", echt)
    _wachtrij_nu()
    assert len(nep.agendas["agenda-a"]) == 1
    assert Dienst.query.one().google_event_id in nep.agendas["agenda-a"]
    assert sync.event_id_voor(Dienst.query.one()) == Dienst.query.one().google_event_id


def test_h2_event_id_is_geldig_voor_google(app, mw):
    import re

    from app.services import sync

    dienst = Dienst(id=12, medewerker_id=mw.id, datum=MAANDAG, volgnummer=2)
    event_id = sync.event_id_voor(dienst)
    assert re.fullmatch(r"[a-v0-9]{5,1024}", event_id)
    assert event_id == sync.event_id_voor(dienst)
    assert event_id != sync.event_id_voor(Dienst(medewerker_id=mw.id, datum=MAANDAG, volgnummer=1))


def test_h2_afspraak_opnieuw_na_verwijderen_werkt(app, gekoppeld, nep):
    """Na wissen en opnieuw invullen bestaat het ID al bij Google (geannuleerd): bijwerken."""
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    _wachtrij_nu()
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "")])
    _wachtrij_nu()
    assert nep.agendas["agenda-a"] == {} and Dienst.query.count() == 0
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "5")])
    _wachtrij_nu()
    assert SyncTaak.query.count() == 0
    assert [a["summary"] for a in nep.agendas["agenda-a"].values()] == ["VW Dag"]


# ---------------------------------------------------------------------------
# M1 · De worker verwijdert een dienst die de planner net opnieuw invulde
# ---------------------------------------------------------------------------

def _in_ander_proces(medewerker_id, datum, wijzig):
    """Wijzig een dienst via een eigen verbinding, zoals een webverzoek in een ander proces."""
    from sqlalchemy.orm import Session

    with Session(db.engine) as sessie:
        dienst = sessie.query(Dienst).filter_by(medewerker_id=medewerker_id, datum=datum).one()
        wijzig(dienst, sessie)
        dienst.versie += 1
        sessie.commit()


def test_m1_worker_verwijdert_opnieuw_ingevulde_dienst_niet(app, gekoppeld, nep, monkeypatch):
    from app.models import Dienstcode
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    _wachtrij_nu()
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "")])  # leeg, afspraak moet weg
    code5 = Dienstcode.query.filter_by(nummer=5).one().id
    echt = nep.verwijder_afspraak

    def tussendoor(agenda_id, event_id):
        echt(agenda_id, event_id)

        def vul_in(dienst, sessie):  # de planner typt intussen code 5
            dienst.dienstcode_id, dienst.begin, dienst.eind = code5, "07:00", "15:30"

        _in_ander_proces(gekoppeld.id, MAANDAG, vul_in)

    monkeypatch.setattr(nep, "verwijder_afspraak", tussendoor)
    _wachtrij_nu()
    db.session.expire_all()
    dienst = Dienst.query.one()  # niet weggegooid
    assert dienst.dienstcode_id == code5
    monkeypatch.setattr(nep, "verwijder_afspraak", echt)
    from app.services import sync_planning

    sync_planning.plan_dag(dienst.medewerker, MAANDAG)
    _wachtrij_nu()
    assert [a["summary"] for a in nep.agendas["agenda-a"].values()] == ["VW Dag"]


def test_m1_worker_overschrijft_nieuwer_event_id_niet(app, gekoppeld, nep, monkeypatch):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4")])
    echt = nep.maak_afspraak

    def tussendoor(agenda_id, body):
        event_id = echt(agenda_id, body)

        def ontkoppel(dienst, sessie):  # bijv. de planner past de dienst aan
            dienst.eind = "18:00"

        _in_ander_proces(gekoppeld.id, MAANDAG, ontkoppel)
        return event_id

    monkeypatch.setattr(nep, "maak_afspraak", tussendoor)
    _wachtrij_nu()
    db.session.expire_all()
    dienst = Dienst.query.one()
    assert dienst.eind == "18:00" and dienst.versie == 2


# ---------------------------------------------------------------------------
# M2 · Een vrije dienstnaam van dienst 2 verdwijnt via het coderaster
# ---------------------------------------------------------------------------

def _dag(mw_id):
    from app.services.weekrooster import diensten_van_dag

    db.session.expire_all()
    return diensten_van_dag(mw_id, MAANDAG)


def test_m2_vrije_tweede_dienst_blijft_bij_andere_eerste_code(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4/3")])
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "dienstnaam", "Cursus", volgnummer=2)])
    d1, d2 = _dag(mw.id)
    from app.services.weekrooster import matrix_code

    assert matrix_code(d1, d2) == "4/"  # het raster toont de vrije naam niet als code
    bijgewerkt, fouten = wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "5/")])
    assert fouten == []
    d1, d2 = _dag(mw.id)
    assert d1.dienstcode.nummer == 5 and d2 is not None and d2.dienstnaam == "Cursus"
    assert bijgewerkt[f"{mw.id}|{MAANDAG.isoformat()}"]["code"] == "5/"


def test_m2_zonder_tweede_deel_verdwijnt_dienst2_wel(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4/3")])
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "dienstnaam", "Cursus", volgnummer=2)])
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "5")])
    d1, d2 = _dag(mw.id)
    assert d1.dienstcode.nummer == 5 and d2 is None


def test_m2_lege_tweede_code_wist_een_tweede_dienst_met_code(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4/3")])
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4/")])
    d1, d2 = _dag(mw.id)
    assert d1.dienstcode.nummer == 4 and d2 is None


# ---------------------------------------------------------------------------
# M3 · De worker draait als PID 1 zonder SIGTERM-handler
# ---------------------------------------------------------------------------

def test_m3_worker_stopt_netjes_na_sigterm(app, klaar, monkeypatch):
    import os
    import signal

    from app import worker

    rondes = []

    def ronde(planning, nu=None):
        rondes.append(1)
        os.kill(os.getpid(), signal.SIGTERM)  # 'docker stop' midden in een ronde

    monkeypatch.setattr(worker, "create_app", lambda: app)
    monkeypatch.setattr(worker, "een_ronde", ronde)
    oud = {s: signal.getsignal(s) for s in (signal.SIGTERM, signal.SIGINT)}
    try:
        worker.main()  # moet na de lopende ronde zelf stoppen (geen oneindige lus)
    finally:
        for sig, handler in oud.items():
            signal.signal(sig, handler)
    assert rondes == [1]


@pytest.mark.parametrize("bestand", ["docker-compose.yml", "README.md"])
def test_m3_init_in_docker_compose(bestand):
    """Een kleine init (tini) als PID 1 geeft signalen door en ruimt zombieprocessen op."""
    import os
    import re

    with open(os.path.join(os.path.dirname(__file__), "..", bestand), encoding="utf-8") as f:
        tekst = f.read()
    for service in ("web", "worker"):
        blok = re.search(rf"^  {service}:\n((?:    .*\n|\n)+)", tekst, re.M)
        assert blok and re.search(r"^    init: true", blok.group(1), re.M), service


# ---------------------------------------------------------------------------
# M4 · Na het terugzetten van een back-up worden de agenda's niet bijgewerkt
# ---------------------------------------------------------------------------

def _gekoppelde_backup():
    from app.services import backup, instellingen

    instellingen.schrijf("setup_voltooid", "1")
    db.session.add_all([Medewerker(naam="Medewerker A", initialen="MA", agenda_modus="B",
                                   agenda_id="agenda-a"),
                        Medewerker(naam="Medewerker B", initialen="MB")])
    db.session.commit()
    return backup.maak_backup("handmatig")


def test_m4_terugzetten_via_scherm_plant_agenda_sync(gemigreerd, client):
    import os

    from .conftest import login, maak_gebruiker

    pad = _gekoppelde_backup()
    maak_gebruiker("beheerder", "beheerder")
    login(client, "beheerder")
    antwoord = client.post("/beheer/backups/terugzetten",
                           data={"naam": os.path.basename(pad), "bevestig": "1"}, follow_redirects=True)
    db.session.remove()
    taken = SyncTaak.query.all()
    assert [(t.soort, t.medewerker_id) for t in taken] == \
        [("volledig", Medewerker.query.filter_by(initialen="MA").one().id)]
    assert "Google Agenda" in antwoord.data.decode()


def test_m4_terugzetten_via_cli_plant_agenda_sync(gemigreerd):
    import os

    pad = _gekoppelde_backup()
    resultaat = gemigreerd.test_cli_runner().invoke(args=["terugzetten", os.path.basename(pad), "--yes"])
    assert resultaat.exit_code == 0, resultaat.output
    db.session.remove()
    assert [t.soort for t in SyncTaak.query.all()] == ["volledig"]
    assert "Google Agenda" in resultaat.output


# ---------------------------------------------------------------------------
# M5 · Inlogvloed: per nieuwe naam elke poging argon2 en databaserijen
# ---------------------------------------------------------------------------

def _mislukte_pogingen_vanaf_ip(aantal, ip="127.0.0.1"):
    from app.models import LoginPoging
    from app.services import klok

    db.session.add_all([LoginPoging(gebruikersnaam=f"onbekend{i}", ip=ip, gelukt=False,
                                    tijdstip=klok.utc_nu()) for i in range(aantal)])
    db.session.commit()


def _tel_rijen():
    from app.models import Logboek, LoginPoging

    return LoginPoging.query.count(), Logboek.query.count()


def test_m5_harde_grens_per_ip_zonder_hash_en_zonder_rijen(app, client, klaar, monkeypatch):
    from app.blueprints import auth

    hashes = []
    monkeypatch.setattr(auth, "controleer_dummy", lambda w: hashes.append(w) or False)
    monkeypatch.setattr(auth, "controleer_wachtwoord", lambda h, w: hashes.append(w) or False)
    _mislukte_pogingen_vanaf_ip(auth.HARDE_GRENS_PER_IP)
    client.post("/login", data={"gebruikersnaam": "nieuw1", "wachtwoord": "x"})  # eerste: één logregel
    voor = _tel_rijen()
    for i in range(5):
        antwoord = client.post("/login", data={"gebruikersnaam": f"nieuwer{i}", "wachtwoord": "x"})
        assert antwoord.status_code == 429
    assert hashes == []  # geen argon2 meer
    assert _tel_rijen() == voor  # geen nieuwe pogingen of logboekregels


def test_m5_tijdens_blokkade_geen_login_mislukt_regels(app, client, klaar):
    from app.blueprints import auth
    from app.models import Logboek

    _mislukte_pogingen_vanaf_ip(auth.MAX_POGINGEN_PER_IP)
    for i in range(3):
        client.post("/login", data={"gebruikersnaam": f"naam{i}", "wachtwoord": "fout"})
    assert Logboek.query.filter_by(actie="Login mislukt").count() == 0


def test_m5_een_kans_per_naam_blijft_tot_de_harde_grens(app, client, klaar):
    from app.blueprints import auth

    from .conftest import login

    _mislukte_pogingen_vanaf_ip(auth.HARDE_GRENS_PER_IP - 1)
    assert login(client, "collega").status_code == 302


# ---------------------------------------------------------------------------
# M6 · Privacy van de deellink: geen contracturen en weektotalen
# ---------------------------------------------------------------------------

@pytest.fixture
def deellink(mw):
    from app.models import Contracturen
    from app.services import instellingen
    from app.services.weekrooster import Wijziging, wijzig_cellen

    db.session.add(Contracturen(medewerker_id=mw.id, jaar=2026, uren=1234))
    instellingen.schrijf("deellink_actief", "1")
    instellingen.schrijf("deellink_token", "geheim-token-voor-de-test")
    db.session.commit()
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4"), Wijziging(mw.id, MAANDAG + timedelta(days=1),
                                                                        "code", "4")])
    return "geheim-token-voor-de-test"


def test_m6_deellink_toont_geen_contracturen_en_weektotalen(app, client, deellink):
    pagina = client.get(f"/deel/{deellink}/week/2026/10").data.decode()
    assert "VW Vroeg" in pagina  # het rooster zelf wel
    assert "1234" not in pagina and "1.234" not in pagina and "data-pcontract" not in pagina
    assert "16,00" not in pagina  # weektotaal (2 x 8 uur)
    assert "Weektotaal" not in pagina and "contracturen" not in pagina.lower()


def test_m6_ingelogd_ziet_contracturen_en_weektotalen_wel(app, als_gebruiker, deellink):
    pagina = als_gebruiker.get("/week/2026/10").data.decode()
    assert 'data-pcontract="1234"' in pagina and "16,00" in pagina


def test_m6_docstring_deellink_klopt():
    from app.blueprints import deel

    assert "contracturen" in deel.__doc__ and "weektotalen" in deel.__doc__


# ---------------------------------------------------------------------------
# L1 · Setup, maak-beheerder en de import valideren invoer niet
# ---------------------------------------------------------------------------

def _setup_tot_stap(client, stap):
    from app.services import setup_code

    from .conftest import WACHTWOORD

    client.get("/setup/")
    client.post("/setup/", data={"code": setup_code.lees_code()})
    if stap > 1:
        client.post("/setup/stap/1", data={"gebruikersnaam": "planner", "weergavenaam": "Planner",
                                           "wachtwoord": WACHTWOORD, "herhaling": WACHTWOORD})


@pytest.mark.parametrize("naam", ["met spatie", "x", "a" * 65, "<script>", "ü-teken"])
def test_l1_setup_weigert_ongeldige_gebruikersnaam(app, client, naam):
    from app.models import Gebruiker

    from .conftest import WACHTWOORD

    _setup_tot_stap(client, 1)
    antwoord = client.post("/setup/stap/1", data={"gebruikersnaam": naam, "weergavenaam": "P",
                                                  "wachtwoord": WACHTWOORD, "herhaling": WACHTWOORD})
    assert antwoord.status_code == 400 and Gebruiker.query.count() == 0


def test_l1_setup_weigert_te_lange_weergavenaam(app, client):
    from app.models import Gebruiker

    from .conftest import WACHTWOORD

    _setup_tot_stap(client, 1)
    antwoord = client.post("/setup/stap/1", data={"gebruikersnaam": "planner", "weergavenaam": "P" * 121,
                                                  "wachtwoord": WACHTWOORD, "herhaling": WACHTWOORD})
    assert antwoord.status_code == 400 and Gebruiker.query.count() == 0


def test_l1_maak_beheerder_weigert_ongeldige_naam(app, klaar, monkeypatch):
    import getpass

    from app.models import Gebruiker

    monkeypatch.setattr(getpass, "getpass", lambda prompt="": "nieuwwachtwoord1")
    resultaat = app.test_cli_runner().invoke(
        args=["maak-beheerder", "--gebruikersnaam", "met spatie", "--weergavenaam", "X"])
    assert resultaat.exit_code != 0 and "Gebruikersnaam" in resultaat.output
    assert Gebruiker.query.filter_by(gebruikersnaam="met spatie").first() is None


def test_l1_setup_stap4_geeft_geldige_initialen_en_begrenst_naam(app, client):
    from app.services.validatie import INITIALEN_PATROON

    _setup_tot_stap(client, 4)
    client.post("/setup/stap/4", data={"medewerkers": "Ömer Øzdemir\nÉva Ångström\n" + "N" * 200})
    medewerkers = Medewerker.query.all()
    assert len(medewerkers) == 3
    assert all(INITIALEN_PATROON.fullmatch(m.initialen) for m in medewerkers)
    assert {m.initialen for m in medewerkers} >= {"OOZ", "EAN"}
    assert all(len(m.naam) <= 120 for m in medewerkers)


def test_l1_import_corrigeert_initialen_en_slaat_ongeldige_codes_over(app, klaar, tmp_path):
    import openpyxl

    from app.models import Dienstcode
    from app.services.excel_import import importeer, lees_bestand
    from app.services.validatie import INITIALEN_PATROON

    boek = openpyxl.Workbook()
    lijsten = boek.active
    lijsten.title = "Lijsten"
    lijsten.cell(2, 2, "J.J.")
    lijsten.cell(2, 3, "Jan Jansen")
    lijsten.cell(3, 2, "ABCDEFGHIJKLM")
    lijsten.cell(3, 3, "Piet " + "P" * 200)
    for rij, nummer in enumerate((0, -3, 7), start=2):
        lijsten.cell(rij, 6, nummer)
        lijsten.cell(rij, 7, f"Code {nummer} " + "x" * 80)
    boek.create_sheet("Kalender")["E2"] = 2026
    pad = str(tmp_path / "x.xlsx")
    boek.save(pad)
    plan = lees_bestand(pad)
    assert [c.nummer for c in plan.codes] == [7]
    assert any("0" in w and "dienstcode" in w.lower() for w in plan.waarschuwingen)
    assert all(INITIALEN_PATROON.fullmatch(m.initialen) for m in plan.medewerkers)
    importeer(plan)
    assert all(INITIALEN_PATROON.fullmatch(m.initialen) for m in Medewerker.query.all())
    assert all(len(m.naam) <= 120 for m in Medewerker.query.all())
    assert len(Dienstcode.query.one().omschrijving) <= 60


def test_l1_beheer_begrenst_lengtes(app, als_beheerder):
    from app.models import Dienstcode, Gebruiker

    antwoord = als_beheerder.post("/beheer/medewerkers/nieuw", data={"naam": "N" * 121, "initialen": "NN"})
    assert antwoord.status_code == 400 and Medewerker.query.count() == 0
    antwoord = als_beheerder.post("/beheer/dienstcodes/nieuw", data={
        "nummer": "77", "omschrijving": "o" * 61, "kleur_achtergrond": "#FFFFFF", "kleur_tekst": "#000000"})
    assert antwoord.status_code == 400 and Dienstcode.query.count() == 0
    antwoord = als_beheerder.post("/beheer/gebruikers/nieuw", data={
        "gebruikersnaam": "nieuw", "weergavenaam": "W" * 121, "rol": "gebruiker",
        "wachtwoord": "langwachtwoord1"})
    assert antwoord.status_code == 400 and Gebruiker.query.filter_by(gebruikersnaam="nieuw").first() is None


# ---------------------------------------------------------------------------
# L2 · Integer-overflow geeft HTTP 500
# ---------------------------------------------------------------------------

GROOT = "99999999999999999999"  # past niet in een SQLite INTEGER (64 bits)


@pytest.mark.parametrize("pad", [
    f"/beheer/medewerkers/{GROOT}", f"/beheer/gebruikers/{GROOT}", f"/beheer/dienstcodes/{GROOT}",
    f"/week/{GROOT}/1", f"/beheer/logboek?pagina={GROOT}", f"/zoeken/?code={GROOT}",
    f"/kalender/?jaar={GROOT}",
])
def test_l2_grote_getallen_in_get_geven_geen_500(app, als_beheerder, mw, pad):
    assert als_beheerder.get(pad).status_code in (200, 400, 404)


def test_l2_grote_getallen_in_formulieren(app, als_beheerder, mw):
    from app.models import Dienstcode, Gebruiker

    assert als_beheerder.get(f"/zoeken/export.csv?code={GROOT}").status_code == 400
    antwoord = als_beheerder.post("/week/2026/10/kopieer", data={"naar": "2026-W11", "medewerker_id": GROOT})
    assert antwoord.status_code in (302, 404)
    antwoord = als_beheerder.post("/beheer/dienstcodes/nieuw", data={"nummer": GROOT, "omschrijving": "X"})
    assert antwoord.status_code == 400 and Dienstcode.query.filter_by(omschrijving="X").first() is None
    antwoord = als_beheerder.post("/beheer/gebruikers/nieuw", data={
        "gebruikersnaam": "nieuw", "weergavenaam": "N", "rol": "gebruiker", "medewerker_id": GROOT,
        "wachtwoord": "langwachtwoord1"})
    assert antwoord.status_code == 400 and Gebruiker.query.filter_by(gebruikersnaam="nieuw").first() is None
    assert als_beheerder.post("/beheer/vakanties/opslaan", data={"naam": "V", "datum_van": "2026-07-01",
                                                         "datum_tot": "2026-07-02", "id": GROOT}
                              ).status_code in (302, 404)


def test_l2_blanco_code_te_groot(app, als_beheerder):
    from .test_regressie import _instellingen_formulier

    pagina = als_beheerder.post("/beheer/instellingen", data=_instellingen_formulier(blanco_code=GROOT),
                                follow_redirects=True)
    assert pagina.status_code in (200, 400) and "blanco-code" in pagina.data.decode()
    from app.services import instellingen

    assert instellingen.lees("blanco_code") != GROOT


@pytest.mark.parametrize("wijziging", [
    {"mw": int(GROOT), "datum": "2026-03-02", "veld": "code", "waarde": "4"},
    {"mw": 1, "datum": "2026-03-02", "veld": "code", "waarde": "4", "versie": int(GROOT)},
    {"mw": 1, "datum": "2026-03-02", "veld": "code", "waarde": "4", "versie2": int(GROOT)},
    {"mw": 1, "datum": "2026-03-02", "veld": "code", "waarde": "4", "versie": -1},
    {"mw": 1, "datum": "2026-03-02", "veld": "code", "waarde": "4", "volgnummer": int(GROOT)},
])
def test_l2_api_cellen_grote_getallen_geven_400(app, als_beheerder, mw, wijziging):
    antwoord = als_beheerder.post("/api/cellen", json={"wijzigingen": [wijziging]})
    assert antwoord.status_code == 400 and antwoord.is_json


def test_l2_api_cellen_grote_code_en_ook_tonen(app, als_beheerder, mw):
    antwoord = als_beheerder.post("/api/cellen", json={"wijzigingen": [
        {"mw": mw.id, "datum": "2026-03-02", "veld": "code", "waarde": GROOT}]})
    assert antwoord.status_code == 200 and antwoord.get_json()["fouten"]
    antwoord = als_beheerder.post("/api/cellen", json={"ook_tonen": [f"{GROOT}|2026-03-02"]})
    assert antwoord.status_code == 400


def test_l2_onverwachte_fout_op_api_geeft_json(app, als_beheerder, mw, monkeypatch):
    from app.blueprints import rooster

    def kapot(*args, **kwargs):
        raise RuntimeError("onverwacht")

    monkeypatch.setattr(rooster, "verwerk_rooster", kapot)
    app.config["PROPAGATE_EXCEPTIONS"] = False
    antwoord = als_beheerder.post("/api/cellen", json={"wijzigingen": []})
    assert antwoord.status_code == 500 and antwoord.is_json and "fout" in antwoord.get_json()


# ---------------------------------------------------------------------------
# L3 · Dienst 2 blijft bestaan terwijl dienst 1 wordt opgeruimd
# ---------------------------------------------------------------------------

def _volgnummers(mw_id):
    db.session.expire_all()
    return sorted(d.volgnummer for d in Dienst.query.filter_by(medewerker_id=mw_id, datum=MAANDAG))


def test_l3_dienst1_blijft_als_dienst2_er_nog_is(app, mw):
    from app.services.weekrooster import Wijziging, matrix_code, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4/3")])
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "/3")])
    assert _volgnummers(mw.id) == [1, 2]  # lege dienst 1 blijft als plaatshouder
    d1, d2 = _dag(mw.id)
    assert d1.is_leeg and d2.dienstcode.nummer == 3 and matrix_code(d1, d2) == "/3"
    # Dienst 2 ook weg: dan gaan ze allebei
    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "")])
    assert _volgnummers(mw.id) == []


def test_l3_alleen_tweede_dienst_op_lege_dag_krijgt_dienst1(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "/3")])
    assert _volgnummers(mw.id) == [1, 2]


def test_l3_worker_ruimt_dienst1_niet_op_naast_dienst2(app, gekoppeld, nep):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "4/3")])
    _wachtrij_nu()
    wijzig_cellen([Wijziging(gekoppeld.id, MAANDAG, "code", "/3")])  # dienst 1 leeg, met afspraak
    _wachtrij_nu()
    assert _volgnummers(gekoppeld.id) == [1, 2]
    assert [a["summary"] for a in nep.agendas["agenda-a"].values()] == ["VW Avond"]


def test_l3_api_toont_lege_plaatshouder_als_null(app, als_beheerder, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "/3")])
    gegevens = als_beheerder.get("/api/v1/week/2026/10").get_json()
    rij = gegevens["medewerkers"][0]
    assert rij["dagen"][0] is None and rij["tweede_diensten"][0]["code"] == 3


# ---------------------------------------------------------------------------
# L4 · Archiveren met een ongeldige datum gebruikt stil 'vandaag'
# ---------------------------------------------------------------------------

def test_l4_archiveren_met_ongeldige_datum_geeft_fout(app, als_beheerder, mw):
    antwoord = als_beheerder.post(f"/beheer/medewerkers/{mw.id}/archiveer", data={"vanaf": "31-02-2026"},
                                  follow_redirects=True)
    assert "ongeldige datum" in antwoord.data.decode().lower()
    db.session.expire_all()
    assert db.session.get(Medewerker, mw.id).gearchiveerd_vanaf is None


def test_l4_archiveren_met_geldige_datum(app, als_beheerder, mw):
    als_beheerder.post(f"/beheer/medewerkers/{mw.id}/archiveer", data={"vanaf": "2026-12-01"})
    db.session.expire_all()
    assert db.session.get(Medewerker, mw.id).gearchiveerd_vanaf == date(2026, 12, 1)


# ---------------------------------------------------------------------------
# L6 · sync_dag buiten de sync-periode
# ---------------------------------------------------------------------------

def test_l6_sync_dag_buiten_periode_maakt_geen_afspraken(app, gekoppeld, nep):
    from app.services import sync_planning
    from app.services.weekrooster import Wijziging, wijzig_cellen

    ver = date.today() + timedelta(days=800)  # ver na de sync-periode (standaard 12 maanden)
    wijzig_cellen([Wijziging(gekoppeld.id, ver, "code", "4")])
    _wachtrij_nu()
    assert nep.agendas["agenda-a"] == {}  # net als sync_volledig: niets buiten de periode
    # Maar een gewiste dienst met een (oude) afspraak wordt daar wel opgeruimd
    dienst = Dienst.query.one()
    nep.agendas["agenda-a"]["oud"] = {"id": "oud"}
    dienst.google_event_id = "oud"
    db.session.commit()
    wijzig_cellen([Wijziging(gekoppeld.id, ver, "code", "")])
    sync_planning.plan_dag(dienst.medewerker, ver)
    _wachtrij_nu()
    assert nep.agendas["agenda-a"] == {} and Dienst.query.count() == 0


# ---------------------------------------------------------------------------
# L7 · Begin = eind: uren, Google-afspraak en overlappen niet consistent
# ---------------------------------------------------------------------------

def test_l7_begin_gelijk_aan_eind_wordt_geweigerd(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4")])  # 07:15-15:45
    _, fouten = wijzig_cellen([Wijziging(mw.id, MAANDAG, "eind", "07:15")])
    assert fouten and "gelijk" in fouten[0]["melding"]
    db.session.expire_all()
    assert Dienst.query.one().eind == "15:45"


def test_l7_dienstcode_met_gelijke_standaardtijden_geweigerd(app, als_beheerder):
    from app.models import Dienstcode

    antwoord = als_beheerder.post("/beheer/dienstcodes/nieuw", data={
        "nummer": "88", "omschrijving": "Raar", "std_begin": "08:00", "std_eind": "08:00"})
    assert antwoord.status_code == 400 and Dienstcode.query.filter_by(nummer=88).first() is None


def test_l7_een_betekenis_overal(app):
    """Begin = eind (bijv. uit een oude import) betekent overal: duur 0."""
    from types import SimpleNamespace

    from app.services.google_agenda import afspraak_voor
    from app.services.urenberekening import bereken_uren
    from app.services.weekrooster import overlappen

    def dienst(begin, eind):
        return SimpleNamespace(begin=begin, eind=eind)

    assert bereken_uren("08:00", "08:00") == 0
    assert not overlappen(dienst("08:00", "08:00"), dienst("10:00", "12:00"))
    nul = SimpleNamespace(begin="08:00", eind="08:00", datum=MAANDAG, dienstcode=None, dienstnaam="X",
                          opmerking_tekst="", id=1, medewerker_id=1)
    body = afspraak_voor(nul).body
    assert body["start"]["dateTime"] == body["end"]["dateTime"]


# ---------------------------------------------------------------------------
# L8 · Telefoonkaart toont opmerkingtijden alleen met een begintijd
# ---------------------------------------------------------------------------

def test_l8_kaart_toont_opmerkingtijd_met_alleen_eind(app):
    from app.services.weekrooster import dienst_naar_dict
    from app.services.weekweergave import kaart

    d = dienst_naar_dict(None)
    d.update(opmerking="Training", opm_begin="", opm_eind="17:00", tweede=None)
    assert "17:00" in str(kaart(1, "2026-03-02", d, "ma", False))


def test_l8_raster_js_zelfde_regel():
    import os

    pad = os.path.join(os.path.dirname(__file__), "..", "app", "static", "js", "raster.js")
    with open(pad, encoding="utf-8") as f:
        js = f.read()
    assert '(g.opm_begin ? " ("' not in js  # alleen begin telde: ook eind moet meetellen


# ---------------------------------------------------------------------------
# L9 · Zoeken en de CSV-export knippen stil af op 5000 resultaten
# ---------------------------------------------------------------------------

def test_l9_afkappen_wordt_gemeld(app, als_gebruiker, mw, monkeypatch):
    from app.services import overzichten
    from app.services.weekrooster import Wijziging, wijzig_cellen

    monkeypatch.setattr(overzichten, "MAX_RESULTATEN", 2)
    wijzig_cellen([Wijziging(mw.id, MAANDAG + timedelta(days=i), "code", "4") for i in range(3)])
    pagina = als_gebruiker.get("/zoeken/?code=4").data.decode()
    assert "alleen de eerste 2" in pagina
    csv = als_gebruiker.get("/zoeken/export.csv?code=4").data.decode()
    assert "alleen de eerste 2" in csv
    assert overzichten.zoek_diensten(code=4, limiet=5) and len(overzichten.zoek_diensten(code=4)) == 2


def test_l9_geen_melding_onder_de_grens(app, als_gebruiker, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4")])
    assert "alleen de eerste" not in als_gebruiker.get("/zoeken/?code=4").data.decode()


# ---------------------------------------------------------------------------
# L10 · ook_tonen/ook_dagen in /api/cellen onbegrensd
# ---------------------------------------------------------------------------

def test_l10_ook_tonen_en_ook_dagen_begrensd(app, als_beheerder, mw):
    veel = [f"{mw.id}|2026-03-02"] * 6000
    assert als_beheerder.post("/api/cellen", json={"ook_tonen": veel}).status_code == 400
    assert als_beheerder.post("/api/cellen", json={"ook_dagen": ["2026-03-02"] * 6000}).status_code == 400
    assert als_beheerder.post("/api/cellen", json={"ook_tonen": "geen lijst"}).status_code == 400
    assert als_beheerder.post("/api/cellen", json={"ook_tonen": [f"{mw.id}|2026-03-02"]}).status_code == 200


# ---------------------------------------------------------------------------
# L11 · Laatste-beheerdercontrole niet atomair
# ---------------------------------------------------------------------------

def test_l11_laatste_beheerder_ook_bij_gelijktijdige_wijziging(app, als_beheerder, monkeypatch):
    from sqlalchemy.orm import Session

    from app.blueprints.beheer import gebruikers
    from app.models import Gebruiker

    from .conftest import maak_gebruiker

    tweede = maak_gebruiker("tweede", "beheerder")
    eerste = Gebruiker.query.filter_by(gebruikersnaam="beheerder").one()
    echt = gebruikers._aantal_actieve_beheerders

    def tussendoor(behalve_id=None):
        aantal = echt(behalve_id)
        with Session(db.engine) as sessie:  # een andere beheerder degradeert 'tweede' tegelijk
            sessie.query(Gebruiker).filter_by(id=tweede.id).update({"rol": "gebruiker"})
            sessie.commit()
        return aantal

    monkeypatch.setattr(gebruikers, "_aantal_actieve_beheerders", tussendoor)
    antwoord = als_beheerder.post(f"/beheer/gebruikers/{eerste.id}", data={
        "gebruikersnaam": "beheerder", "weergavenaam": "B", "rol": "gebruiker", "actief": "1"})
    assert antwoord.status_code == 400
    db.session.expire_all()
    assert Gebruiker.query.filter_by(rol="beheerder", actief=True).count() >= 1


# ---------------------------------------------------------------------------
# L12 · Nachtelijke back-up alleen in het geheugen onthouden
# ---------------------------------------------------------------------------

def test_l12_na_herstart_geen_tweede_nachtelijke_backup(app, klaar, monkeypatch):
    from app import worker
    from app.services import backup, klok

    from .test_regressie import Klok

    nu = klok.nu().replace(hour=3)
    monkeypatch.setattr(klok, "nu", Klok(nu))  # elke back-up een eigen naam (tijd loopt door)
    worker.een_ronde(worker.Planning(), nu)
    assert len(backup.lijst_backups()) == 1
    worker.een_ronde(worker.Planning(), nu)  # worker herstart: nieuw (leeg) geheugen
    assert len(backup.lijst_backups()) == 1


# ---------------------------------------------------------------------------
# L13 · Dienstcode hernoemen laat aanvullingen achter de oude naam staan
# ---------------------------------------------------------------------------

def test_l13_hernoemen_werkt_aanvulling_bij(app, als_beheerder, mw):
    from app.models import Dienstcode
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4"),
                   Wijziging(mw.id, MAANDAG, "dienstnaam", "VW Vroeg tot 12:00")])
    code = Dienstcode.query.filter_by(nummer=4).one()
    als_beheerder.post(f"/beheer/dienstcodes/{code.id}", data={
        "nummer": "4", "omschrijving": "VW Ochtend", "std_begin": "07:15", "std_eind": "15:45",
        "vet": "1", "actief": "1", "in_agenda": "1"})
    db.session.expire_all()
    dienst = Dienst.query.one()
    assert dienst.dienstnaam == "VW Ochtend tot 12:00" and dienst.dienstcode.nummer == 4


def test_l13_alleen_hoofdletters_anders_is_geen_aanvulling(app, mw):
    from app.services.weekrooster import Wijziging, wijzig_cellen

    wijzig_cellen([Wijziging(mw.id, MAANDAG, "code", "4"),
                   Wijziging(mw.id, MAANDAG, "dienstnaam", "vw vroeg")])
    db.session.expire_all()
    dienst = Dienst.query.one()
    assert dienst.dienstnaam_override == "" and dienst.dienstcode.nummer == 4


# ---------------------------------------------------------------------------
# L14 · Waarschuwing 'Alle uren herberekenen' bij een eigen roostervrije dag
# ---------------------------------------------------------------------------

def test_l14_eigen_roostervrije_dag_geeft_herbereken_melding(app, als_beheerder):
    from app.models import Feestdag
    from app.services import instellingen

    instellingen.schrijf("toeslag_feestdag", "2")
    db.session.commit()
    antwoord = als_beheerder.post("/beheer/feestdagen/nieuw", data={"naam": "Teamdag", "datum": "2026-06-10"},
                                  follow_redirects=True)
    assert "Alle uren herberekenen" in antwoord.data.decode()
    dag = Feestdag.query.filter_by(naam="Teamdag").one()
    antwoord = als_beheerder.post(f"/beheer/feestdagen/{dag.id}/verwijder", follow_redirects=True)
    assert "Alle uren herberekenen" in antwoord.data.decode()


def test_l14_zonder_feestdagtoeslag_geen_melding(app, als_beheerder):
    antwoord = als_beheerder.post("/beheer/feestdagen/nieuw", data={"naam": "Teamdag", "datum": "2026-06-10"},
                                  follow_redirects=True)
    assert "Alle uren herberekenen" not in antwoord.data.decode()


# ---------------------------------------------------------------------------
# L15 · De tijdzonecache geldt per proces (60 s)
# ---------------------------------------------------------------------------

def test_l15_andere_processen_zien_een_nieuwe_tijdzone_snel(app, klaar, monkeypatch):
    from sqlalchemy.orm import Session

    from app.models import Instelling
    from app.services import klok

    tijd = [1000.0]
    monkeypatch.setattr(klok.time, "monotonic", lambda: tijd[0])
    klok.wis_cache()
    assert klok.tijdzone_naam() == "Europe/Amsterdam"
    with Session(db.engine) as sessie:  # een ander proces (Beheer) wijzigt de tijdzone
        sessie.merge(Instelling(sleutel="tijdzone", waarde="Europe/London"))
        sessie.commit()
    tijd[0] += 6
    assert klok.tijdzone_naam() == "Europe/London"
