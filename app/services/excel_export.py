"""Het rooster exporteren naar MS Excel (.xlsx), in de opbouw van het oude Excel-rooster.

Zo kan het bestand weer ingelezen worden met Beheer → Importeren (zie de docstring van
excel_import.py). Bladen:
- W1..W53:      één blad per ISO-week, zoals het weekrooster: per dag drie kolommen
                (D, G, J, M, P, S, V), per medewerker vier regels (opmerking, opmerkingtijden,
                dienstnaam, begin/eind/uren), contracturen en weektotaal in Z, het code-raster
                in AB..AI. De tweede dienst van een dag staat rechts, vanaf AK ('2e dienst').
- Lijsten:      medewerkers met contracturen, dienstcodes, toeslagen (N2/N3).
- Urenoverzicht zoals /overzicht/uren.
- Vakanties en Kalender (E2 = het jaar).

Wat NIET terugkomt bij een import: kleuren van dienstcodes en kleurregels (de import gebruikt
de kleuren uit het voorbeeldpakket of grijs), e-mailadressen, archiefdatums en agenda-
koppelingen, feestdagen en uitgezette feestdagen, en handmatig gewiste automatische
dagopmerkingen. Een periode moet binnen één ISO-jaar vallen: dan blijft de export
herimporteerbaar (één blad per weeknummer).

Tekst die met = + - of @ begint, wordt als tekst weggeschreven (nooit als formule).
"""

import io
import logging
from copy import copy
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, GradientFill, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload

from ..models import Dienst, Dienstcode, Medewerker, Vakantie
from . import instellingen
from .excel_import import CODE_KOLOMMEN, DAG_KOLOMMEN, TWEEDE_KOLOMMEN
from .kalender import DAGNAMEN_KORT, MAX_JAAR, MIN_JAAR, aantal_weken, eerste_en_laatste_dag_isojaar
from .overzichten import uren_overzicht
from .weekrooster import dagopmerkingen, kleurregels, matrix_code, medewerkers_voor_periode

log = logging.getLogger(__name__)
FORMULE_TEKENS = ("=", "+", "-", "@")
DUN = Side(style="thin", color="BFBFBF")
RAND = Border(left=DUN, right=DUN, top=DUN, bottom=DUN)
VET = Font(bold=True)
KOP = PatternFill("solid", fgColor="D9E1F2")
WEEKEND = PatternFill("solid", fgColor="F2F2F2")
MIDDEN = Alignment(horizontal="center", vertical="center")
# Gecentreerd over de lege cellen ernaast, zonder samenvoegen (dat is in openpyxl erg traag)
OVER_DRIE = Alignment(horizontal="centerContinuous", vertical="center")
TIJD = {"number_format": "hh:mm", "border": RAND}
UREN = {"number_format": "0.00", "border": RAND}


class ExportFout(ValueError):
    """Ongeldige keuze voor de export (periode, jaar)."""


def periode_voor(jaar: int | None, van: date | None, tot: date | None) -> tuple[int, date, date]:
    """(ISO-jaar, van, tot) van de export. Een periode moet binnen één ISO-jaar vallen."""
    if van is None and tot is None:
        if jaar is None or not MIN_JAAR <= jaar <= MAX_JAAR:
            raise ExportFout(f"Kies een jaar tussen {MIN_JAAR} en {MAX_JAAR}.")
        eerste, laatste = eerste_en_laatste_dag_isojaar(jaar)
        return jaar, eerste, laatste
    if van is None or tot is None or tot < van:
        raise ExportFout("Kies een geldige periode (van en tot, tot niet vóór van).")
    iso_jaar = van.isocalendar()[0]
    if tot.isocalendar()[0] != iso_jaar:
        raise ExportFout("De periode moet binnen één jaar (ISO-weken) vallen; exporteer per jaar.")
    if not MIN_JAAR <= iso_jaar <= MAX_JAAR:
        raise ExportFout(f"Kies een periode tussen {MIN_JAAR} en {MAX_JAAR}.")
    return iso_jaar, van, tot


def bestandsnaam(jaar: int, van: date, tot: date, medewerker: Medewerker | None) -> str:
    eerste, laatste = eerste_en_laatste_dag_isojaar(jaar)
    naam = f"rooster-{jaar}" if (van, tot) == (eerste, laatste) else f"rooster-{van:%Y%m%d}-{tot:%Y%m%d}"
    if medewerker is not None:
        naam += f"-{medewerker.initialen.lower()}"
    return naam + ".xlsx"


# ---------------------------------------------------------------------------
# Hulpjes voor cellen
# ---------------------------------------------------------------------------

# Opmaak per combinatie één keer opbouwen en daarna kopiëren: elke losse toekenning
# (font, fill, ...) kost in openpyxl een opzoeking in de stijllijsten van het werkboek.
# De cache hoort bij het werkboek (de stijlnummers verwijzen naar diens lijsten).


def _zet_als(blad, rij: int, kolom: int, waarde, stijl: dict | None = None) -> None:
    """Als _zet, maar een lege waarde slaan we over (geen lege, opgemaakte cel: scheelt tijd)."""
    if waarde is not None and waarde != "":
        _zet(blad, rij, kolom, waarde, stijl)


def _zet(blad, rij: int, kolom: int, waarde, stijl: dict | None = None):
    """Schrijf een waarde; tekst die als formule gelezen kan worden blijft tekst."""
    cel = blad.cell(rij, kolom)
    cel.value = waarde
    formule = isinstance(waarde, str) and waarde.startswith(FORMULE_TEKENS)
    if formule:
        cel.data_type = "s"  # openpyxl maakt anders een formule van '=...'
    if stijl or formule:
        cache = blad.parent.__dict__.setdefault("_rooster_stijlen", {})
        sleutel = (formule,) + tuple((k, id(v)) for k, v in (stijl or {}).items())
        if sleutel in cache:
            cel._style = copy(cache[sleutel][0])
        else:
            for naam, stijlwaarde in (stijl or {}).items():
                setattr(cel, naam, stijlwaarde)
            if formule:
                cel.quotePrefix = True  # Excel toont het als tekst
            # De stijlobjecten zelf ook bewaren: zo blijft hun id() uniek zolang de cache bestaat
            cache[sleutel] = (copy(cel._style), tuple((stijl or {}).values()))
    return cel


def _tijd(tekst: str | None) -> time | None:
    return time.fromisoformat(tekst) if tekst else None


class _Stijlen:
    """Opvulling en lettertype per dienstcode en kleurregel, één keer gemaakt (snelheid)."""

    def __init__(self) -> None:
        self.codes: dict[int, dict] = {}
        self.regels = {}
        for sleutel, regel in kleurregels().items():
            if regel.kleur_achtergrond2:
                kleuren = (_kleur(regel.kleur_achtergrond), _kleur(regel.kleur_achtergrond2))
                vulling = GradientFill(stop=kleuren)
            else:
                vulling = PatternFill("solid", fgColor=_kleur(regel.kleur_achtergrond))
            self.regels[sleutel] = {"fill": vulling, "font": Font(color=_kleur(regel.kleur_tekst))}

    def code(self, code: Dienstcode | None) -> dict:
        if code is None:
            return {}
        if code.id not in self.codes:
            self.codes[code.id] = {
                "fill": PatternFill("solid", fgColor=_kleur(code.kleur_achtergrond)),
                "font": Font(color=_kleur(code.kleur_tekst), bold=code.vet, italic=code.cursief),
            }
        return self.codes[code.id]

    def raster(self, dienst1: Dienst | None, dienst2: Dienst | None) -> dict:
        """Code-cel: kleur van de code; bij twee diensten links/rechts (verloop)."""
        if dienst2 is None:
            return self.code(dienst1.dienstcode if dienst1 else None)
        kleuren = [_kleur(d.dienstcode.kleur_achtergrond) if d is not None and d.dienstcode else "FFFFFF"
                   for d in (dienst1, dienst2)]
        return {"fill": GradientFill(stop=tuple(kleuren)), "font": VET}

    def opmerking(self, tekst: str) -> dict:
        return self.regels.get((tekst or "").strip().casefold(), {})


def _kleur(waarde: str | None) -> str:
    """'#AABBCC' -> 'AABBCC' (openpyxl wil geen '#')."""
    return (waarde or "#FFFFFF").lstrip("#").upper()[:6] or "FFFFFF"


# ---------------------------------------------------------------------------
# De export
# ---------------------------------------------------------------------------

def maak_export(jaar: int, van: date, tot: date, medewerker: Medewerker | None = None) -> bytes:
    """Bouw het .xlsx-bestand en geef de inhoud (bytes)."""
    eerste, laatste = eerste_en_laatste_dag_isojaar(jaar)
    medewerkers = [medewerker] if medewerker is not None else medewerkers_voor_periode(van, tot)
    ids = [m.id for m in medewerkers]
    diensten = (Dienst.query.options(joinedload(Dienst.dienstcode))
                .filter(Dienst.datum >= van, Dienst.datum <= tot, Dienst.medewerker_id.in_(ids)).all())
    per_dag: dict[tuple[int, date], dict[int, Dienst]] = {}
    for dienst in diensten:
        if not dienst.is_leeg or dienst.volgnummer == 1:
            per_dag.setdefault((dienst.medewerker_id, dienst.datum), {})[dienst.volgnummer] = dienst
    dagen_jaar = [eerste + timedelta(days=i) for i in range((laatste - eerste).days + 1)]
    opmerkingen = dagopmerkingen(dagen_jaar)
    stijlen = _Stijlen()

    boek = Workbook()
    boek.remove(boek.active)
    weken = sorted({d.isocalendar()[1] for d in dagen_jaar if van <= d <= tot})
    for week in weken:
        _weekblad(boek.create_sheet(f"W{week}"), jaar, week, van, tot, medewerkers, per_dag,
                  opmerkingen, stijlen)
    _lijsten(boek.create_sheet("Lijsten"), jaar, medewerkers)
    _urenoverzicht(boek.create_sheet("Urenoverzicht"), jaar, medewerker)
    _vakanties(boek.create_sheet("Vakanties"), van, tot)
    kalender = boek.create_sheet("Kalender")
    _zet(kalender, 1, 5, "Jaar", {"font": VET})
    _zet(kalender, 2, 5, jaar)
    _zet(kalender, 1, 1, f"Rooster {jaar}: {van:%d-%m-%Y} t/m {tot:%d-%m-%Y}", {"font": VET})
    uitvoer = io.BytesIO()
    boek.save(uitvoer)
    log.debug("Excel-export: %s weken, %s medewerkers, %s diensten", len(weken), len(medewerkers),
              len(diensten))
    return uitvoer.getvalue()


def _weekblad(blad, jaar, week, van, tot, medewerkers, per_dag, opmerkingen, stijlen) -> None:
    maandag = date.fromisocalendar(jaar, week, 1)
    dagen = [maandag + timedelta(days=i) for i in range(7)]
    _zet(blad, 1, 2, f"Week {week} ({jaar})", {"font": Font(bold=True, size=13)})
    _zet(blad, 2, 2, "Naam", {"font": VET})
    _zet(blad, 2, 26, "Uren", {"font": VET})
    _zet(blad, 1, 28, "Code-raster", {"font": VET})
    _zet(blad, 3, TWEEDE_KOLOMMEN[0], "2e dienst van de dag (bij twee diensten)", {"font": VET})
    for i, dag in enumerate(dagen):
        weekend = {"fill": WEEKEND} if dag.weekday() >= 5 else {"fill": KOP}
        for kolom in (DAG_KOLOMMEN[i], TWEEDE_KOLOMMEN[i]):
            _zet(blad, 1, kolom, DAGNAMEN_KORT[i], {"font": VET, "alignment": OVER_DRIE, **weekend})
            for k in (kolom + 1, kolom + 2):
                _zet(blad, 1, k, None, weekend)
                _zet(blad, 2, k, None, weekend)
            _zet(blad, 2, kolom, datetime.combine(dag, time()),
                 {"number_format": "dd-mm-yy", "font": VET, "alignment": OVER_DRIE, **weekend})
        _zet(blad, 2, CODE_KOLOMMEN[i], DAGNAMEN_KORT[i].upper(), {"font": VET, "alignment": MIDDEN})
        if van <= dag <= tot:
            opm = opmerkingen.get(dag, {}).get("tekst", "")
            _zet(blad, 3, DAG_KOLOMMEN[i], opm or None, {"alignment": OVER_DRIE})

    for n, medewerker in enumerate(medewerkers):
        basis = 4 + 4 * n
        _zet(blad, basis, 2, medewerker.naam, {"font": VET})
        _zet(blad, 6 + 2 * n, 28, medewerker.initialen, {"font": VET})
        contract = medewerker.contracturen_voor(jaar)
        _zet(blad, basis, 26, contract)
        totaal = 0.0
        heeft = False
        for i, dag in enumerate(dagen):
            kolom, k2 = DAG_KOLOMMEN[i], TWEEDE_KOLOMMEN[i]
            # Alleen cellen met inhoud krijgen een rand: lege cellen opmaken kostte het
            # grootste deel van de tijd (en maakt het bestand flink groter)
            if not van <= dag <= tot:
                continue
            per_vn = per_dag.get((medewerker.id, dag), {})
            dienst1, dienst2 = per_vn.get(1), per_vn.get(2)
            if dienst1 is not None and dienst1.is_leeg and dienst2 is None:
                dienst1 = None
            if dienst1 is None and dienst2 is None:
                continue
            heeft = True
            code = matrix_code(dienst1, dienst2)
            if code:
                _zet(blad, 6 + 2 * n, CODE_KOLOMMEN[i], int(code) if code.isdigit() else code,
                     {"alignment": MIDDEN, **stijlen.raster(dienst1, dienst2)})
            if dienst1 is not None:
                opmerking = dienst1.opmerking_tekst
                _zet_als(blad, basis, kolom, opmerking, {**stijlen.opmerking(opmerking), "border": RAND})
                _zet_als(blad, basis + 1, kolom, _tijd(dienst1.opmerking_begin), TIJD)
                _zet_als(blad, basis + 1, kolom + 1, _tijd(dienst1.opmerking_eind), TIJD)
                totaal += _dienst(blad, basis, kolom, dienst1, stijlen)
            if dienst2 is not None:
                totaal += _dienst(blad, basis, k2, dienst2, stijlen)
        if heeft:
            _zet(blad, basis + 2, 26, round(totaal, 2), {"font": VET, "number_format": "0.00"})
    _opmaak_weekblad(blad)


def _dienst(blad, basis: int, kolom: int, dienst: Dienst, stijlen: _Stijlen) -> float:
    """Dienstnaam (regel c) en begin/eind/uren (regel d) van één dienst. Geeft de uren."""
    stijl = stijlen.code(dienst.dienstcode)
    _zet(blad, basis + 2, kolom, dienst.dienstnaam or None, {**stijl, "alignment": OVER_DRIE, "border": RAND})
    for k in (kolom + 1, kolom + 2):  # kleur over de hele breedte van de dag
        _zet(blad, basis + 2, k, None, {**stijl, "alignment": OVER_DRIE, "border": RAND})
    _zet_als(blad, basis + 3, kolom, _tijd(dienst.begin), TIJD)
    _zet_als(blad, basis + 3, kolom + 1, _tijd(dienst.eind), TIJD)
    _zet_als(blad, basis + 3, kolom + 2, dienst.uren_berekend, UREN)
    return dienst.uren_berekend or 0.0


def _opmaak_weekblad(blad) -> None:
    blad.column_dimensions["A"].width = 2
    blad.column_dimensions["B"].width = 22
    for kolom in DAG_KOLOMMEN + TWEEDE_KOLOMMEN:
        for k in range(3):
            blad.column_dimensions[get_column_letter(kolom + k)].width = 7
    blad.column_dimensions["Z"].width = 9
    blad.freeze_panes = "C4"
    blad.sheet_view.zoomScale = 90


def _lijsten(blad, jaar: int, medewerkers: list[Medewerker]) -> None:
    for kolom, kop in ((2, "Initialen"), (3, "Personeel"), (4, f"Contract {jaar}"), (6, "Dienst"),
                       (7, "Omschrijving"), (8, "Van"), (9, "Tot"), (10, "Uren"), (13, "Toeslag")):
        _zet(blad, 1, kolom, kop, {"font": VET, "fill": KOP})
    for rij, medewerker in enumerate(medewerkers, start=2):
        _zet(blad, rij, 2, medewerker.initialen)
        _zet(blad, rij, 3, medewerker.naam)
        _zet(blad, rij, 4, medewerker.contracturen_voor(jaar))
    stijlen = _Stijlen()
    for rij, code in enumerate(Dienstcode.query.order_by(Dienstcode.nummer).all(), start=2):
        _zet(blad, rij, 6, code.nummer)
        _zet(blad, rij, 7, code.omschrijving, stijlen.code(code))
        _zet(blad, rij, 8, _tijd(code.std_begin), {"number_format": "hh:mm"})
        _zet(blad, rij, 9, _tijd(code.std_eind), {"number_format": "hh:mm"})
        _zet(blad, rij, 10, code.std_uren)
    toeslagen = instellingen.toeslagen()
    _zet(blad, 2, 13, "Zaterdag")
    _zet(blad, 2, 14, toeslagen["factor_zaterdag"])
    _zet(blad, 3, 13, "Zondag")
    _zet(blad, 3, 14, toeslagen["factor_zondag"])
    for kolom, breedte in (("B", 10), ("C", 26), ("G", 22)):
        blad.column_dimensions[kolom].width = breedte


def _urenoverzicht(blad, jaar: int, medewerker: Medewerker | None) -> None:
    weken = list(range(1, aantal_weken(jaar) + 1))
    koppen = ["Naam", "Initialen"] + [f"W{w}" for w in weken] + ["Totaal", "Contracturen", "Verschil"]
    for kolom, kop in enumerate(koppen, start=1):
        _zet(blad, 1, kolom, kop, {"font": VET, "fill": KOP})
    rijen = [r for r in uren_overzicht(jaar) if medewerker is None or r.medewerker.id == medewerker.id]
    for rij, r in enumerate(rijen, start=2):
        _zet(blad, rij, 1, r.medewerker.naam)
        _zet(blad, rij, 2, r.medewerker.initialen)
        for kolom, week in enumerate(weken, start=3):
            _zet(blad, rij, kolom, r.per_week.get(week), {"number_format": "0.00"})
        _zet(blad, rij, len(weken) + 3, round(r.gewerkt, 2), {"number_format": "0.00", "font": VET})
        _zet(blad, rij, len(weken) + 4, r.contracturen, {"number_format": "0.00"})
        _zet(blad, rij, len(weken) + 5, None if r.verschil is None else round(r.verschil, 2),
             {"number_format": "0.00"})
    blad.column_dimensions["A"].width = 24
    blad.freeze_panes = "C2"


def _vakanties(blad, van: date, tot: date) -> None:
    for kolom, kop in enumerate(("Vakantie", "Datum van", "Datum tot"), start=1):
        _zet(blad, 1, kolom, kop, {"font": VET, "fill": KOP})
    vakanties = (Vakantie.query.filter(Vakantie.datum_tot >= van, Vakantie.datum_van <= tot)
                 .order_by(Vakantie.datum_van).all())
    for rij, vakantie in enumerate(vakanties, start=2):
        _zet(blad, rij, 1, vakantie.naam)
        _zet(blad, rij, 2, datetime.combine(vakantie.datum_van, time()), {"number_format": "dd-mm-yyyy"})
        _zet(blad, rij, 3, datetime.combine(vakantie.datum_tot, time()), {"number_format": "dd-mm-yyyy"})
    blad.column_dimensions["A"].width = 24
    for kolom in ("B", "C"):
        blad.column_dimensions[kolom].width = 12
