"""Import uit het oude Excel-rooster (.xlsm), zodat lopende roosterdata niet verloren gaat.

Opbouw van het oude bestand:
- Lijsten:   B = initialen, C = naam, D = contracturen (vanaf rij 2)
             F = dienstnummer, G = omschrijving, H = van, I = tot, J = uren
             N2 = toeslag zaterdag, N3 = toeslag zondag
- Vakanties: A = naam, B = van, C = tot (vanaf rij 2)
- Kalender:  E2 = jaar
- W1..W53:   rij 2 = datums (D, G, J, M, P, S, V), rij 3 = dagopmerkingen
             per medewerker een blok van 4 rijen vanaf rij 4:
               +0 opmerking, +1 opmerkingtijden, +2 dienstnaam, +3 begin/eind/uren
             naam in kolom B, contracturen in Z, weektotaal in Z (+2)
             code-raster: AB = initialen, AC..AI = codes ma..zo op rij 6 + 2n
- De waarde 15 (soms als datum 1900-01-15) betekent 'leeg'.

Uitbreidingen voor de eigen Excel-export (app/services/excel_export.py), zodat die weer
ingelezen kan worden; een oud bestand heeft ze niet en wordt gelezen zoals altijd:
- een code-cel met twee codes ('4/7', '/3' of '4/'): dienst 1 en dienst 2;
- tweede dienst per dag in AK..BE (per dag drie kolommen vanaf AK, AN, AQ, ...):
  +2 dienstnaam, +3 begin/eind/uren, net als dienst 1 in D..V;
- het code-raster ook na de 11e medewerker, als AB de initialen van die medewerker heeft;
- sinds 1.6.0 formules (zie excel_export.py). openpyxl bewaart bij formules geen uitkomst; een
  bestand dat nog niet in Excel is opgeslagen heeft dus lege waarden. Een tweede lees-pass
  (zonder data_only) vindt die formulecellen: tijden uit een opzoekformule = de standaardtijden
  van de code, uren uit een formule = niet 'zelf ingevuld', een weektotaal zonder waarde telt
  niet mee in de controle. Is het bestand in Excel opgeslagen, dan gelden de waarden;
- in Lijsten de rekenregels van het bestand: N4 feestdagtoeslag, N5 opmerkingtijden meetellen,
  N6 pauzeaftrek aan, N9:O13 de pauzestaffel, plus het blad Feestdagen. De controles (weektotalen,
  zelf ingevulde uren) rekenen daarmee; een oud bestand: alleen N2/N3 en > 5,5 uur -> 0,5 pauze.

Werkwijze: eerst een droogloop (lees_bestand) met een voorbeeld van wat er gaat
gebeuren, daarna pas definitief importeren (importeer).
Wachtwoorden en rechten uit de bladen 'Beveiliging' en 'Rechten' worden bewust NIET gelezen.
"""

import logging
import math
import os
import re
import time as _time
from dataclasses import asdict, dataclass, field, replace
from datetime import date, datetime, time, timedelta
from types import SimpleNamespace

from flask import current_app
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import joinedload

from ..extensions import db
from ..models import Contracturen, Dagopmerking, Dienst, Dienstcode, Medewerker, Vakantie
from . import instellingen, klok, logboek
from .feestdagen import zorg_voor_jaar
from .kalender import (
    MAX_JAAR,
    MIN_JAAR,
    aantal_weken,
    dagen_van_week,
    eerste_en_laatste_dag_isojaar,
    maandag_van_week,
)
from .medewerkers import uniek_voorstel, voorstel_initialen
from .rooster import UrenContext, markeer_bijgewerkt, uren_voor
from .roosteracties import Actie, Inhoud, Telling, plan_agenda, som, voer_uit
from .tijden import is_cijfers
from .urenberekening import STANDAARD_PAUZE, Staffel, bereken_uren, dagfactor
from .validatie import MAX_NAAM, MAX_OMSCHRIJVING, initialen_fout, is_codenummer
from .voorbeeldpakket import DIENSTCODES
from .weekrooster import automatische_dagopmerkingen, dagopmerkingen

EXCEL_BLANCO = 15
DAG_KOLOMMEN = [4, 7, 10, 13, 16, 19, 22]  # D, G, J, M, P, S, V
CODE_KOLOMMEN = list(range(29, 36))  # AC..AI
TWEEDE_KOLOMMEN = [37 + 3 * i for i in range(7)]  # AK, AN, AQ, AT, AW, AZ, BC: tweede dienst
OUD_CODERASTER = 11  # het oude bestand had een code-raster voor 11 medewerkers
MAX_BLOKKEN = 30
UPLOAD_BEWAREN_SECONDEN = 24 * 3600
# Grenzen tegen 'zip-' en 'XML-bommen' (een klein bestand dat uitgepakt enorm wordt).
# Een echt jaarrooster is uitgepakt een paar MB met zo'n 60 bladen.
MAX_UITGEPAKT = 100 * 1024 * 1024  # bytes, alle onderdelen samen
MAX_ONDERDELEN = 2000
MAX_BLADEN = 120
TEKST_TIJD = re.compile(r"(\d{1,2})[:.](\d{2})(?::\d{2})?")  # een tijd als tekst: '7:15', '07.15', '15:45:00'
log = logging.getLogger(__name__)


class ImportFout(Exception):
    """Het bestand is niet te lezen als oud rooster."""


def import_map() -> str:
    """Map voor geüploade Excel-bestanden (alleen tot de import klaar is)."""
    pad = os.path.join(current_app.config["DATA_MAP"], "import")
    os.makedirs(pad, mode=0o700, exist_ok=True)
    return pad


def ruim_oude_uploads_op(nu: float | None = None) -> int:
    """Verwijder geüploade bestanden die er langer dan een dag staan (import afgebroken)."""
    grens = (nu or _time.time()) - UPLOAD_BEWAREN_SECONDEN
    verwijderd = 0
    for naam in os.listdir(import_map()):
        pad = os.path.join(import_map(), naam)
        if os.path.isfile(pad) and os.path.getmtime(pad) < grens:
            os.remove(pad)
            verwijderd += 1
    return verwijderd


# ---------------------------------------------------------------------------
# Hulpfuncties voor celwaarden
# ---------------------------------------------------------------------------

def _is_leeg(waarde) -> bool:
    """Leeg, of de Excel-'blanco' 15 (ook als datum 1900-01-15)."""
    if waarde is None:
        return True
    if isinstance(waarde, str):
        return waarde.strip() in ("", str(EXCEL_BLANCO))
    if isinstance(waarde, datetime) and waarde.year < 1901:
        return True
    if isinstance(waarde, (int, float)) and waarde == EXCEL_BLANCO:
        return True
    return False


def _tijd(waarde) -> str | None:
    """Excel-tijd -> 'HH:MM' (of None)."""
    if _is_leeg(waarde):
        return None
    if isinstance(waarde, datetime):
        waarde = waarde.time()
    if isinstance(waarde, time):
        minuten = round(waarde.hour * 60 + waarde.minute + waarde.second / 60)
    elif isinstance(waarde, (int, float)) and 0 <= waarde < 1:
        minuten = round(waarde * 1440)  # Excel bewaart tijd als deel van een dag
    elif isinstance(waarde, str) and (gevonden := TEKST_TIJD.fullmatch(waarde.strip())):
        # Een tijd die als tekst in de cel staat ('15:45' of '7.15'); anders telde de dienst
        # zonder uren (zo miste in het rooster van 2026 een dienst van 8 uur)
        uur, minuut = int(gevonden.group(1)), int(gevonden.group(2))
        if uur > 24 or minuut > 59:
            return None
        minuten = uur * 60 + minuut
    else:
        return None
    minuten %= 1440
    return f"{minuten // 60:02d}:{minuten % 60:02d}"


def _tekst(waarde) -> str:
    if _is_leeg(waarde) or isinstance(waarde, (datetime, time)):
        return ""
    if isinstance(waarde, float) and waarde.is_integer():
        waarde = int(waarde)
    return str(waarde).strip()


def _getal(waarde) -> float | None:
    """Getal uit een cel; leeg, tekst, oneindig of NaN -> None."""
    if isinstance(waarde, bool) or waarde is None:
        return None
    try:
        getal = float(waarde) if isinstance(waarde, (int, float)) else float(str(waarde).replace(",", "."))
    except (ValueError, OverflowError):
        return None
    return getal if math.isfinite(getal) else None


def _datum(waarde) -> date | None:
    if isinstance(waarde, datetime):
        return waarde.date()
    if isinstance(waarde, date):
        return waarde
    return None


@dataclass
class Bestandsregels:
    """De rekenregels van het Excel-bestand (zie de docstring bovenaan), voor de controles."""

    zaterdag: float = 1.5
    zondag: float = 2.0
    feestdag: float | None = None
    feestdagen: frozenset = frozenset()
    opmerkingtijden: bool = False
    pauze: Staffel = STANDAARD_PAUZE

    def factor(self, datum: date) -> float:
        return dagfactor(datum, self.zaterdag, self.zondag, datum in self.feestdagen, self.feestdag)

    def uren(self, d) -> float | None:
        """Wat het bestand voor deze dienst berekent (zonder zelf ingevulde uren)."""
        factor = self.factor(d.datum)
        uren = bereken_uren(d.begin, d.eind, factor, self.pauze)
        if self.opmerkingtijden and d.volgnummer == 1:
            extra = bereken_uren(d.opm_begin, d.opm_eind, factor, self.pauze)
            if extra is not None:
                uren = (uren or 0) + extra
        return uren


def _handmatige_uren(d, regels: "Bestandsregels", app_uren: float | None = None) -> float | None:
    """Uren die in Excel met de hand in de urenkolom zijn gezet, of None.

    Dat zijn:
    - uren bij een dienst zonder tijden (bijv. een cursus of 'TW');
    - uren die afwijken van wat de tijden opleveren. Voorbeeld: dienst 07:15-13:00
      met daarna een training 13:00-17:00 op de opmerkingregel; de planner typte
      dan de uren van de hele dag (9,25). Die nemen we over, precies zoals in Excel.
    Komen de uren overeen met wat de app zelf berekent (app_uren, bijvoorbeeld met een
    feestdagtoeslag), dan zijn ze niet met de hand ingevuld (export en weer import). app_uren
    is zonder opmerkingtijden: anders hangt het van die instelling af (zie _gewenste_inhoud).
    regels: de rekenregels van het bestand (het oude Excel kende alleen > 5,5 -> 0,5 pauze).
    Uren uit een formule (export sinds 1.6.0) staan als None in d.excel_uren: niet zelf ingevuld.
    """
    if d.excel_uren is None or (app_uren is not None and abs(app_uren - d.excel_uren) <= 0.001):
        return None
    if not d.begin and not d.eind:
        return d.excel_uren or None
    berekend = regels.uren(d)
    if berekend is not None and abs(berekend - d.excel_uren) > 0.001:
        return d.excel_uren
    return None


def _nieuwe_uren(d, regels: "Bestandsregels") -> float | None:
    """Uren zoals de webapp ze na de import heeft (handmatige uren gaan voor)."""
    handmatig = _handmatige_uren(d, regels)
    if handmatig is not None:
        return handmatig
    return regels.uren(d)


# ---------------------------------------------------------------------------
# Het plan (resultaat van de droogloop)
# ---------------------------------------------------------------------------

@dataclass
class ImportMedewerker:
    naam: str
    initialen: str
    contracturen: float | None
    bestaand_id: int | None = None
    # Hoe deze medewerker gekoppeld wordt: 'naam' (bestaande medewerker met dezelfde naam),
    # 'initialen' (alleen de initialen zijn bezet: er komt een NIEUWE medewerker) of 'nieuw'
    koppeling: str = "nieuw"
    toelichting: str = ""


@dataclass
class ImportCode:
    nummer: int
    omschrijving: str
    begin: str | None
    eind: str | None
    uren: float | None


@dataclass
class ImportDienst:
    naam: str  # naam van de medewerker (sleutel naar ImportMedewerker)
    datum: date
    code: int | None
    dienstnaam: str
    begin: str | None
    eind: str | None
    opmerking: str
    opm_begin: str | None
    opm_eind: str | None
    excel_uren: float | None
    volgnummer: int = 1  # 2 = tweede dienst van de dag


@dataclass
class ImportPlan:
    jaar: int
    jaar_in_bestand: int | None = None  # Kalender!E2 (ter controle)
    toeslag_zaterdag: float | None = None
    toeslag_zondag: float | None = None
    # Overige rekenregels van het bestand (export sinds 1.6.0); een oud bestand: zoals de VBA
    toeslag_feestdag: float | None = None
    feestdagen: frozenset = frozenset()
    opmerkingtijden: bool = False
    pauze: Staffel = STANDAARD_PAUZE
    medewerkers: list[ImportMedewerker] = field(default_factory=list)
    codes: list[ImportCode] = field(default_factory=list)
    vakanties: list[tuple[str, date, date]] = field(default_factory=list)
    diensten: list[ImportDienst] = field(default_factory=list)
    dagopmerkingen: dict[date, str] = field(default_factory=dict)
    weken: list[int] = field(default_factory=list)
    # Controle: weektotalen uit Excel (kolom Z) per (naam, week)
    excel_weektotalen: dict[tuple[str, int], float] = field(default_factory=dict)
    waarschuwingen: list[str] = field(default_factory=list)
    # Naam in Excel -> naam in de app, als die alleen in leestekens/spaties verschilt
    # (zie _bepaal_koppelingen); de weekbladen gebruiken de naam uit Excel
    alias: dict[str, str] = field(default_factory=dict)

    @property
    def regels(self) -> Bestandsregels:
        return Bestandsregels(self.toeslag_zaterdag or 1.5, self.toeslag_zondag or 2.0,
                              self.toeslag_feestdag, self.feestdagen, self.opmerkingtijden, self.pauze)

    @property
    def afwijkend(self) -> int:
        """Aantal diensten waarvan de tijden afwijken van de standaard van de code."""
        standaard = {c.nummer: (c.begin, c.eind) for c in self.codes}
        return sum(1 for d in self.diensten
                   if d.code in standaard and (d.begin, d.eind) != standaard[d.code])

    def handmatige_uren(self) -> list[str]:
        """Diensten waarvan de uren in Excel met de hand afwijken van de tijden.

        Die worden als 'zelf ingevulde uren' overgenomen (zie _handmatige_uren).
        """
        bestand = self.regels
        regels = []
        for d in self.diensten:
            if not (d.begin or d.eind):
                continue
            handmatig = _handmatige_uren(d, bestand)
            if handmatig is not None:
                berekend = bestand.uren(d)
                extra = f", opmerking '{d.opmerking}'" if d.opmerking else ""
                regels.append(f"{d.datum:%d-%m-%Y} {d.naam}: {d.begin}-{d.eind}{extra} – "
                              f"tijden geven {berekend:.2f}, Excel {handmatig:.2f} (overgenomen)")
        return regels

    def dubbele_diensten(self) -> list[str]:
        """(medewerker, datum)-combinaties die meer dan eens in het bestand staan."""
        gezien: dict[tuple[str, date, int], int] = {}
        for d in self.diensten:
            sleutel = (d.naam, d.datum, d.volgnummer)
            gezien[sleutel] = gezien.get(sleutel, 0) + 1
        return [f"{naam} op {datum:%d-%m-%Y}{' (dienst 2)' if vn == 2 else ''} ({aantal}×)"
                for (naam, datum, vn), aantal in sorted(gezien.items(), key=lambda x: (x[0][1], x[0][0]))
                if aantal > 1]

    def weektotaal_verschillen(self) -> list[str]:
        regels = self.regels
        nieuw: dict[tuple[str, int], float] = {}
        for d in self.diensten:
            sleutel = (d.naam, d.datum.isocalendar()[1])
            nieuw[sleutel] = nieuw.get(sleutel, 0) + (_nieuwe_uren(d, regels) or 0)
        verschillen = []
        for (naam, week), excel in sorted(self.excel_weektotalen.items(), key=lambda x: (x[0][1], x[0][0])):
            eigen = nieuw.get((naam, week), 0)
            if abs(eigen - excel) > 0.001:
                verschillen.append(f"W{week} {naam}: Excel {excel:.2f}, nieuw {eigen:.2f}")
        return verschillen


# ---------------------------------------------------------------------------
# Lezen (droogloop)
# ---------------------------------------------------------------------------

def controleer_zip(pad: str) -> None:
    """Controleer het bestand vóór openpyxl het uitpakt (ImportFout als het niet veilig is).

    - een xlsx/xlsm is een zip: uitgepakte grootte en aantal onderdelen zijn begrensd;
    - Excel gebruikt nooit een DOCTYPE of ENTITY in zijn XML; die worden geweigerd
      (zo kan een 'billion laughs'-bestand het geheugen niet vullen);
    - het aantal bladen is begrensd.
    """
    import zipfile

    try:
        with zipfile.ZipFile(pad) as archief:
            onderdelen = archief.infolist()
            if len(onderdelen) > MAX_ONDERDELEN or sum(o.file_size for o in onderdelen) > MAX_UITGEPAKT:
                raise ImportFout("Het bestand is uitgepakt te groot (of bevat te veel onderdelen) "
                                 "voor een rooster; het is niet ingelezen.")
            for onderdeel in onderdelen:
                if onderdeel.filename.lower().endswith((".xml", ".rels", ".vml")):
                    with archief.open(onderdeel) as bestand:
                        begin = bestand.read(64 * 1024).upper()
                    if b"<!DOCTYPE" in begin or b"<!ENTITY" in begin:
                        raise ImportFout("Het bestand is niet veilig om in te lezen (het bevat "
                                         "XML-definities die Excel zelf nooit maakt).")
            bladen = [o for o in onderdelen if re.fullmatch(r"xl/worksheets/[^/]+\.xml", o.filename)]
            if len(bladen) > MAX_BLADEN:
                raise ImportFout(f"Het bestand heeft te veel bladen ({len(bladen)}; hooguit {MAX_BLADEN}).")
    except zipfile.BadZipFile as fout:
        raise ImportFout(f"Het bestand kan niet gelezen worden: {fout}") from fout


def _jaar_in_kalender(boek) -> int | None:
    """Het jaar uit Kalender!E2 (of None als dat er niet is)."""
    if "Kalender" not in boek.sheetnames:
        return None
    waarde = _getal(boek["Kalender"]["E2"].value)
    return int(waarde) if waarde else None


def controleer_bestand(pad: str) -> None:
    """Snelle controle direct na het uploaden: is het een veilig, leesbaar rooster? (ImportFout)"""
    import openpyxl

    controleer_zip(pad)
    try:
        boek = openpyxl.load_workbook(pad, read_only=True, keep_vba=False)
        bladen = boek.sheetnames
        boek.close()
    except Exception as fout:  # elk leesprobleem is een importfout
        raise ImportFout(f"Het bestand kan niet gelezen worden: {fout}") from fout
    if "Lijsten" not in bladen:
        raise ImportFout("Dit lijkt geen rooster: het blad 'Lijsten' ontbreekt.")


def jaar_uit_bestand(pad: str) -> int | None:
    """Voorstel voor het jaar: Kalender!E2, als dat een geldig jaar is. Leest alleen dat blad."""
    import openpyxl

    try:
        controleer_zip(pad)
        boek = openpyxl.load_workbook(pad, data_only=True, read_only=True, keep_vba=False)
        try:
            jaar = _jaar_in_kalender(boek)
        finally:
            boek.close()
    except Exception:  # alleen een voorstel; fouten meldt de droogloop
        return None
    return jaar if jaar is not None and MIN_JAAR <= jaar <= MAX_JAAR else None


def jaar_uit_naam(bestandsnaam: str) -> int | None:
    """Voorstel voor het jaar uit de bestandsnaam, bijv. 'Rooster 2027.xlsm' -> 2027."""
    for kandidaat in re.findall(r"(?<!\d)(\d{4})(?!\d)", bestandsnaam or ""):
        if MIN_JAAR <= int(kandidaat) <= MAX_JAAR:
            return int(kandidaat)
    return None


def lees_bestand(pad: str, jaar: int | None = None) -> ImportPlan:
    """Lees het Excel-bestand en maak een plan voor het gekozen jaar. Er wordt nog niets opgeslagen.

    jaar: het jaar dat de beheerder kiest ('Rooster voor jaar'). Zonder keuze geldt
    Kalender!E2; staat dat er ook niet, dan volgt een ImportFout (nooit stil het
    huidige jaar). Wijkt het gekozen jaar af van E2 of van de datums in de weekbladen,
    dan komt er een waarschuwing; weekbladen van een ander jaar worden overgeslagen.
    """
    import openpyxl

    controleer_zip(pad)
    try:
        # data_only: de laatst berekende waarden (datums, totalen) in plaats van formules
        boek = openpyxl.load_workbook(pad, data_only=True, keep_vba=False)
    except Exception as fout:  # elk leesprobleem is een importfout
        raise ImportFout(f"Het bestand kan niet gelezen worden: {fout}") from fout
    if "Lijsten" not in boek.sheetnames:
        raise ImportFout("Dit lijkt geen oud rooster: het blad 'Lijsten' ontbreekt.")

    in_bestand = _jaar_in_kalender(boek)
    geldig_in_bestand = in_bestand is not None and MIN_JAAR <= in_bestand <= MAX_JAAR
    if jaar is None:
        if in_bestand is not None and not geldig_in_bestand:
            raise ImportFout(f"Het jaar in Kalender!E2 ({in_bestand}) is ongeldig; verwacht een jaar "
                             f"tussen {MIN_JAAR} en {MAX_JAAR}.")
        if in_bestand is None:
            raise ImportFout("Kies voor welk jaar dit rooster is (het bestand heeft geen jaar in "
                             "Kalender!E2).")
        jaar = in_bestand
    if not MIN_JAAR <= jaar <= MAX_JAAR:
        raise ImportFout(f"Kies een jaar tussen {MIN_JAAR} en {MAX_JAAR}.")
    plan = ImportPlan(jaar=jaar, jaar_in_bestand=in_bestand)
    if in_bestand is not None and in_bestand != jaar:
        plan.waarschuwingen.append(
            f"Let op: je importeert voor {jaar}, maar in het bestand (Kalender!E2) staat {in_bestand}. "
            f"Weekbladen met datums uit een ander jaar worden overgeslagen.")
    _lees_lijsten(boek["Lijsten"], plan)
    _bepaal_koppelingen(plan)
    if "Vakanties" in boek.sheetnames:
        _lees_vakanties(boek["Vakanties"], plan)
    if "Feestdagen" in boek.sheetnames:
        plan.feestdagen = frozenset(d for rij in range(2, 400)
                                    if (d := _datum(boek["Feestdagen"].cell(rij, 1).value)))
    # Export sinds 1.6.0: welke cellen zijn formules? (zonder Excel hebben ze geen waarde)
    formules = _formulecellen(pad) if "Rekenhulp" in boek.sheetnames else {}

    weekbladen = sorted(
        (int(naam[1:]), naam) for naam in boek.sheetnames if re.fullmatch(r"W\d{1,2}", naam))
    for week, naam in weekbladen:
        cel = boek[naam].cell(2, DAG_KOLOMMEN[0]).value
        datum = None if _is_leeg(cel) else _datum(cel)
        if datum is not None and datum.isocalendar()[0] != plan.jaar:
            plan.waarschuwingen.append(
                f"{naam} overgeslagen: de datums in het blad ({datum:%d-%m-%Y}) horen bij "
                f"{datum.isocalendar()[0]}, niet bij {plan.jaar}.")
            continue
        if week < 1 or week > aantal_weken(plan.jaar):
            if _blad_heeft_diensten(boek[naam]):
                plan.waarschuwingen.append(f"{naam} overgeslagen: {plan.jaar} heeft geen week {week}.")
            continue
        _lees_weekblad(boek[naam], week, plan, formules.get(naam, set()))
    if dubbel := plan.dubbele_diensten():
        plan.waarschuwingen.append(
            "Dubbele diensten (zelfde medewerker en dag staan er meer dan eens in): "
            + "; ".join(dubbel) + ". Pas het Excel-bestand aan; zo kan het niet geïmporteerd worden.")
    log.debug("Excel gelezen: jaar %s, %s medewerkers, %s codes, %s diensten, %s weken, "
              "%s waarschuwingen", plan.jaar, len(plan.medewerkers), len(plan.codes),
              len(plan.diensten), len(plan.weken), len(plan.waarschuwingen))
    return plan


def _formulecellen(pad: str) -> dict[str, set[tuple[int, int]]]:
    """Per weekblad de cellen (rij, kolom) met een formule, t/m kolom BE (zonder de rekenhulp)."""
    import openpyxl

    boek = openpyxl.load_workbook(pad, read_only=True, keep_vba=False)
    try:
        resultaat = {}
        for naam in boek.sheetnames:
            if re.fullmatch(r"W\d{1,2}", naam):
                rijen = boek[naam].iter_rows(max_col=TWEEDE_KOLOMMEN[-1] + 2)
                resultaat[naam] = {(cel.row, cel.column) for rij in rijen for cel in rij
                                   if getattr(cel, "data_type", "") == "f"}
        return resultaat
    finally:
        boek.close()


def _bepaal_koppelingen(plan: ImportPlan) -> None:
    """Bepaal per Excel-medewerker of hij aan een bestaande medewerker gekoppeld wordt.

    Een gelijke naam koppelt automatisch. Ook een naam die alleen in leestekens, spaties of
    hoofdletters verschilt ('A, Wouw. v.d.' en 'A. Wouw, v.d.'), als precies één medewerker
    zo heet: anders kwam er na het aanpassen van de naam in de app bij elke import een tweede
    medewerker met alle diensten bij (en telden zijn uren dubbel in het overzicht).
    Zijn alleen de initialen gelijk, dan is het waarschijnlijk iemand anders: er komt een
    nieuwe medewerker (met unieke initialen) en de droogloop toont een waarschuwing.
    """
    alle = Medewerker.query.all()
    in_bestand = {im.naam for im in plan.medewerkers}
    for im in plan.medewerkers:
        bestaand = next((m for m in alle if m.naam == im.naam), None)
        if bestaand is not None:
            im.bestaand_id, im.koppeling = bestaand.id, "naam"
            im.toelichting = "bestaande medewerker (zelfde naam)"
            continue
        sleutel = _naamsleutel(im.naam)
        bijna = [m for m in alle if sleutel and _naamsleutel(m.naam) == sleutel and m.naam not in in_bestand]
        if len(bijna) == 1:
            bestaand = bijna[0]
            plan.alias[im.naam] = bestaand.naam
            plan.waarschuwingen.append(
                f"'{im.naam}' uit Excel is gekoppeld aan de bestaande medewerker '{bestaand.naam}' "
                "(de naam verschilt alleen in leestekens of spaties).")
            im.naam = bestaand.naam
            im.bestaand_id, im.koppeling = bestaand.id, "naam"
            im.toelichting = "bestaande medewerker (zelfde naam op leestekens en spaties na)"
            continue
        andere = Medewerker.query.filter_by(initialen=im.initialen).first() if im.initialen else None
        if andere is not None:
            im.koppeling = "initialen"
            im.toelichting = f"nieuw; initialen {im.initialen} zijn al van {andere.naam}"
            plan.waarschuwingen.append(
                f"'{im.naam}' heeft dezelfde initialen ({im.initialen}) als de bestaande medewerker "
                f"'{andere.naam}'. Er wordt een nieuwe medewerker aangemaakt met andere initialen. "
                "Is het dezelfde persoon? Pas dan eerst de naam in de app of in Excel aan.")
        else:
            im.toelichting = "nieuwe medewerker"


def _naamsleutel(naam: str) -> str:
    """'A, Wouw. v.d.' -> 'awouwvd': alleen letters en cijfers, zonder hoofdletters."""
    return re.sub(r"[\W_]", "", naam.casefold())


def _lees_lijsten(blad, plan: ImportPlan) -> None:
    for rij in range(2, 200):
        naam = _tekst(blad.cell(rij, 3).value)[:MAX_NAAM]
        if not naam:
            continue
        # Initialen volgens dezelfde regel als in Beheer (A-Z en 0-9, hooguit 10 tekens)
        ruw = _tekst(blad.cell(rij, 2).value).upper()
        initialen = re.sub(r"[^A-Z0-9]", "", ruw)[:10] or voorstel_initialen(naam)
        if ruw and initialen != ruw:
            plan.waarschuwingen.append(f"Initialen '{ruw}' van '{naam}' zijn ongeldig (alleen letters en "
                                       f"cijfers, maximaal 10); '{initialen}' wordt gebruikt.")
        plan.medewerkers.append(ImportMedewerker(
            naam=naam,
            initialen=initialen if not initialen_fout(initialen) else "",
            contracturen=_getal(blad.cell(rij, 4).value),
        ))
    for rij in range(2, 200):
        nummer = _getal(blad.cell(rij, 6).value)
        omschrijving = _tekst(blad.cell(rij, 7).value)[:MAX_OMSCHRIJVING]
        if nummer is None or int(nummer) == EXCEL_BLANCO or not omschrijving:
            continue
        if nummer != int(nummer) or not is_codenummer(int(nummer)):
            plan.waarschuwingen.append(f"Dienstcode {nummer:g} ({omschrijving}) overgeslagen: het nummer "
                                       "moet een positief geheel getal zijn.")
            continue
        plan.codes.append(ImportCode(
            nummer=int(nummer), omschrijving=omschrijving,
            begin=_tijd(blad.cell(rij, 8).value), eind=_tijd(blad.cell(rij, 9).value),
            uren=_getal(blad.cell(rij, 10).value),
        ))
    plan.toeslag_zaterdag = _getal(blad["N2"].value)
    plan.toeslag_zondag = _getal(blad["N3"].value)
    if _tekst(blad["M6"].value).startswith("Pauzeaftrek"):  # export sinds 1.6.0
        plan.toeslag_feestdag = _getal(blad["N4"].value)
        plan.opmerkingtijden = _getal(blad["N5"].value) == 1
        regels = [(g, a) for rij in range(9, 14)
                  if (g := _getal(blad.cell(rij, 14).value)) is not None
                  and (a := _getal(blad.cell(rij, 15).value)) is not None]
        waarde, fouten = instellingen.controleer_pauze(_getal(blad["N6"].value) == 1, regels)
        if fouten:
            plan.waarschuwingen.append("De pauzeregels in Lijsten zijn ongeldig (" + " ".join(fouten)
                                       + "); de controle rekent met meer dan 5,5 uur -> 0,5 pauze.")
        else:
            plan.pauze = tuple(waarde["regels"]) if waarde["aan"] else ()


def _lees_vakanties(blad, plan: ImportPlan) -> None:
    for rij in range(2, 200):
        naam = _tekst(blad.cell(rij, 1).value)
        van, tot = _datum(blad.cell(rij, 2).value), _datum(blad.cell(rij, 3).value)
        if naam and van and tot:
            plan.vakanties.append((naam, van, tot))


def _blad_heeft_diensten(blad) -> bool:
    return any(not _is_leeg(blad.cell(6 + 2 * n, k).value)
               for n in range(MAX_BLOKKEN // 3) for k in CODE_KOLOMMEN)


def _code(tekst: str) -> int | None:
    getal = _getal(tekst) if tekst.strip() else None
    if getal is None or getal != int(getal) or not is_codenummer(int(getal)):
        return None  # geen geldig codenummer: de dienstnaam telt
    return None if int(getal) == EXCEL_BLANCO else int(getal)


def _lees_codes(waarde) -> tuple[int | None, int | None]:
    """Code-cel -> (code dienst 1, code dienst 2). '4' -> (4, None); '4/7' -> (4, 7); '/3' -> (None, 3)."""
    if _is_leeg(waarde):
        return None, None
    if isinstance(waarde, str) and "/" in waarde:
        eerste, _, tweede = waarde.partition("/")
        return _code(eerste), _code(tweede)
    return _code(str(waarde)), None


def _lees_weekblad(blad, week: int, plan: ImportPlan, formules: set[tuple[int, int]] = frozenset()) -> None:
    cel = blad.cell(2, DAG_KOLOMMEN[0]).value
    maandag = (None if _is_leeg(cel) else _datum(cel)) or maandag_van_week(plan.jaar, week)
    if maandag.isocalendar()[:2] != (plan.jaar, week):
        plan.waarschuwingen.append(
            f"W{week}: datum in het blad ({maandag:%d-%m-%Y}) past niet bij week {week}; "
            "de ISO-week wordt aangehouden.")
        maandag = maandag_van_week(plan.jaar, week)
    dagen = [maandag + timedelta(days=i) for i in range(7)]
    plan.weken.append(week)

    for i, kolom in enumerate(DAG_KOLOMMEN):
        tekst = _tekst(blad.cell(3, kolom).value)
        if tekst:
            plan.dagopmerkingen[dagen[i]] = tekst

    namen = {m.naam for m in plan.medewerkers}
    initialen = {m.naam: m.initialen for m in plan.medewerkers}
    standaard = {c.nummer: (c.begin, c.eind) for c in plan.codes}

    def tijd(rij: int, kolom: int, code: int | None, positie: int) -> str | None:
        """Een tijd; een opzoekformule zonder waarde = de standaardtijd van de code."""
        waarde = blad.cell(rij, kolom).value
        if waarde is None and (rij, kolom) in formules:
            return standaard.get(code, (None, None))[positie]
        return _tijd(waarde)

    def uren(rij: int, kolom: int) -> float | None:
        """Uren uit een formule zijn nooit 'zelf ingevuld' (ook niet met een waarde uit Excel)."""
        return None if (rij, kolom) in formules else _getal(blad.cell(rij, kolom).value)

    for n in range(MAX_BLOKKEN):
        basis = 4 + 4 * n
        naam = _tekst(blad.cell(basis, 2).value)[:MAX_NAAM]
        naam = plan.alias.get(naam, naam)  # zelfde persoon, in de app anders gespeld
        if is_cijfers(naam):  # bijv. 0 uit een formule naar een lege cel
            naam = ""
        if not naam:
            if n > 12 and all(_is_leeg(blad.cell(basis + k, 2).value) for k in range(8)):
                break
            continue
        if naam not in namen:
            plan.waarschuwingen.append(f"W{week}: '{naam}' staat niet in Lijsten; overgeslagen.")
            continue
        excel_totaal = _getal(blad.cell(basis + 2, 26).value)  # kolom Z
        heeft_dienst = False
        # Code-raster: in het oude bestand voor 11 medewerkers; daarna alleen als AB klopt
        met_raster = n < OUD_CODERASTER or (
            initialen.get(naam) and _tekst(blad.cell(6 + 2 * n, 28).value).upper() == initialen[naam])
        for i, kolom in enumerate(DAG_KOLOMMEN):
            code1, code2 = _lees_codes(blad.cell(6 + 2 * n, CODE_KOLOMMEN[i]).value if met_raster else None)
            dienst = ImportDienst(
                naam=naam, datum=dagen[i],
                code=code1,
                dienstnaam=_tekst(blad.cell(basis + 2, kolom).value),
                begin=tijd(basis + 3, kolom, code1, 0),
                eind=tijd(basis + 3, kolom + 1, code1, 1),
                opmerking=_tekst(blad.cell(basis, kolom).value),
                opm_begin=_tijd(blad.cell(basis + 1, kolom).value),
                opm_eind=_tijd(blad.cell(basis + 1, kolom + 1).value),
                excel_uren=uren(basis + 3, kolom + 2),
            )
            if (dienst.code is not None or dienst.dienstnaam or dienst.begin or dienst.eind
                    or dienst.opmerking or dienst.opm_begin or dienst.opm_eind):
                plan.diensten.append(dienst)
                heeft_dienst = True
            # Tweede dienst (alleen in een export uit de app; zie de docstring bovenaan)
            k2 = TWEEDE_KOLOMMEN[i]
            tweede = ImportDienst(
                naam=naam, datum=dagen[i], code=code2, volgnummer=2,
                dienstnaam=_tekst(blad.cell(basis + 2, k2).value),
                begin=tijd(basis + 3, k2, code2, 0),
                eind=tijd(basis + 3, k2 + 1, code2, 1),
                opmerking="", opm_begin=None, opm_eind=None,
                excel_uren=uren(basis + 3, k2 + 2),
            )
            if tweede.code is not None or tweede.dienstnaam or tweede.begin or tweede.eind:
                plan.diensten.append(tweede)
                heeft_dienst = True
        if heeft_dienst and excel_totaal is not None:
            plan.excel_weektotalen[(naam, week)] = excel_totaal


# ---------------------------------------------------------------------------
# Keuzes bij het importeren: wat er overschreven wordt en welke algemene gegevens
# ---------------------------------------------------------------------------

MODUS_ALLES = "alles"
MODUS_GEDEELTELIJK = "gedeeltelijk"
MODUS_AANVULLEN = "aanvullen"
MODI = {
    MODUS_ALLES: "Alles in het gekozen jaar",
    MODUS_GEDEELTELIJK: "Gedeeltelijk: alleen gekozen medewerkers en/of periode",
    MODUS_AANVULLEN: "Alleen lege dagen aanvullen (niets overschrijven)",
}


@dataclass
class ImportKeuzes:
    """Wat de import mag wijzigen. Standaard: zie standaard().

    modus:         alles / gedeeltelijk / aanvullen (zie MODI)
    medewerkers:   namen uit het plan (leeg = iedereen); niet bij 'alles'
    van, tot:      periode binnen het gekozen jaar (leeg = het hele jaar); niet bij 'alles'
    dagopmerkingen_bij_selectie: dagopmerkingen ook overnemen als er medewerkers gekozen zijn
    contracturen, toeslagen, vakanties, codes: algemene gegevens overnemen (ja/nee)
    """

    modus: str = MODUS_ALLES
    medewerkers: tuple[str, ...] = ()
    van: date | None = None
    tot: date | None = None
    dagopmerkingen_bij_selectie: bool = False
    contracturen: bool = True
    toeslagen: bool = True
    vakanties: bool = True
    codes: bool = True

    @classmethod
    def standaard(cls, plan: "ImportPlan") -> "ImportKeuzes":
        """Toeslagen gelden voor alle jaren: alleen standaard overnemen voor het huidige jaar."""
        return cls(toeslagen=plan.jaar == klok.vandaag().year)

    def genormaliseerd(self) -> "ImportKeuzes":
        """Bij 'alles' tellen de filters niet mee."""
        if self.modus == MODUS_ALLES:
            return replace(self, medewerkers=(), van=None, tot=None)
        return replace(self, medewerkers=tuple(sorted(set(self.medewerkers))))

    def controleer(self, plan: "ImportPlan") -> list[str]:
        """Foutmeldingen (leeg = in orde). Wordt bij het opslaan opnieuw gedaan."""
        fouten = []
        if self.modus not in MODI:
            fouten.append("Kies een geldige overschrijfmodus.")
        namen = {m.naam for m in plan.medewerkers}
        if onbekend := [n for n in self.medewerkers if n not in namen]:
            fouten.append("Onbekende medewerker(s): " + ", ".join(onbekend) + ".")
        eerste, laatste = eerste_en_laatste_dag_isojaar(plan.jaar)
        for datum in (self.van, self.tot):
            if datum is not None and not eerste <= datum <= laatste:
                fouten.append(f"De periode moet binnen {plan.jaar} vallen "
                              f"({eerste:%d-%m-%Y} t/m {laatste:%d-%m-%Y}).")
                break
        if self.van and self.tot and self.van > self.tot:
            fouten.append("De begindatum van de periode ligt na de einddatum.")
        return fouten

    def als_dict(self) -> dict:
        """Voor in de sessie (alleen JSON-waarden)."""
        gegevens = asdict(self)
        gegevens["medewerkers"] = list(self.medewerkers)
        gegevens["van"] = self.van.isoformat() if self.van else ""
        gegevens["tot"] = self.tot.isoformat() if self.tot else ""
        return gegevens

    def beschrijving(self, jaar: int) -> str:
        """Voor het logboek: jaar, modus, medewerkers en periode."""
        delen = [f"jaar {jaar}", f"modus {self.modus}"]
        if self.modus != MODUS_ALLES:
            delen.append("medewerkers: " + (", ".join(self.medewerkers) or "alle"))
            van = f"{self.van:%d-%m-%Y}" if self.van else "begin jaar"
            tot = f"{self.tot:%d-%m-%Y}" if self.tot else "eind jaar"
            delen.append(f"periode {van} t/m {tot}")
        uit = [naam for naam in ("contracturen", "toeslagen", "vakanties", "codes")
               if not getattr(self, naam)]
        if uit:
            delen.append("niet overgenomen: " + ", ".join(uit))
        return ", ".join(delen)


# ---------------------------------------------------------------------------
# Het effect van de import (droogloop) en het definitief importeren
# ---------------------------------------------------------------------------

@dataclass
class ImportEffect:
    """Wat de import precies gaat doen met de gekozen keuzes (voor de droogloop)."""

    per_medewerker: dict[str, Telling] = field(default_factory=dict)
    dagopmerkingen: list[tuple[date, str, str]] = field(default_factory=list)  # (dag, oud, nieuw)
    toeslagen: list[tuple[str, str, str]] = field(default_factory=list)  # (sleutel, oud, nieuw)
    nieuwe_codes: list[ImportCode] = field(default_factory=list)
    vakanties: list[tuple[str, date, date]] = field(default_factory=list)
    contracturen: list[tuple[str, float | None, float]] = field(default_factory=list)
    nieuwe_medewerkers: list[str] = field(default_factory=list)
    acties: list[Actie] = field(default_factory=list)

    @property
    def totaal(self) -> Telling:
        return som(self.per_medewerker.values())


def _bereik(plan: ImportPlan, keuzes: ImportKeuzes) -> tuple[list[str], list[date]]:
    """(namen, dagen) die de import mag raken.

    Dagen: alleen de weken die in het bestand staan (H1), binnen de gekozen periode.
    """
    eerste, laatste = eerste_en_laatste_dag_isojaar(plan.jaar)
    van, tot = keuzes.van or eerste, keuzes.tot or laatste
    dagen = sorted({d for week in plan.weken for d in dagen_van_week(plan.jaar, week) if van <= d <= tot})
    namen = [m.naam for m in plan.medewerkers if not keuzes.medewerkers or m.naam in keuzes.medewerkers]
    return namen, dagen


def _codes_na_import(plan: ImportPlan, keuzes: ImportKeuzes) -> dict[int, tuple[str, str | None, str | None]]:
    """Dienstcodes zoals ze er na de import zijn: nummer -> (omschrijving, std_begin, std_eind)."""
    codes = {c.nummer: (c.omschrijving, c.std_begin, c.std_eind) for c in Dienstcode.query.all()}
    if keuzes.codes:
        for ic in plan.codes:
            codes.setdefault(ic.nummer, (ic.omschrijving, ic.begin, ic.eind))
    return codes


def _gewenste_inhoud(d: ImportDienst, codes: dict, per_naam: dict, regels: Bestandsregels,
                     context: UrenContext) -> Inhoud:
    """De inhoud die een dienst uit het bestand in de app krijgt."""
    from .weekrooster import begint_met_dienstnaam

    nummer = d.code if d.code in codes else None
    if nummer is None and d.dienstnaam:
        # Geen (bekende) code maar wel een naam: code op naam zoeken als die uniek is
        kandidaten = per_naam.get(d.dienstnaam.casefold(), [])
        nummer = kandidaten[0] if len(kandidaten) == 1 else None
    dienstnaam = d.dienstnaam[:60]
    if nummer is not None:
        omschrijving = codes[nummer][0]
        # Een aanvulling achter de dienstnaam ('VW Vroeg tot 12:00') blijft bewaard
        dienstnaam = dienstnaam if (dienstnaam.casefold() != omschrijving.casefold()
                                    and begint_met_dienstnaam(dienstnaam, omschrijving)) else ""
    # Wat de app zelf voor deze dienst berekent (met feestdagen en instellingen), zonder de
    # opmerkingtijden: die instelling kan later uit. Typte de planner in Excel het dagtotaal
    # (dienst + training op de opmerkingregel) en telde de app dat bij de import toevallig ook,
    # dan raakten die uren anders kwijt zodra 'Opmerkingtijden meetellen' uit ging (4 uur bij
    # J. Hoskam op 17-11-2026).
    proef = SimpleNamespace(uren_handmatig=None, datum=d.datum, begin=d.begin, eind=d.eind,
                            opmerking_begin=None, opmerking_eind=None)
    uren = _handmatige_uren(d, regels, app_uren=uren_voor(proef, context))
    if d.volgnummer == 1:
        opmerking = (d.opmerking[:120], d.opm_begin, d.opm_eind)
    else:
        opmerking = ("", None, None)  # de opmerking hoort bij de dag (dienst 1)
    return Inhoud(nummer, dienstnaam, d.begin, d.eind, *opmerking, uren)


def effect(plan: ImportPlan, keuzes: ImportKeuzes | None = None) -> ImportEffect:
    """Bereken precies wat de import met deze keuzes doet. Er wordt niets opgeslagen.

    importeer() voert exact deze acties uit, zodat de droogloop klopt met het resultaat.
    """
    keuzes = (keuzes or ImportKeuzes.standaard(plan)).genormaliseerd()
    resultaat = ImportEffect()
    namen, dagen = _bereik(plan, keuzes)
    dagenset = set(dagen)
    for naam in namen:
        resultaat.per_medewerker[naam] = Telling()

    bestaande_mw = {m.naam: m for m in Medewerker.query.filter(Medewerker.naam.in_(namen)).all()} \
        if namen else {}
    resultaat.nieuwe_medewerkers = [n for n in namen if n not in bestaande_mw]

    # Algemene gegevens
    if keuzes.toeslagen:
        for sleutel, waarde in (("toeslag_zaterdag", plan.toeslag_zaterdag),
                                ("toeslag_zondag", plan.toeslag_zondag)):
            if waarde and instellingen.lees(sleutel) != str(waarde):
                resultaat.toeslagen.append((sleutel, instellingen.lees(sleutel), str(waarde)))
    if keuzes.codes:
        bestaande_codes = {nummer for (nummer,) in db.session.query(Dienstcode.nummer)}
        resultaat.nieuwe_codes = [c for c in plan.codes if c.nummer not in bestaande_codes]
    if keuzes.vakanties:
        jaar_van, jaar_tot = date(plan.jaar, 1, 1), date(plan.jaar, 12, 31)
        for naam, van, tot in plan.vakanties:
            if van <= jaar_tot and tot >= jaar_van \
                    and not Vakantie.query.filter_by(naam=naam, datum_van=van).first():
                resultaat.vakanties.append((naam, van, tot))
    if keuzes.contracturen:
        for im in plan.medewerkers:
            if im.naam not in namen or im.contracturen is None:
                continue
            medewerker = bestaande_mw.get(im.naam)
            oud = next((c.uren for c in medewerker.contracturen if c.jaar == plan.jaar), None) \
                if medewerker else None
            if oud != im.contracturen:
                resultaat.contracturen.append((im.naam, oud, im.contracturen))
    if not dagen or not namen:
        return resultaat

    # Diensten
    regels = plan.regels
    codes = _codes_na_import(plan, keuzes)
    per_naam: dict[str, list[int]] = {}
    for nummer, (omschrijving, _, _) in codes.items():
        per_naam.setdefault(omschrijving.casefold(), []).append(nummer)
    context = UrenContext(dagen[0], dagen[-1])
    gewenst: dict[tuple[str, date], dict[int, Inhoud]] = {}
    for d in plan.diensten:
        if d.naam in resultaat.per_medewerker and d.datum in dagenset:
            gewenst.setdefault((d.naam, d.datum), {})[d.volgnummer] = \
                _gewenste_inhoud(d, codes, per_naam, regels, context)
    naam_van_id = {m.id: m.naam for m in bestaande_mw.values()}
    bestaand: dict[tuple[str, date], dict[int, Dienst]] = {}
    if naam_van_id:
        for dienst in (Dienst.query.options(joinedload(Dienst.dienstcode))
                       .filter(Dienst.medewerker_id.in_(naam_van_id),
                               Dienst.datum >= dagen[0], Dienst.datum <= dagen[-1]).all()):
            if dienst.datum in dagenset:
                bestaand.setdefault((naam_van_id[dienst.medewerker_id], dienst.datum), {})[
                    dienst.volgnummer] = dienst

    for sleutel in sorted(set(gewenst) | set(bestaand), key=lambda s: (s[1], namen.index(s[0]))):
        naam, datum = sleutel
        oud, nieuw = bestaand.get(sleutel, {}), gewenst.get(sleutel, {})
        gevuld = {vn for vn, dienst in oud.items() if not dienst.is_leeg}
        telling = resultaat.per_medewerker[naam]
        for vn in sorted(set(oud) | set(nieuw)):
            dienst, inhoud = oud.get(vn), nieuw.get(vn)
            if keuzes.modus == MODUS_AANVULLEN and gevuld:
                soort = "overgeslagen" if inhoud else ""
            elif inhoud is None:
                soort = "verwijderd" if vn in gevuld else ""
            elif vn not in gevuld:
                soort = "nieuw"
            else:
                soort = "gelijk" if Inhoud.van_dienst(dienst) == inhoud else "vervangen"
            if soort:
                setattr(telling, soort, getattr(telling, soort) + 1)
                resultaat.acties.append(Actie(naam, datum, vn, soort, dienst, inhoud))

    # Dagopmerkingen: zelfde periode; bij gekozen medewerkers alleen als dat aangevinkt is
    if not keuzes.medewerkers or keuzes.dagopmerkingen_bij_selectie:
        huidig = dagopmerkingen(dagen)
        for dag, tekst in sorted(plan.dagopmerkingen.items()):
            if dag not in dagenset or huidig[dag]["tekst"] == tekst:
                continue
            if keuzes.modus == MODUS_AANVULLEN and huidig[dag]["handmatig"]:
                continue
            resultaat.dagopmerkingen.append((dag, huidig[dag]["tekst"], tekst))
    return resultaat


def importeer(plan: ImportPlan, keuzes: ImportKeuzes | None = None) -> dict:
    """Schrijf het plan naar de database, in één transactie (alles of niets).

    Alleen het bereik uit de keuzes wordt geraakt (zie effect): andere weken, jaren,
    medewerkers en dagen blijven precies zoals ze zijn. Bestaande medewerkers (zelfde
    naam) en codes (zelfde nummer) worden hergebruikt; bestaande codes nooit gewijzigd.
    Gaat er iets mis, dan wordt alles teruggedraaid en volgt een ImportFout.
    """
    keuzes = (keuzes or ImportKeuzes.standaard(plan)).genormaliseerd()
    if dubbel := plan.dubbele_diensten():
        raise ImportFout("Er is niets geïmporteerd: dubbele diensten in het bestand ("
                         + "; ".join(dubbel) + ").")
    if fouten := keuzes.controleer(plan):
        raise ImportFout("Er is niets geïmporteerd: " + " ".join(fouten))
    try:
        resultaat, geraakt = _importeer(plan, keuzes)
        db.session.commit()
    except IntegrityError as fout:
        db.session.rollback()
        raise ImportFout("Er is niets geïmporteerd: het bestand bevat gegevens die botsen "
                         f"(bijvoorbeeld een dubbele dienst). Details: {fout.orig}") from fout
    except Exception:
        db.session.rollback()
        raise
    plan_agenda(geraakt, "de import")
    log.info("Excel-import klaar: %s", resultaat)
    return resultaat


def _importeer(plan: ImportPlan, keuzes: ImportKeuzes) -> tuple[dict, dict[Medewerker, set[date]]]:
    """Het eigenlijke importeren; er wordt hier nergens gecommit."""
    # Feestdagen vooraf aanmaken (zonder commit), zodat de berekeningen hieronder
    # nooit halverwege iets opslaan. Een week kan over de jaargrens lopen.
    for jaar in (plan.jaar - 1, plan.jaar, plan.jaar + 1):
        zorg_voor_jaar(jaar, commit=False)
    uitkomst = effect(plan, keuzes)
    resultaat = {"medewerkers": 0, "codes": 0, "vakanties": 0, "diensten": 0,
                 "vervangen": 0, "verwijderd": 0, "dagopmerkingen": 0}

    for sleutel, oud, nieuw in uitkomst.toeslagen:
        logboek.log("Instelling gewijzigd", "Excel-import", veld=sleutel, oud=oud, nieuw=nieuw)
        instellingen.schrijf(sleutel, nieuw)

    # Nieuwe dienstcodes (kleuren uit het voorbeeldpakket als het nummer daar in staat)
    pakket = {c[0]: c for c in DIENSTCODES}
    for ic in uitkomst.nieuwe_codes:
        uit_pakket = pakket.get(ic.nummer)
        db.session.add(Dienstcode(
            nummer=ic.nummer, omschrijving=ic.omschrijving[:60], std_begin=ic.begin,
            std_eind=ic.eind, std_uren=ic.uren,
            kleur_achtergrond=uit_pakket[5] if uit_pakket else "#D9D9D9",
            kleur_tekst=uit_pakket[6] if uit_pakket else "#000000",
            vet=True, cursief=uit_pakket[7] if uit_pakket else False,
            hele_dag_zonder_tijden=ic.begin is None,
        ))
        logboek.log("Dienstcode toegevoegd", f"Excel-import: {ic.nummer} {ic.omschrijving}")
        resultaat["codes"] += 1
    db.session.flush()

    # Medewerkers: alleen op naam hergebruiken (zie _bepaal_koppelingen)
    medewerkers: dict[str, Medewerker] = {}
    volgorde = db.session.query(db.func.max(Medewerker.volgorde)).scalar() or 0
    for im in plan.medewerkers:
        if im.naam not in uitkomst.per_medewerker:
            continue  # niet gekozen
        medewerker = Medewerker.query.filter_by(naam=im.naam).first()
        if medewerker is None:
            volgorde += 1
            initialen = im.initialen if im.initialen and not Medewerker.query.filter_by(
                initialen=im.initialen).first() else uniek_voorstel(im.naam)
            medewerker = Medewerker(naam=im.naam[:120], initialen=initialen or f"M{volgorde}",
                                    volgorde=volgorde)
            db.session.add(medewerker)
            db.session.flush()
            logboek.log("Medewerker toegevoegd", "Excel-import", medewerker=medewerker.naam)
            resultaat["medewerkers"] += 1
        medewerkers[im.naam] = medewerker
    for naam, oud, nieuw in uitkomst.contracturen:
        medewerker = medewerkers[naam]
        bestaand = next((c for c in medewerker.contracturen if c.jaar == plan.jaar), None)
        if bestaand:
            bestaand.uren = nieuw
        else:
            medewerker.contracturen.append(Contracturen(jaar=plan.jaar, uren=nieuw))
        logboek.log("Contracturen gewijzigd", f"Excel-import, {plan.jaar}", medewerker=naam,
                    veld="contracturen", oud="" if oud is None else oud, nieuw=nieuw)

    for naam, van, tot in uitkomst.vakanties:
        db.session.add(Vakantie(naam=naam[:80], datum_van=van, datum_tot=tot))
        logboek.log("Vakantie toegevoegd", f"Excel-import: {naam} {van:%d-%m-%Y} t/m {tot:%d-%m-%Y}")
        resultaat["vakanties"] += 1
    db.session.flush()

    # Diensten: per actie uit de droogloop (gedeeld met roosterpatronen, zie roosteracties.py)
    aantallen, geraakt = voer_uit(uitkomst.acties, lambda actie: medewerkers[actie.sleutel], "Excel-import",
                                  lambda actie: f"Dienst {actie.soort}")
    resultaat["diensten"] = aantallen["nieuw"]
    resultaat["vervangen"] = aantallen["vervangen"]
    resultaat["verwijderd"] = aantallen["verwijderd"]

    # Dagopmerkingen: alleen bewaren als ze afwijken van de automatische tekst
    if uitkomst.dagopmerkingen:
        dagen = [dag for dag, _, _ in uitkomst.dagopmerkingen]
        automatisch = automatische_dagopmerkingen(min(dagen), max(dagen))
        for dag, oud, tekst in uitkomst.dagopmerkingen:
            bestaand = Dagopmerking.query.filter_by(datum=dag).first()
            if automatisch.get(dag, "") == tekst:
                if bestaand:
                    db.session.delete(bestaand)
            elif bestaand:
                bestaand.tekst = tekst[:120]
            else:
                db.session.add(Dagopmerking(datum=dag, tekst=tekst[:120], handmatig=True))
            logboek.log("Dagopmerking gewijzigd", "Excel-import", datum=dag, veld="dagopmerking",
                        oud=oud, nieuw=tekst)
            resultaat["dagopmerkingen"] += 1
            # De dagopmerking staat in de omschrijving van agenda-afspraken
            for dienst in Dienst.query.filter(Dienst.datum == dag, Dienst.google_event_id != "").all():
                geraakt.setdefault(dienst.medewerker, set()).add(dag)

    logboek.log("Excel-import", keuzes.beschrijving(plan.jaar) + "; "
                + ", ".join(f"{k}: {v}" for k, v in resultaat.items()))
    markeer_bijgewerkt()
    db.session.flush()
    return resultaat, geraakt
